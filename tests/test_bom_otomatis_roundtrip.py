"""Round-trip BOM_AKSESORIS + BOM_OTOMATIS (berkas FOKUS) lewat API nyata."""
import io
import os
import sys

import openpyxl
import requests

API = os.environ.get("API_URL") or open("/app/frontend/.env").read().split("REACT_APP_BACKEND_URL=")[1].splitlines()[0].strip()
TOKEN = open("/tmp/tok").read().strip()
H = {"Authorization": f"Bearer {TOKEN}"}


def post_xlsx(path, data=None, **params):
    files = {"file": ("x.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")} if data else None
    r = requests.post(f"{API}{path}", headers=H, files=files, params=params, timeout=120)
    assert r.status_code == 200, (path, r.status_code, r.text[:300])
    return r


# 1. FOKUS tanpa berkas → sheet BOM_AKSESORIS & BOM_OTOMATIS ada
wb = openpyxl.load_workbook(io.BytesIO(post_xlsx("/api/rahaza/master/gap-fokus").content))
print("sheets:", wb.sheetnames)
assert "BOM_AKSESORIS" in wb.sheetnames and "BOM_OTOMATIS" in wb.sheetnames
ws = wb["BOM_AKSESORIS"]
model_rows = [r for r in ws.iter_rows(values_only=True, min_row=2) if r[0]]
print("BOM_AKSESORIS baris model tanpa aksesoris:", len(model_rows))

# 2. Bangun berkas unggahan: 1 kelompok di BOM_AKSESORIS + 1 kelompok berbeda di BOM_OTOMATIS, pakai model & material nyata
ref = list(wb["REF_AKSESORIS"].iter_rows(values_only=True, min_row=2))
acc = [r for r in ref if r[4] and float(r[4]) > 0 and (r[3] or "").lower() == "pcs"][:2]
assert len(acc) == 2, "butuh 2 aksesoris pcs berharga"
models_in_file = sorted({r[0] for r in model_rows})[:2]
assert len(models_in_file) == 2, models_in_file
up = openpyxl.Workbook()
w1 = up.active
w1.title = "BOM_AKSESORIS"
hdr = ["kode_model", "nama_model", "kode_material", "nama_material", "qty_per_pcs", "satuan", "keterangan", "varian", "varian_tersedia", "yang_perlu_diisi"]
w1.append(hdr)
w1.append([models_in_file[0], "", acc[0][0], acc[0][1], "1 pcs", "", "", "", "", ""])
w1.append(["", "", acc[1][0], acc[1][1], "2", "pcs", "", "", "", ""])
w1.append(["", "", "KODE-TIDAK-ADA-XYZ", "salah", "1", "pcs", "", "", "", ""])  # → issue kode_tak_dikenal (BOM_AKSESORIS baris 4)
w2 = up.create_sheet("BOM_OTOMATIS")
w2.append(hdr)
w2.append([models_in_file[1], "", acc[0][0], acc[0][1], "3 pcs", "", "", "", "", ""])
w2.append(["", "", "KODE-TIDAK-ADA-ABC", "salah", "1", "pcs", "", "", "", ""])  # → issue di BOM_OTOMATIS baris 3
buf = io.BytesIO()
up.save(buf)
data = buf.getvalue()

# 3. Pratinjau: kedua kelompok terbaca, label sumber benar
p = post_xlsx("/api/rahaza/master/fill-preview", data).json()
print("totals:", p["totals"])
print("warnings:", p["warnings"])
groups = {g["model_code"] for g in p["bom_groups"]}
assert groups == set(models_in_file), (groups, models_in_file)
assert p["totals"]["bom_lines"] == 3, p["totals"]
assert any(w.startswith("BOM_AKSESORIS baris 4:") for w in p["warnings"]), p["warnings"]
assert any(w.startswith("BOM_OTOMATIS baris 3:") for w in p["warnings"]), p["warnings"]
assert any(i.get("where") == "BOM_OTOMATIS baris 3" for i in p["bom_issues"])

# 4. FOKUS dari berkas ini: baris bermasalah muncul lagi (kode_material kuning → BOM_AKSESORIS)
wb2 = openpyxl.load_workbook(io.BytesIO(post_xlsx("/api/rahaza/master/gap-fokus", data).content))
rows_aks = [r for r in wb2["BOM_AKSESORIS"].iter_rows(values_only=True, min_row=2)]
rows_oto = [r for r in wb2["BOM_OTOMATIS"].iter_rows(values_only=True, min_row=2)]
codes_aks = {r[2] for r in rows_aks}
print("FOKUS ulang — BOM_AKSESORIS:", len(rows_aks), "BOM_OTOMATIS:", len(rows_oto))
assert "KODE-TIDAK-ADA-XYZ" in codes_aks and "KODE-TIDAK-ADA-ABC" in codes_aks, codes_aks
# kelompok bermasalah ditulis utuh (baris benar ikut) dan model tidak muncul lagi sebagai 'tanpa aksesoris'
assert acc[0][0] in codes_aks
assert sum(1 for r in rows_aks if r[0] == models_in_file[0]) == 1

# 5. SISA (gap-workbook) juga membaca BOM_OTOMATIS
wb3 = openpyxl.load_workbook(io.BytesIO(post_xlsx("/api/rahaza/master/gap-workbook", data).content))
sisa = [r for r in wb3["BOM_AKSESORIS"].iter_rows(values_only=True, min_row=2)]
assert any(r[2] == "KODE-TIDAK-ADA-ABC" for r in sisa), "kelompok BOM_OTOMATIS harus ikut di Laporan SISA"

# 6. Terapkan (scope=bom) → kedua model dapat aksesoris; idempoten
a1 = post_xlsx("/api/rahaza/master/fill-apply", data, scope="bom").json()
print("apply#1:", {k: a1[k] for k in ("bom_models", "bom_groups", "boms_touched", "bom_lines_appended", "boms_unchanged")})
assert a1["bom_models"] == 2 and a1["boms_touched"] > 0
a2 = post_xlsx("/api/rahaza/master/fill-apply", data, scope="bom").json()
print("apply#2:", {k: a2[k] for k in ("boms_touched", "boms_unchanged")})
assert a2["boms_touched"] == 0 and a2["boms_unchanged"] == a1["boms_touched"] + a1["boms_unchanged"]
print("OK — round-trip BOM_AKSESORIS + BOM_OTOMATIS lolos")
