"""FOKUS: kelompok tanpa varian yang variannya bisa disarankan dari nama bahan pembeda → sheet BOM_OTOMATIS (biru)."""
import io
import os

import openpyxl
import requests
from pymongo import MongoClient

API = open("/app/frontend/.env").read().split("REACT_APP_BACKEND_URL=")[1].splitlines()[0].strip()
H = {"Authorization": f"Bearer {open('/tmp/tok').read().strip()}"}
env = dict(l.split("=", 1) for l in open("/app/backend/.env").read().splitlines() if "=" in l)
db = MongoClient(env["MONGO_URL"].strip('"'))[env["DB_NAME"].strip('"')]

# model dengan ≥2 warna aktif
pipe = [{"$match": {"active": True}}, {"$group": {"_id": "$model_id", "colors": {"$addToSet": "$color_name"}}}, {"$match": {"colors.1": {"$exists": True}}}, {"$limit": 1}]
grp = next(db.rahaza_model_variants.aggregate(pipe))
model = db.rahaza_models.find_one({"id": grp["_id"]}, {"_id": 0, "code": 1, "name": 1})
c1, c2 = [c for c in grp["colors"] if c][:2]
acc = list(db.rahaza_materials.find({"type": {"$nin": ["fg", "fabric"]}, "unit": "pcs", "unit_cost": {"$gt": 0}, "active": {"$ne": False}}, {"_id": 0, "code": 1, "name": 1}).limit(2))
print("model", model["code"], "warna", c1, "/", c2, "acc", [a["code"] for a in acc])

up = openpyxl.Workbook()
w = up.active
w.title = "BOM_AKSESORIS"
w.append(["kode_model", "nama_model", "kode_material", "nama_material", "qty_per_pcs", "satuan", "keterangan", "varian", "varian_tersedia"])
w.append([model["code"], model["name"], acc[0]["code"], f"Kancing warna {c1}", "1 pcs", "", "", "", ""])
w.append([model["code"], model["name"], acc[1]["code"], f"Kancing warna {c2}", "1 pcs", "", "", "", ""])
buf = io.BytesIO()
up.save(buf)
data = buf.getvalue()
files = {"file": ("x.xlsx", data, "application/octet-stream")}
p = requests.post(f"{API}/api/rahaza/master/fill-preview", headers=H, files=files, timeout=120).json()
print("issues:", [(i["kategori"], i.get("where")) for i in p["bom_issues"]])
assert p["bom_issue_counts"].get("kelompok_tanpa_varian") == 2, p["bom_issue_counts"]

r = requests.post(f"{API}/api/rahaza/master/gap-fokus", headers=H, files={"file": ("x.xlsx", data, "application/octet-stream")}, timeout=120)
wb = openpyxl.load_workbook(io.BytesIO(r.content))
oto = [x for x in wb["BOM_OTOMATIS"].iter_rows(min_row=2) if x[0].value]
aks = [x for x in wb["BOM_AKSESORIS"].iter_rows(min_row=2) if x[0].value == model["code"]]
print("BOM_OTOMATIS:", [(x[0].value, x[7].value, x[7].fill.fgColor.rgb, x[9].value[:60]) for x in oto])
print("BOM_AKSESORIS model ini:", len(aks))
assert len(oto) == 2 and all(x[7].fill.fgColor.rgb.endswith("DDEBF7") for x in oto), "kelompok varian-disarankan harus di BOM_OTOMATIS dengan sel biru"
assert {x[7].value for x in oto} == {c1, c2}
assert len(aks) == 0

# unggah balik berkas FOKUS apa adanya → kelompok BOM_OTOMATIS terbaca dengan varian tersaran
p2 = requests.post(f"{API}/api/rahaza/master/fill-preview", headers=H, files={"file": ("f.xlsx", r.content, "application/octet-stream")}, timeout=120).json()
mine = [g for g in p2["bom_groups"] if g["model_code"] == model["code"]]
print("groups setelah unggah balik FOKUS:", [(g["model_code"], g["target_label"], g["where"]) for g in mine])
assert len(mine) == 2 and all(g["where"].startswith("BOM_OTOMATIS") for g in mine)
print("OK — saran varian → BOM_OTOMATIS → unggah balik terbaca")
