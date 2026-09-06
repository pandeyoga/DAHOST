"""Regression and safety tests for master client package scripts (dry-run only)."""

import asyncio
import hashlib
import json
import os
import subprocess
import sys
from pathlib import Path

import pytest
import requests
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS_DIR = ROOT / "scripts"
BACKEND_DIR = ROOT / "backend"

sys.path.insert(0, str(SCRIPTS_DIR))
sys.path.insert(0, str(BACKEND_DIR))

import build_master_client_package as bmp  # noqa: E402
import import_master_template as imt  # noqa: E402
from master_autofix_rules import arithmetic  # noqa: E402
from master_client_documents import write_answers, write_questions  # noqa: E402
from master_client_questions import build_questions, catalog_analysis  # noqa: E402
from master_import_bom import convert_line  # noqa: E402
from master_review_safety import ReadOnlyCollection, ReadOnlyDatabase, file_hash, validate_paths  # noqa: E402


MASTER_SOURCE = ROOT / "private" / "master_sources" / "template_client.xlsx"
EXAMPLE_SOURCE = ROOT / "data_import" / "CONTOH_TERISI_MASTER_DA.xlsx"


# --- Numeric parser and BOM conversion regressions ---
def test_num_and_is_num_finite_values_regression():
    assert imt.num(0.3) == 0.3
    assert imt.num(0.24) == 0.24
    assert imt.num(200.0) == 200.0
    assert imt.num(129000.0) == 129000.0
    assert imt.num("0,24") == 0.24
    assert imt.num("95.000") == 95000.0
    assert imt.is_num(0.3) is True
    assert imt.is_num("1 karung") is False


def test_num_handles_nan_infinity_and_bool_safely():
    assert imt.num(float("nan")) == 0.0
    assert imt.num(float("inf")) == 0.0
    assert imt.num(float("-inf")) == 0.0
    assert imt.num(True) == 0.0
    assert imt.is_num(float("nan")) is False
    assert imt.is_num(float("inf")) is False
    assert imt.is_num(True) is False


def test_convert_line_cm_to_kg_requires_gsm_and_width():
    class API:
        @staticmethod
        def num(v):
            return imt.num(v)

        @staticmethod
        def is_num(v):
            return imt.is_num(v)

        @staticmethod
        def s(v):
            return imt.s(v)

        errors = []
        warnings = []

        @classmethod
        def err(cls, sheet, row, msg):
            cls.errors.append((sheet, row, msg))

        @classmethod
        def warn(cls, sheet, row, msg):
            cls.warnings.append((sheet, row, msg))

    row = {"__row": 22, "qty_per_pcs": 100, "satuan": "cm", "keterangan": "uji"}
    material = {"id": "m1", "code": "FAB-1", "name": "Fabric", "base_uom": "kg", "type": "fabric", "gsm": None, "width_cm": None}
    out = convert_line(row, material, API)
    assert out is None
    assert any("memerlukan gramasi_gsm dan lebar_cm positif" in e[2] for e in API.errors)


# --- Catalog normalization and duplicate conflict behavior ---
def _make_catalog_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "14_KATALOG_JUAL"
    ws.append(["kode_akun", "sku", "harga_jual", "harga_coret", "tautan_produk", "aktif", "nama_tampil"])
    ws.append([" shp-01 ", " sku-a ", 100000, 120000, "https://a", "ya", "Nama A"])
    ws.append(["SHP-01", "SKU-A", 100000, 120000, "https://a", "ya", "Nama A"])
    ws.append(["SHP-01", "SKU-A", 101000, 120000, "https://b", "ya", "Nama B"])
    ws.append(["", "SKU-Z", 50000, 70000, "", "ya", "x"])  # incomplete
    return wb


def test_catalog_analysis_normalizes_case_whitespace_and_detects_conflicts():
    wb = _make_catalog_workbook()
    try:
        data = catalog_analysis(wb)
    finally:
        wb.close()
    assert data["rows"] == 4
    assert data["unique_account_sku"] == 1
    assert len(data["duplicate_groups"]) == 1
    assert data["conflicting_groups"] == 1
    assert len(data["incomplete_rows"]) == 1


def test_build_questions_includes_six_groups_and_actual_rows():
    wb = Workbook()
    ws_bom = wb.active
    ws_bom.title = "10_BOM"
    ws_bom.append(["kode_model", "qty_per_pcs", "satuan"])
    ws_bom.append(["DA-2101", 465, "kg"])
    ws_cat = wb.create_sheet("14_KATALOG_JUAL")
    ws_cat.append(["kode_akun", "sku", "harga_jual", "harga_coret", "tautan_produk", "aktif", "nama_tampil"])
    ws_cat.append(["SHP-01", "SKU-1", 100, 200, "https://x", "ya", "N"])
    ws_cat.append(["SHP-01", "SKU-1", 90, 200, "https://y", "ya", "M"])

    review = {
        "kesalahan": [["10_BOM", 2, "qty ekstrem 465 kg/p"], ["14_KATALOG_JUAL", 2, "duplikat"]],
        "peringatan": [],
    }
    catalog = catalog_analysis(wb)
    qs = build_questions(wb, review, catalog)
    wb.close()

    assert len(qs) == 6
    assert qs[1]["id"] == "Q2"
    assert "465" in qs[1]["pertanyaan"]
    row_refs = [f["baris"] for f in qs[5]["temuan"]]
    assert 2 in row_refs


# --- Safety wrappers and path guards ---
def test_read_only_wrappers_allow_reads_and_block_mutations():
    class DummyCollection:
        def find(self):
            return "ok-find"

        def find_one(self):
            return "ok-find_one"

        def insert_one(self):
            return "blocked"

    class DummyDB(dict):
        pass

    db = DummyDB(test=DummyCollection())
    ro_db = ReadOnlyDatabase(db)
    ro_coll = ro_db["test"]
    assert isinstance(ro_coll, ReadOnlyCollection)
    assert ro_coll.find() == "ok-find"
    with pytest.raises(RuntimeError):
        ro_coll.insert_one()


def test_validate_paths_rejects_public_uploads_existing_and_nonexistent(tmp_path):
    with pytest.raises(ValueError):
        validate_paths(ROOT / "does_not_exist.xlsx", tmp_path / "out")

    with pytest.raises(ValueError):
        validate_paths(EXAMPLE_SOURCE, ROOT / "frontend" / "public" / "x")

    with pytest.raises(ValueError):
        validate_paths(EXAMPLE_SOURCE, ROOT / "uploads" / "x")

    existing_dest = tmp_path / "already"
    existing_dest.mkdir()
    with pytest.raises(FileExistsError):
        validate_paths(EXAMPLE_SOURCE, existing_dest)


def test_package_rejects_negative_historical_errors(tmp_path):
    with pytest.raises(ValueError):
        asyncio.run(bmp.package(EXAMPLE_SOURCE, tmp_path / "pkg-neg", historical_errors=-1))


def test_failure_during_build_cleans_staging_and_no_completed_output(monkeypatch, tmp_path):
    async def boom(*args, **kwargs):
        raise RuntimeError("forced test failure")

    monkeypatch.setattr(bmp, "_build_package", boom)
    dest = tmp_path / "pkg-fail"
    with pytest.raises(RuntimeError):
        asyncio.run(bmp.package(EXAMPLE_SOURCE, dest))
    assert not dest.exists()
    leftovers = [p for p in tmp_path.iterdir() if p.name.startswith(f".{dest.name}-")]
    assert leftovers == []


def test_cli_rejects_apply_flag():
    proc = subprocess.run(
        [sys.executable, str(SCRIPTS_DIR / "build_master_client_package.py"), str(EXAMPLE_SOURCE), "/tmp/ignored", "--apply"],
        capture_output=True,
        text=True,
        check=False,
    )
    assert proc.returncode != 0
    assert "unrecognized arguments: --apply" in proc.stderr


# --- Package generation and artifact integrity ---
def test_real_package_generation_private_artifacts_and_manifest(tmp_path):
    if not MASTER_SOURCE.exists():
        pytest.skip("master source file not present")

    dest = tmp_path / "review-master"
    summary = asyncio.run(bmp.package(MASTER_SOURCE, dest, historical_errors=1990))

    expected_files = {
        "MASTER_CLIENT_AUTOFIX.xlsx",
        "REVIEW_SEBELUM_AUTOFIX.xlsx",
        "REVIEW_SETELAH_AUTOFIX.xlsx",
        "MASTER_CLIENT_AUTOFIX_ULANG.xlsx",
        "HASIL.json",
        "KESALAHAN.json",
        "PERTANYAAN_KLIEN.json",
        "KATALOG_DUPLIKAT.json",
        "PERTANYAAN_KLIEN.md",
        "JAWABAN_KLIEN.xlsx",
        "README.md",
        "MANIFEST.json",
    }
    produced = {p.name for p in dest.iterdir() if p.is_file()}
    assert produced == expected_files
    assert summary["catalog_rows"] == 692
    assert summary["catalog_unique_account_sku"] == 676
    assert summary["catalog_duplicate_groups"] == 16
    assert summary["catalog_conflicting_groups"] == 12
    assert summary["apply_performed"] is False
    assert summary["database_access"] == "read_only"
    assert summary['database_unchanged'] is True
    assert len(summary['database_review_snapshot']) == 17
    assert summary['changes_second_run'] == 0
    assert summary['data_idempotent'] is True
    assert summary['data_sha256'] == summary['data_sha256_second_run']
    assert summary['source_unchanged'] is True
    assert 'DITAHAN' in summary['phase_7']
    questions = json.loads((dest / 'PERTANYAAN_KLIEN.json').read_text())
    assert sum(len(q['temuan']) for q in questions) == summary['errors_after'] + summary['warnings_after']
    wb = load_workbook(dest / 'JAWABAN_KLIEN.xlsx', read_only=True)
    try:
        assert wb['PERTANYAAN'].max_row == 7
        assert wb['TEMUAN'].max_row == 1 + summary['errors_after'] + summary['warnings_after']
        assert wb['KATALOG_DUPLIKAT'].max_row == 33
    finally:
        wb.close()
    wb = load_workbook(dest / 'MASTER_CLIENT_AUTOFIX.xlsx', read_only=True)
    try:
        headers = [cell.value for cell in wb['10_BOM'][1]]
        quantity_column = headers.index('qty_per_pcs') + 1
        assert all(wb['10_BOM'].cell(row, quantity_column).value == 465 for row in range(92, 101))
    finally:
        wb.close()

    # Source integrity + destination permissions.
    assert file_hash(MASTER_SOURCE) == summary["source_sha256"]
    assert (dest.stat().st_mode & 0o777) == 0o700
    for p in dest.iterdir():
        if p.is_file():
            assert (p.stat().st_mode & 0o777) == 0o600

    # Manifest excludes itself and matches all other files.
    manifest = json.loads((dest / "MANIFEST.json").read_text(encoding="utf-8"))
    assert "MANIFEST.json" not in manifest["files"]
    expected_manifest_files = produced - {"MANIFEST.json"}
    assert set(manifest["files"].keys()) == expected_manifest_files
    for name in expected_manifest_files:
        blob = (dest / name).read_bytes()
        assert manifest["files"][name]["sha256"] == hashlib.sha256(blob).hexdigest()
        assert manifest["files"][name]["bytes"] == len(blob)


def test_example_package_is_clean_and_no_false_465_warning(tmp_path):
    dest = tmp_path / "review-example"
    summary = asyncio.run(bmp.package(EXAMPLE_SOURCE, dest))
    assert summary["errors_after"] == 0
    assert summary["warnings_after"] == 0
    assert summary["historical_errors"] is None

    questions = json.loads((dest / "PERTANYAAN_KLIEN.json").read_text(encoding="utf-8"))
    assert len(questions) == 6
    q2 = [q for q in questions if q["id"] == "Q2"][0]
    assert "465" not in q2["pertanyaan"]


def test_no_recognized_master_sheet_fails_safely(tmp_path):
    src = tmp_path / "unknown.xlsx"
    wb = Workbook()
    wb.active.title = "UNKNOWN"
    wb.active.append(["x", "y"])
    wb.active.append([1, 2])
    wb.save(src)
    wb.close()

    with pytest.raises(ValueError, match="Tidak ditemukan sheet master yang dikenali"):
        asyncio.run(bmp.package(src, tmp_path / "out-unknown"))


# --- Generated docs/workbooks safety: no formula injection, NIK redaction ---
def test_generated_documents_redact_nik_and_store_formula_text_as_string(tmp_path):
    summary = {
        "source": "x.xlsx",
        "source_sha256": "abc",
        "errors_after": 1,
        "warnings_after": 0,
    }
    questions = [
        {
            "id": "Q1",
            "judul": "Topik",
            "pertanyaan": "=1+1",
            "status": "PERLU JAWABAN",
            "temuan": [{"jenis": "kesalahan", "sheet": "02_KARYAWAN", "baris": 2, "pesan": "NIK 1234567890123456"}],
        }
    ] + [
        {"id": f"Q{i}", "judul": "x", "pertanyaan": "x", "status": "TINJAU BISNIS", "temuan": []}
        for i in range(2, 7)
    ]
    catalog = {"duplicate_groups": []}

    write_questions(tmp_path, summary, questions)
    write_answers(tmp_path, questions, catalog)

    text = (tmp_path / "PERTANYAAN_KLIEN.md").read_text(encoding="utf-8")
    assert "[NIK KTP disamarkan]" in text

    wb = load_workbook(tmp_path / "JAWABAN_KLIEN.xlsx", data_only=False)
    try:
        ws = wb["PERTANYAAN"]
        assert ws["C2"].value == "=1+1"
        assert ws["C2"].data_type == "s"
    finally:
        wb.close()


# --- External API/URL sanity checks in scope ---
def test_external_api_health_200():
    base_url = os.environ.get("REACT_APP_BACKEND_URL") or (ROOT / "frontend" / ".env").read_text().split("=")[1].splitlines()[0]
    response = requests.get(f"{base_url.rstrip('/')}/api/health", timeout=20)
    assert response.status_code == 200


def test_importer_dry_run_does_not_mutate_real_db_counts():
    from dotenv import load_dotenv
    from motor.motor_asyncio import AsyncIOMotorClient

    load_dotenv(BACKEND_DIR / ".env")
    client = AsyncIOMotorClient(os.environ["MONGO_URL"], serverSelectionTimeoutMS=10000)
    db = client[os.environ["DB_NAME"]]
    collections = [
        "rahaza_locations",
        "rahaza_employees",
        "rahaza_payroll_profiles",
        "rahaza_colors",
        "rahaza_sizes",
        "rahaza_processes",
        "rahaza_materials",
        "rahaza_models",
        "rahaza_boms",
        "vendor_partners",
        "dewi_maklon_clients",
        "marketing_platform_accounts",
        "marketing_catalog_items",
        "marketing_kol_creators",
        "marketing_livehosts",
    ]
    async def run_checks():
        async def snapshot():
            result = {}
            for name in collections:
                documents = await db[name].find({}, {'_id': 0}).to_list(None)
                rows = sorted(json.dumps(doc, sort_keys=True, default=str) for doc in documents)
                result[name] = hashlib.sha256(json.dumps(rows).encode()).hexdigest()
            return result
        before = await snapshot()
        code = await imt.main(EXAMPLE_SOURCE, False, set())
        after = await snapshot()
        return code, before, after

    try:
        code, before, after = asyncio.run(run_checks())
    finally:
        client.close()
    assert code == 0
    assert after == before
