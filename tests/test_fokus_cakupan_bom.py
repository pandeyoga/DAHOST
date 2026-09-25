"""Berkas FOKUS harus MENCAKUP SEMUA model yang papan Kelengkapan Data R&D tandai BOM-nya kurang
(temuan owner 2026-09-22: berkas lama melewatkan model tanpa BOM / varian tanpa BOM / model dihentikan).

Jalankan:  cd /app/backend && set -a && . .env && set +a && python ../tests/test_fokus_cakupan_bom.py
Mengubah data (apply Ona + Heidi) → DB dipulihkan ke seed go-live di akhir.
"""
import io
import os
import subprocess
import sys

import openpyxl
import requests

API = os.environ.get("API_URL") or "http://localhost:8001"
FAILS = []


def check(name, cond, info=""):
    print(("PASS" if cond else "FAIL"), name, "" if cond else info)
    if not cond:
        FAILS.append(name)


def login(email, pw):
    r = requests.post(f"{API}/api/auth/login", json={"email": email, "password": pw}, timeout=30)
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['token']}"}


def fokus(h):
    r = requests.post(f"{API}/api/rahaza/master/gap-fokus", headers=h, timeout=120)
    assert r.status_code == 200, r.text[:200]
    return openpyxl.load_workbook(io.BytesIO(r.content))


def rows(ws):
    return [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]


H = login("admin@garment.com", "Admin@123")
board = requests.get(f"{API}/api/dewi/rnd/completeness", headers=H, timeout=60).json()
flagged = {r["code"] for r in board["rows"] if "bom" in r["missing"] or "accessories" in r["missing"]}
discontinued = {r["code"] for r in board["rows"] if r.get("discontinued")}
check("papan: model dihentikan tidak dihitung kurang (missing=[], score=None)",
      all(not r["missing"] and r["score"] is None for r in board["rows"] if r.get("discontinued")))
check("papan: discontinued_models = 6 (GIA, Luvia, Maudy, Erlyna, Airyn, Jeslyn)", board.get("discontinued_models") == 6 and len(discontinued) == 6, str(discontinued))

wb = fokus(H)
check("sheet RINGKASAN_MODEL ada setelah PETUNJUK", wb.sheetnames[:2] == ["PETUNJUK", "RINGKASAN_MODEL"], str(wb.sheetnames))
ring = rows(wb["RINGKASAN_MODEL"])
ring_codes = {r[0] for r in ring}
check("RINGKASAN = model kurang BOM di papan + model dihentikan", ring_codes == flagged | discontinued,
      f"hanya di papan: {flagged - ring_codes}; hanya di ringkasan: {ring_codes - flagged - discontinued}")
bom_codes = {r[0] for r in rows(wb["BOM_AKSESORIS"])}
check("SEMUA model kurang BOM di papan punya kelompok di BOM_AKSESORIS", flagged <= bom_codes, str(flagged - bom_codes))
check("model dihentikan TIDAK di BOM_AKSESORIS", not (discontinued & bom_codes))
by_code = {r[0]: r for r in ring}
check("DA-2201 Ona: 'belum punya BOM sama sekali (7 varian)'", "belum punya BOM sama sekali (7" in (by_code.get("DA-2201") or [""] * 8)[7])
check("DA-1509 Hanny: '2 dari 13 varian belum punya BOM'", "2 dari 13" in (by_code.get("DA-1509") or [""] * 8)[7])
check("DA-2112 Heidi: 'belum punya varian/SKU'", "belum punya varian" in (by_code.get("DA-2112") or [""] * 8)[7])
check("DA-2104 GIA: dihentikan (abu-abu)", "nonaktif" in (by_code.get("DA-2104") or [""] * 8)[7])
hanny = [r for r in rows(wb["BOM_AKSESORIS"]) if r[0] == "DA-1509"]
check("Hanny: 1 kelompok untuk semua varian (tak satu pun BOM ber-aksesoris; 2 varian tanpa BOM dibuat saat unggah)", len(hanny) == 1 and not hanny[0][7] and "belum punya BOM dibuat otomatis" in (hanny[0][9] or ""), str(hanny))
lyora = [r for r in rows(wb["BOM_AKSESORIS"]) if r[0] == "DA-1101"]
check("Lyora: 1 dari 6 BOM tanpa aksesoris → 1 kelompok varian BURGUNDY disalin (biru) dari varian lain", len(lyora) == 1 and "BURGUNDY" in (lyora[0][7] or "") and lyora[0][2], str(lyora))
check("papan: Lyora accessories ✗ (dinilai per varian)", "accessories" in next((r["missing"] for r in board["rows"] if r["code"] == "DA-1101"), []))
vb_codes = {r[0] for r in rows(wb["VARIAN_BARU"])}
check("model tanpa SKU ada di VARIAN_BARU DAN BOM_AKSESORIS", {"DA-2112", "DA-3512"} <= vb_codes & bom_codes)

# unggah apa adanya = aman (tidak ada perubahan)
buf = io.BytesIO()
wb.save(buf)
r = requests.post(f"{API}/api/rahaza/master/fill-preview?scope=all", headers=H, files={"file": ("f.xlsx", buf.getvalue())}, timeout=120)
check("fill-preview berkas kosong → 200, 0 kelompok BOM", r.status_code == 200 and len(r.json()["bom_groups"]) == 0, r.text[:200])

# simulasi klien: Ona (nol BOM) isi 1 aksesoris; Heidi (nol SKU) isi VARIAN_BARU + aksesoris → satu unggahan
acc = next(r for r in rows(wb["REF_AKSESORIS"]) if r[2] == "accessory" and (r[4] or 0) > 0 and r[3] == "pcs")
for row in wb["BOM_AKSESORIS"].iter_rows(min_row=2):
    if row[0].value in ("DA-2201", "DA-2112"):
        row[2].value, row[4].value, row[5].value = acc[0], 2, "pcs"
for row in wb["VARIAN_BARU"].iter_rows(min_row=2):
    if row[0].value == "DA-2112":
        row[3].value, row[4].value = "HTM", "ALLSIZE"
buf = io.BytesIO()
wb.save(buf)
r = requests.post(f"{API}/api/rahaza/master/fill-apply?scope=all", headers=H, files={"file": ("f.xlsx", buf.getvalue())}, timeout=300)
a = r.json()
check("fill-apply: Ona 7 BOM dasar + Heidi 1 SKU + 1 BOM", r.status_code == 200 and a.get("bom_base_created") == 8 and a.get("variants_created") == 1, str(a)[:300])
after = requests.get(f"{API}/api/dewi/rnd/completeness", headers=H, timeout=60).json()
ok = {r["code"]: (r["flags"]["bom"], r["flags"]["accessories"]) for r in after["rows"] if r["code"] in ("DA-2201", "DA-2112")}
check("papan setelah unggah: Ona & Heidi BOM ✓ aksesoris ✓", ok == {"DA-2201": (True, True), "DA-2112": (True, True)}, str(ok))
ring2 = {r[0] for r in rows(fokus(H)["RINGKASAN_MODEL"])}
check("RINGKASAN menyusut 47 → 45 (Ona & Heidi hilang)", len(ring2) == len(ring) - 2 and not ({"DA-2201", "DA-2112"} & ring2), str(len(ring2)))

subprocess.run(["bash", "/app/scripts/seed_golive_restore.sh", "--force"], check=True, capture_output=True, timeout=300)
subprocess.run([sys.executable, "/app/scripts/seed_test_accounts.py"], check=True, capture_output=True, timeout=120, cwd="/app/backend")
print("\nSEMUA PASS" if not FAILS else f"\nGAGAL: {FAILS}")
sys.exit(1 if FAILS else 0)
