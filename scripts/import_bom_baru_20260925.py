"""Impor BOM BARU dari sheet `2-BOM` (EXPORT_SKU_BOM_DA.xlsx, 2026-09-25).

Aturan (owner 2026-09-25):
  * Baris `potongan_kain` (CUT-…) di Excel DIABAIKAN. Potongan dibuat/dipakai lewat fitur otomatis
    (`core.cut_panel_master.ensure_panel_for_variant`) → tepat SATU baris potongan per BOM, tidak konflik.
  * Model / varian (SKU) / FG yang belum ada dibuat; warna, ukuran, kategori & aksesoris WAJIB sudah ada di master.
  * BOM per SKU = 1 pcs potongan (otomatis) + aksesoris dari Excel. BOM lama SKU yang sama ditimpa (versi ganda dinonaktifkan).
  * Hangtag A-HTG-0003 + Pin A-PIN-0001 1 pcs ditambahkan bila belum ada (aturan owner: di SEMUA SKU aktif).
  * Idempoten: dijalankan 2x hasilnya sama.

VPS:
  docker compose --env-file .env exec -T backend python /app/scripts/import_bom_baru_20260925.py             # pratinjau
  docker compose --env-file .env exec -T backend python /app/scripts/import_bom_baru_20260925.py --terapkan   # tulis
"""
import asyncio
import os
import re
import sys
import uuid
from collections import Counter, defaultdict
from datetime import datetime, timezone

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "backend"))
os.chdir(os.path.join(HERE, "..", "backend"))
from dotenv import load_dotenv  # noqa: E402
load_dotenv(".env")
from motor.motor_asyncio import AsyncIOMotorClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from core import product_master as pm  # noqa: E402
from core.bom_uom import annotate_line  # noqa: E402
from core.cut_panel_master import bom_line_for_panel, ensure_panel_for_variant, is_panel_line  # noqa: E402
from utils.variant_ssot import ensure_fg_material  # noqa: E402

FILE = os.path.join(HERE, "vps", "data", "bom_baru_20260925.xlsx")
SHEET = "2-BOM"
USER = {"id": "system", "name": "import_bom_baru_20260925"}
ALWAYS = ("A-HTG-0003", "A-PIN-0001")


def now():
    return datetime.now(timezone.utc)


def clean(v):
    return re.sub(r"[\u200b\u00a0]", "", str(v or "")).strip()


def key(v):
    return re.sub(r"\s+", "", clean(v)).upper()


def num(v):
    if isinstance(v, (int, float)):
        return float(v)
    m = re.search(r"\d+(?:[.,]\d+)?", clean(v))
    return float(m.group(0).replace(",", ".")) if m else None


def read_rows():
    ws = load_workbook(FILE, read_only=True, data_only=True)[SHEET]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [clean(h) for h in rows[0]]
    return [dict(zip(hdr, r)) for r in rows[1:] if r and any(r[:20])]


async def main(apply: bool):
    db = AsyncIOMotorClient(os.environ["MONGO_URL"])[os.environ["DB_NAME"]]
    rows = read_rows()
    stat, warn = Counter(), []
    stat["baris_excel"] = len(rows)
    stat["baris_potongan_diabaikan"] = sum(1 for r in rows if clean(r.get("jenis_komponen")) == "potongan_kain")

    colors = {key(c["code"]): c async for c in db.rahaza_colors.find({"active": {"$ne": False}}, {"_id": 0})}
    sizes = {}
    async for s in db.rahaza_sizes.find({"active": {"$ne": False}}, {"_id": 0}):
        sizes[key(s["code"])] = s
        sizes.setdefault(key(s.get("name")), s)
    cats = {key(c["name"]): c async for c in db.rahaza_product_categories.find({}, {"_id": 0})}
    mats = {}
    codes = {key(r.get("kode_material")) for r in rows if clean(r.get("jenis_komponen")) != "potongan_kain"} | set(ALWAYS)
    async for m in db.rahaza_materials.find({"code": {"$in": list(codes)}, "active": {"$ne": False}}, {"_id": 0}):
        mats[key(m["code"])] = m

    # kelompokkan per model → per SKU
    models = defaultdict(lambda: {"name": "", "cat": "", "skus": defaultdict(lambda: {"color": "", "size": "", "lines": []})})
    for r in rows:
        mc = key(r.get("kode_model"))
        md = models[mc]
        md["name"] = md["name"] or clean(r.get("nama_model"))
        md["cat"] = md["cat"] or clean(r.get("kategori"))
        sk = md["skus"][(key(r.get("kode_warna")), key(r.get("ukuran")))]
        sk["color"], sk["size"] = key(r.get("kode_warna")), key(r.get("ukuran"))
        if clean(r.get("jenis_komponen")) != "potongan_kain":
            qty = num(r.get("qty"))
            if not qty:
                warn.append(f"{clean(r.get('sku'))}: qty {key(r.get('kode_material'))} KOSONG di Excel → baris dilewati (isi manual di BOM editor)")
                continue
            sk["lines"].append((key(r.get("kode_material")), qty, clean(r.get("satuan")) or "pcs"))

    touched_models = []
    for mc, md in models.items():
        cat = cats.get(key(md["cat"]))
        if not cat:
            warn.append(f"{mc}: kategori '{md['cat']}' tidak ada di Master Kategori → model dilewati")
            continue
        model = await db.rahaza_models.find_one(pm.live_model_filter({"code": mc}), {"_id": 0, "sop_steps": 0})
        if not model:
            stat["model_baru"] += 1
            model = {"id": str(uuid.uuid4()), "code": mc, "name": md["name"], "description": "", "bundle_size": 30,
                     "material_kg_per_pcs": 0.0, "yarn_kg_per_pcs": 0.0, "sop_steps": [], "reference_videos": [],
                     "reference_images": [], "active": True, "base_hpp": 0.0, "retail_price": 0.0, "weight_gram": 0.0,
                     "hpp_rnd": 0.0, "hpp": 0.0, "hpp_source": "none", "created_from": USER["name"],
                     "created_at": now(), "updated_at": now()}
            pm.apply_category(model, cat)
            if apply:
                await db.rahaza_models.insert_one(dict(model))
        else:
            stat["model_sudah_ada"] += 1
        touched_models.append(model["id"])

        for (cc, sc), sk in md["skus"].items():
            color, size = colors.get(cc), sizes.get(sc)
            if not color or not size:
                warn.append(f"{mc}-{cc}-{sc}: warna/ukuran tidak ada di master → dilewati")
                continue
            sku = f"{mc}-{color['code'].upper()}-{size['code'].upper()}"
            variant = await db.rahaza_model_variants.find_one(
                {"model_id": model["id"], "size_id": size["id"], "color_id": color["id"], "active": True}, {"_id": 0})
            if not variant:
                stat["varian_baru"] += 1
                variant = {"id": str(uuid.uuid4()), "model_id": model["id"], "model_code": mc, "model_name": model["name"],
                           "size_id": size["id"], "size_code": size["code"], "color_id": color["id"],
                           "color_code": color["code"], "color_name": color["name"], "color_hex": color.get("hex"),
                           "sku": sku, "barcode": "", "notes": "", "active": True, "created_from": USER["name"],
                           "created_at": now(), "updated_at": now()}
                if apply:
                    await db.rahaza_model_variants.insert_one(dict(variant))
            else:
                stat["varian_sudah_ada"] += 1
            if apply:
                await ensure_fg_material(db, variant, user=USER)
                panel, new_panel = await ensure_panel_for_variant(db, variant, model)
                stat["potongan_baru"] += 1 if new_panel else 0
            else:
                panel = {"id": "(dry-run)", "code": f"CUT-{mc}-{color['code']}-{size['code']}", "name": "Potongan", "unit_cost": 0}

            # baris aksesoris (gabung kode kembar dgn satuan sama)
            merged, order = {}, []
            for code, qty, unit in sk["lines"] + [(c, 1.0, "pcs") for c in ALWAYS]:
                if (code, unit) in merged:
                    if code in ALWAYS and qty == 1.0 and unit == "pcs":
                        continue
                    merged[(code, unit)] += qty
                    warn.append(f"{sku}: {code} tercantum >1x → qty dijumlah")
                    continue
                merged[(code, unit)] = qty
                order.append((code, unit))
            lines = []
            for code, unit in order:
                m = mats.get(code)
                if not m:
                    warn.append(f"{sku}: aksesoris {code} tidak ada di master → dilewati")
                    continue
                lines.append(annotate_line({"material_id": m["id"], "code": m["code"], "name": m["name"],
                                            "material_type": m.get("type") or "accessory", "qty": merged[(code, unit)],
                                            "unit": unit, "notes": ""}, m))
            new_materials = [annotate_line({**bom_line_for_panel(panel), "material_type": "fabric", "is_cut_panel": True}, panel)] + lines

            bkey = {"model_id": model["id"], "size_id": size["id"], "color_code": color["code"].upper()}
            boms = await db.rahaza_boms.find(bkey, {"_id": 0}).to_list(50)
            live = [b for b in boms if b.get("active") is not False] or boms
            existing = max(live, key=lambda b: b.get("version") or 0) if live else None
            if existing:
                old = [(ln.get("code"), ln.get("qty"), ln.get("unit")) for ln in existing.get("materials") or []]
                same = old == [(ln.get("code"), ln.get("qty"), ln.get("unit")) for ln in new_materials]
                dropped = [ln.get("code") for ln in existing.get("materials") or [] if is_panel_line(ln) and ln.get("code") != panel["code"]]
                if dropped:
                    warn.append(f"{sku}: baris potongan lain dihapus {dropped} (konflik)")
                stat["bom_sama" if same else "bom_ditimpa"] += 1
                if apply and (not same or existing.get("active") is False):
                    await db.rahaza_boms.update_one({"id": existing["id"]}, {
                        "$set": {"materials": new_materials, "active": True, "is_active": True, "updated_at": now(),
                                 "accessories_source": f"{USER['name']} ({SHEET})"},
                        "$push": {"change_log": {"at": now(), "by": USER["name"], "note": "impor BOM baru 2-BOM",
                                                 "lines_before": len(old), "lines_after": len(new_materials)}}})
                for b in boms:
                    if b["id"] != existing["id"] and b.get("active") is not False:
                        stat["bom_versi_ganda_nonaktif"] += 1
                        if apply:
                            await db.rahaza_boms.update_one({"id": b["id"]}, {"$set": {
                                "active": False, "is_active": False, "deactivated_reason": "impor BOM baru: versi ganda",
                                "updated_at": now()}})
            else:
                stat["bom_baru"] += 1
                if apply:
                    await db.rahaza_boms.insert_one({
                        "id": str(uuid.uuid4()), **bkey, "color": color["name"], "version": 1, "is_active": True,
                        "active": True, "materials": new_materials, "notes": "", "created_by": USER["id"],
                        "created_at": now(), "updated_at": now(), "accessories_source": f"{USER['name']} ({SHEET})",
                        "change_log": []})

    if apply:
        from core.product_costing import apply_model_cost
        for mid in touched_models:
            try:
                await apply_model_cost(db, mid, USER)
                stat["hpp_dihitung_ulang"] += 1
            except Exception as e:  # noqa: BLE001
                warn.append(f"HPP model {mid} gagal dihitung: {e}")

    print(("PRATINJAU — belum ada yang ditulis. Jalankan dengan --terapkan untuk menulis.\n" if not apply else "✓ DITERAPKAN\n")
          + "\n".join(f"  {k:<28} {v}" for k, v in stat.items()))
    if warn:
        print(f"\nCatatan ({len(warn)}):")
        for w in dict.fromkeys(warn):
            print("  - " + w)


if __name__ == "__main__":
    asyncio.run(main("--terapkan" in sys.argv))
