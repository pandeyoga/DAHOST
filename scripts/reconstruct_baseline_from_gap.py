#!/usr/bin/env python3
"""scripts/reconstruct_baseline_from_gap.py — BASELINE PREVIEW dari berkas DATA_YANG_PERLU_DIISI_DA.xlsx.

Dipakai HANYA bila workbook MASTER_DATA_DA_PERBAIKAN*.xlsx dan backup VPS tidak tersedia di workspace.
Berkas gap dibuat sistem dari DB VPS, jadi memuat master yang cukup untuk menguji importir BOM:
  REF_LOKASI → 01_LOKASI (peran sesuai keputusan owner di prepare_golive_workbook.PERAN)
  STOK_AWAL_FG → 08_MODEL + 03_WARNA + 04_UKURAN + 09_BARANG_JADI (645 SKU)
  STOK_AWAL_MATERIAL / REF_AKSESORIS → 06_MATERIAL_KAIN + 07_AKSESORIS (harga per satuan dasar)
  TOKO → 13_AKUN_TOKO ; BOM_AKSESORIS.nama_model → nama 20 model tanpa SKU
Yang TIDAK bisa direkonstruksi: BOM kain, isi kemasan (pack_size), berat, techpack, SOP, katalog per toko.

    python3 scripts/reconstruct_baseline_from_gap.py /app/private/golive/DATA_YANG_PERLU_DIISI_DA.xlsx [--apply]
"""
from __future__ import annotations

import asyncio
import os
import subprocess
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))
sys.path.insert(0, str(ROOT / "scripts"))
from master_template_spec import SHEETS  # noqa: E402
from prepare_golive_workbook import PERAN  # noqa: E402

OUT = ROOT / "private" / "golive" / "MASTER_REKONSTRUKSI_DARI_GAP.xlsx"


def rows(wb, name):
    if name not in wb.sheetnames:
        return []
    return [tuple(r) + (None,) * (10 - len(r)) for r in wb[name].iter_rows(values_only=True, min_row=2) if r and r[0] not in (None, "")]


def build(src: Path) -> Path:
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    fg = rows(wb, "STOK_AWAL_FG")
    mats = rows(wb, "STOK_AWAL_MATERIAL")
    ref = {r[0]: r for r in rows(wb, "REF_AKSESORIS")}
    toko = rows(wb, "TOKO")
    bom = rows(wb, "BOM_AKSESORIS")
    models: dict[str, str] = {}
    colors: dict[str, str] = {}
    sizes: dict[str, str] = {}
    skus = []
    for sku, name, mcode, warna, ukuran, *_ in fg:
        rest = sku[len(mcode) + 1:]
        color_code = rest[: -(len(ukuran) + 1)] if rest.endswith("-" + ukuran) else rest.split("-")[0]
        colors.setdefault(color_code.upper(), (warna or color_code).title() if warna else color_code)
        sizes.setdefault(ukuran.upper(), ukuran)
        models.setdefault(mcode, name.split(" [")[0] if name else mcode)
        skus.append((sku, name, mcode, color_code.upper(), ukuran.upper()))
    for r in bom:
        if r[0] and r[1]:
            models.setdefault(str(r[0]).strip(), str(r[1]).strip())
    out = openpyxl.Workbook()
    out.remove(out.active)

    def sheet(name, recs):
        ws = out.create_sheet(name)
        cols = [c[0] for c in SHEETS[name]["kolom"]]
        ws.append(cols)
        for rec in recs:
            ws.append([rec.get(c) for c in cols])

    lokasi = rows(wb, "REF_LOKASI")
    sheet("01_LOKASI", [{"kode": r[0], "nama": r[1], "tipe": r[2] or "gudang", "aktif": "ya", "peran": PERAN.get(r[0], "")}
                        for r in lokasi])
    sheet("03_WARNA", [{"kode": k, "nama": v, "urutan": i + 1} for i, (k, v) in enumerate(sorted(colors.items()))])
    order = {"S": 1, "M": 2, "L": 3, "XL": 4, "XXL": 5, "2XL": 5, "3XL": 6, "ALLSIZE": 9}
    sheet("04_UKURAN", [{"kode": k, "nama": v, "urutan": order.get(k, 7)} for k, v in sorted(sizes.items(), key=lambda kv: order.get(kv[0], 7))])
    sheet("06_MATERIAL_KAIN", [{"kode": r[0], "nama": r[1], "jenis": "fabric", "satuan_dasar": r[3], "harga_per_satuan": r[6] or 0}
                               for r in mats if r[2] == "fabric"])
    sheet("07_AKSESORIS", [{"kode": r[0], "nama": r[1], "satuan_dasar": r[3], "kategori": "",
                            "harga_per_satuan": (ref.get(r[0]) or (None, None, None, None, r[6]))[4] or r[6] or 0}
                           for r in mats if r[2] != "fabric"])
    sheet("08_MODEL", [{"kode": k, "nama": v, "kategori": "", "harga_jual_dasar": 0} for k, v in sorted(models.items())])
    sheet("09_BARANG_JADI", [{"sku": s, "nama": n, "kode_model": m, "kode_warna": c, "kode_ukuran": u, "satuan": "pcs", "harga_jual": 0}
                             for s, n, m, c, u in skus])
    sheet("13_AKUN_TOKO", [{"kode_akun": r[0], "nama_akun": r[1], "platform": (r[2] or "shopee").lower(), "status": "active"} for r in toko])
    OUT.parent.mkdir(parents=True, exist_ok=True)
    out.save(OUT)
    print(f"→ {OUT}: {len(models)} model · {len(colors)} warna · {len(sizes)} ukuran · {len(skus)} SKU · {len(mats)} material · {len(toko)} toko")
    return OUT


async def sync():
    from dotenv import load_dotenv
    from motor.motor_asyncio import AsyncIOMotorClient
    load_dotenv(ROOT / "backend" / ".env")
    db = AsyncIOMotorClient(os.environ["MONGO_URL"])[os.environ["DB_NAME"]]
    from core.master_sync import sync_product_masters
    rep = await sync_product_masters(db, {"id": "system", "name": "reconstruct_baseline"})
    print("sync:", rep.get("variants"))
    for c in ("rahaza_models", "rahaza_model_variants", "rahaza_materials", "rahaza_boms", "marketing_platform_accounts"):
        print(f"  {c}: {await db[c].count_documents({})}")


if __name__ == "__main__":
    src = Path(sys.argv[1])
    dest = build(src)
    args = [sys.executable, str(ROOT / "scripts" / "import_master_template.py"), str(dest)]
    if "--apply" in sys.argv:
        args.append("--apply")
    rc = subprocess.call(args)
    if rc == 0 and "--apply" in sys.argv:
        asyncio.run(sync())
    sys.exit(rc)
