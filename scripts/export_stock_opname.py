"""Ekspor daftar MASTER untuk STOCK OPNAME klien — 1 sheet per kelompok:
AKSESORIS · KAIN_ROLL · POTONGAN_CUTTING · FG (barang jadi). Kolom kuning = diisi klien.

VPS:
  docker compose --env-file .env exec -T backend python /app/scripts/export_stock_opname.py
  docker compose --env-file .env cp backend:/app/backups/STOCK_OPNAME_<tanggal>.xlsx ./
"""
import asyncio
import os
import sys
from datetime import datetime, timezone

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "backend"))
os.chdir(os.path.join(HERE, "..", "backend"))
from dotenv import load_dotenv  # noqa: E402
load_dotenv(".env")
from motor.motor_asyncio import AsyncIOMotorClient  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from core.stock_service import onhand_map  # noqa: E402

OUT_DIR = sys.argv[sys.argv.index("--out") + 1] if "--out" in sys.argv else "/app/backups"
HEAD = PatternFill("solid", fgColor="1F3A5F")
INPUT = PatternFill("solid", fgColor="FFF2CC")
THIN = Border(*(Side(style="thin", color="BFBFBF"),) * 4)

GROUPS = [
    ("AKSESORIS", {"type": "accessory"}, False),
    ("KAIN_ROLL", {"type": "fabric", "is_cut_panel": {"$ne": True}}, True),
    ("POTONGAN_CUTTING", {"is_cut_panel": True}, False),
    ("FG", {"type": "fg"}, False),
]
BASE_COLS = ["No", "Kode", "Nama", "Kategori", "Warna", "Ukuran", "Model", "Satuan Dasar", "Stok Sistem"]
INPUT_COLS = ["Lokasi Hitung", "QTY FISIK"]


async def main():
    db = AsyncIOMotorClient(os.environ["MONGO_URL"])[os.environ["DB_NAME"]]
    locs = [f"{loc['code']} - {loc['name']}" async for loc in db.rahaza_locations.find(
        {"active": {"$ne": False}, "type": {"$ne": "kantor"}}, {"_id": 0, "code": 1, "name": 1}).sort("code", 1)]
    wb = Workbook()
    guide = wb.active
    guide.title = "PETUNJUK"
    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    lines = [
        f"STOCK OPNAME — daftar master per {today}",
        "",
        "1. Isi HANYA kolom KUNING: 'Lokasi Hitung' (pilih dari daftar) dan 'QTY FISIK' (hasil hitung, satuan = kolom 'Satuan Dasar').",
        "2. Barang yang sama ada di >1 lokasi: salin barisnya (Kode sama), isi lokasi & qty masing-masing.",
        "3. Kain roll: isi QTY FISIK dalam satuan dasar (mis. kg/yard) + 'Jumlah Roll'.",
        "4. Barang tidak ada/0: isi QTY FISIK = 0 (jangan dikosongkan). Baris kosong = belum dihitung.",
        "5. Jangan ubah kolom Kode & material_id (dipakai untuk impor balik).",
        "",
        "Sheet: AKSESORIS · KAIN_ROLL · POTONGAN_CUTTING · FG (barang jadi)",
    ]
    for i, t in enumerate(lines, 1):
        guide.cell(i, 1, t).font = Font(bold=(i == 1), size=13 if i == 1 else 11)
    guide.column_dimensions["A"].width = 120

    lists = wb.create_sheet("_LOKASI")
    for i, loc in enumerate(locs, 1):
        lists.cell(i, 1, loc)
    lists.sheet_state = "hidden"

    summary = []
    for title, q, is_roll in GROUPS:
        rows = await db.rahaza_materials.find({**q, "active": {"$ne": False}}, {"_id": 0}).sort("code", 1).to_list(20000)
        stock = await onhand_map([r["id"] for r in rows], db=db) if rows else {}
        ws = wb.create_sheet(title)
        cols = BASE_COLS + INPUT_COLS + (["Jumlah Roll"] if is_roll else []) + ["Catatan", "material_id"]
        inputs = {c for c in cols if c in INPUT_COLS or c in ("Jumlah Roll", "Catatan")}
        for j, c in enumerate(cols, 1):
            cell = ws.cell(1, j, c)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = HEAD
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = THIN
        for i, r in enumerate(rows, 2):
            vals = [i - 1, r.get("code"), r.get("name"), r.get("category_name") or r.get("category") or "",
                    r.get("color") or r.get("color_name") or "", r.get("size_code") or r.get("size") or "",
                    r.get("model_code") or r.get("style_sku") or "", r.get("base_uom") or r.get("unit") or "",
                    float(stock.get(r["id"], 0) or 0), None, None] + ([None] if is_roll else []) + [None, r["id"]]
            for j, v in enumerate(vals, 1):
                cell = ws.cell(i, j, v)
                cell.border = THIN
                if cols[j - 1] in inputs:
                    cell.fill = INPUT
        last = max(len(rows) + 1, 2)
        li = cols.index("Lokasi Hitung") + 1
        qi = cols.index("QTY FISIK") + 1
        from openpyxl.utils import get_column_letter as L
        if locs:
            dv = DataValidation(type="list", formula1=f"=_LOKASI!$A$1:$A${len(locs)}", allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{L(li)}2:{L(li)}{last + 500}")
        dq = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                            error="QTY harus angka ≥ 0", showErrorMessage=True)
        ws.add_data_validation(dq)
        dq.add(f"{L(qi)}2:{L(qi)}{last + 500}")
        for j, w in enumerate([5, 26, 48, 16, 16, 9, 12, 10, 12, 26, 12] + ([11] if is_roll else []) + [24, 38], 1):
            ws.column_dimensions[L(j)].width = w
        ws.column_dimensions[L(len(cols))].hidden = True
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = f"A1:{L(len(cols))}{last}"
        summary.append(f"{title}: {len(rows)} item")

    os.makedirs(OUT_DIR, exist_ok=True)
    path = os.path.join(OUT_DIR, f"STOCK_OPNAME_{today}.xlsx")
    wb.save(path)
    print("✓ " + path + "\n  " + "\n  ".join(summary))


if __name__ == "__main__":
    asyncio.run(main())
