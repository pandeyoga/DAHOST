"""Tarik SEMUA SKU + BOM-nya (berisi maupun kosong) langsung dari DB ke Excel.

Jalankan:  cd /app/backend && set -a && . .env && set +a && python ../scripts/export_sku_bom.py [--out PATH]
Sheet:
  DETAIL_BOM     satu baris per komponen BOM; SKU tanpa BOM = 1 baris kolom komponen kosong;
                 model tanpa SKU = 1 baris kolom SKU kosong (status TANPA_SKU)
  RINGKASAN_SKU  satu baris per SKU: status BOM, jumlah baris kain/aksesoris, biaya BOM
  RINGKASAN_MODEL satu baris per model: kelengkapan per varian (logika = scripts/laporan_kekosongan_bom.py)
  KETERANGAN     arti kolom & status
"""
from __future__ import annotations

import os
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pymongo import MongoClient

FILL_HDR = PatternFill("solid", fgColor="1F3A5F")
FILL_NO_BOM = PatternFill("solid", fgColor="F8CBAD")     # oranye: SKU tanpa BOM
FILL_NO_ACC = PatternFill("solid", fgColor="FFF2CC")     # kuning: BOM ada tapi tanpa aksesoris
FILL_NO_SKU = PatternFill("solid", fgColor="D9D9D9")     # abu: model tanpa SKU
FILL_INACTIVE = PatternFill("solid", fgColor="EDEDED")

DETAIL_COLS = [
    "kode_model", "nama_model", "kategori", "model_aktif", "hpp_model_db",
    "sku", "kode_warna", "nama_warna", "ukuran", "sku_aktif",
    "status_bom", "bom_id", "versi_bom", "bom_aktif", "jumlah_baris_bom",
    "no_baris", "jenis_komponen", "kode_material", "nama_material",
    "qty", "satuan", "qty_dasar", "satuan_dasar", "harga_per_satuan_dasar", "biaya_baris",
    "harga_master_saat_ini", "satuan_master", "material_terhubung", "kain_sumber", "catatan_baris",
]
SUMMARY_COLS = [
    "kode_model", "nama_model", "kategori", "sku", "kode_warna", "nama_warna", "ukuran", "sku_aktif",
    "status_bom", "bom_id", "baris_kain", "baris_aksesoris", "baris_total", "biaya_bom", "hpp_model_db",
]
MODEL_COLS = [
    "kode_model", "nama_model", "kategori", "model_aktif", "status_kelengkapan", "varian_total", "varian_aktif",
    "bom_aktif", "varian_tanpa_bom", "bom_tanpa_aksesoris", "hpp_model_db",
]


def _is_acc(ln: dict) -> bool:
    return (ln.get("material_type") or ln.get("type")) == "accessory" and not ln.get("is_cut_panel")


def _jenis(ln: dict) -> str:
    if ln.get("is_cut_panel"):
        return "potongan_kain"
    return ln.get("material_type") or ln.get("type") or ""


def _num(v):
    try:
        return float(v) if v not in (None, "") else None
    except (TypeError, ValueError):
        return None


def build(db):
    models = list(db.rahaza_models.find({}, {"_id": 0}).sort("code", 1))
    variants = defaultdict(list)
    for v in db.rahaza_model_variants.find({}, {"_id": 0}):
        variants[v["model_id"]].append(v)
    boms = defaultdict(list)
    for b in db.rahaza_boms.find({}, {"_id": 0}):
        boms[b["model_id"]].append(b)
    materials = {m["code"]: m for m in db.rahaza_materials.find({}, {"_id": 0, "code": 1, "unit_cost": 1, "unit": 1, "base_uom": 1, "active": 1})}
    mat_by_id = {m["id"]: m for m in db.rahaza_materials.find({}, {"_id": 0, "id": 1, "code": 1, "unit_cost": 1, "unit": 1, "base_uom": 1})}

    detail, summary, model_rows = [], [], []
    stats = Counter()

    for m in models:
        mid = m["id"]
        base = {"kode_model": m.get("code"), "nama_model": m.get("name"), "kategori": m.get("category_name") or m.get("category"),
                "model_aktif": "YA" if m.get("active") is not False else "TIDAK", "hpp_model_db": _num(m.get("hpp"))}
        mv = sorted(variants.get(mid, []), key=lambda v: (v.get("color_code") or "", v.get("size_code") or ""))
        mb = boms.get(mid, [])
        active_boms = [b for b in mb if b.get("active") is not False]
        by_key = defaultdict(list)
        for b in mb:
            by_key[((b.get("color_code") or "").upper(), b.get("size_id"))].append(b)

        # --- ringkasan model (logika sama dengan laporan_kekosongan_bom.py) ---
        act = [v for v in mv if v.get("active") is not False]
        keys = {((b.get("color_code") or "").upper(), b.get("size_id")) for b in active_boms}
        no_bom = [v["sku"] for v in act if ((v.get("color_code") or "").upper(), v.get("size_id")) not in keys]
        bom_no_acc = [b.get("color_code") for b in active_boms if not any(_is_acc(ln) for ln in b.get("materials") or [])]
        if m.get("active") is False:
            st = "MODEL_NONAKTIF"
        elif mv and not act:
            st = "F_DIHENTIKAN"
        elif not mv:
            st = "A_TANPA_SKU"
        elif not active_boms:
            st = "B_ADA_SKU_NOL_BOM"
        elif no_bom:
            st = "C_SEBAGIAN_VARIAN_TANPA_BOM"
        elif len(bom_no_acc) == len(active_boms):
            st = "D_BOM_TANPA_AKSESORIS"
        elif bom_no_acc:
            st = "D2_SEBAGIAN_BOM_TANPA_AKSESORIS"
        else:
            st = "E_LENGKAP"
        stats[st] += 1
        model_rows.append({**base, "status_kelengkapan": st, "varian_total": len(mv), "varian_aktif": len(act), "bom_aktif": len(active_boms),
                           "varian_tanpa_bom": ", ".join(no_bom), "bom_tanpa_aksesoris": ", ".join(x or "" for x in bom_no_acc)})

        if not mv:
            detail.append({**base, "status_bom": "TANPA_SKU", "jumlah_baris_bom": 0})
            summary.append({**base, "status_bom": "TANPA_SKU", "baris_kain": 0, "baris_aksesoris": 0, "baris_total": 0, "biaya_bom": None})
            continue

        for v in mv:
            vbase = {**base, "sku": v.get("sku"), "kode_warna": v.get("color_code"), "nama_warna": v.get("color_name"),
                     "ukuran": v.get("size_code"), "sku_aktif": "YA" if v.get("active") is not False else "TIDAK"}
            cands = by_key.get(((v.get("color_code") or "").upper(), v.get("size_id")), [])
            actives = [b for b in cands if b.get("active") is not False]
            bom = max(actives, key=lambda b: b.get("version") or 0) if actives else (max(cands, key=lambda b: b.get("version") or 0) if cands else None)

            if bom is None:
                detail.append({**vbase, "status_bom": "TANPA_BOM", "jumlah_baris_bom": 0})
                summary.append({**vbase, "status_bom": "TANPA_BOM", "baris_kain": 0, "baris_aksesoris": 0, "baris_total": 0, "biaya_bom": None})
                continue

            lines = bom.get("materials") or []
            n_acc = sum(1 for ln in lines if _is_acc(ln))
            n_fab = len(lines) - n_acc
            if bom.get("active") is False:
                status = "BOM_NONAKTIF"
            elif not lines:
                status = "BOM_KOSONG"
            elif n_acc == 0:
                status = "BOM_TANPA_AKSESORIS"
            else:
                status = "BOM_LENGKAP"
            bbase = {**vbase, "status_bom": status, "bom_id": bom.get("id"), "versi_bom": bom.get("version"),
                     "bom_aktif": "YA" if bom.get("active") is not False else "TIDAK", "jumlah_baris_bom": len(lines)}
            total = 0.0
            if not lines:
                detail.append(bbase)
            for i, ln in enumerate(lines, 1):
                qb, uc = _num(ln.get("qty_base")), _num(ln.get("unit_cost_base"))
                biaya = round(qb * uc, 2) if qb is not None and uc is not None else None
                total += biaya or 0
                mat = materials.get(ln.get("code")) or mat_by_id.get(ln.get("material_id")) or {}
                detail.append({**bbase, "no_baris": i, "jenis_komponen": _jenis(ln), "kode_material": ln.get("code"), "nama_material": ln.get("name"),
                               "qty": _num(ln.get("qty")), "satuan": ln.get("unit"), "qty_dasar": qb, "satuan_dasar": ln.get("unit_base"),
                               "harga_per_satuan_dasar": uc, "biaya_baris": biaya,
                               "harga_master_saat_ini": _num(mat.get("unit_cost")), "satuan_master": mat.get("base_uom") or mat.get("unit"),
                               "material_terhubung": "TIDAK" if ln.get("unlinked") else ("YA" if mat else "TIDAK ADA DI MASTER"),
                               "kain_sumber": ln.get("source_material_code") or "", "catatan_baris": ln.get("notes") or ln.get("uom_note") or ""})
            summary.append({**bbase, "baris_kain": n_fab, "baris_aksesoris": n_acc, "baris_total": len(lines), "biaya_bom": round(total, 2)})

    return detail, summary, model_rows, stats


def _write_sheet(ws, cols, rows, status_key="status_bom"):
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = FILL_HDR
        c.alignment = Alignment(vertical="center", wrap_text=True)
    fills = {"TANPA_BOM": FILL_NO_BOM, "BOM_KOSONG": FILL_NO_BOM, "BOM_TANPA_AKSESORIS": FILL_NO_ACC, "TANPA_SKU": FILL_NO_SKU,
             "BOM_NONAKTIF": FILL_INACTIVE, "A_TANPA_SKU": FILL_NO_SKU, "B_ADA_SKU_NOL_BOM": FILL_NO_BOM, "C_SEBAGIAN_VARIAN_TANPA_BOM": FILL_NO_BOM,
             "D_BOM_TANPA_AKSESORIS": FILL_NO_ACC, "D2_SEBAGIAN_BOM_TANPA_AKSESORIS": FILL_NO_ACC, "F_DIHENTIKAN": FILL_INACTIVE, "MODEL_NONAKTIF": FILL_INACTIVE}
    for r in rows:
        ws.append([r.get(c) for c in cols])
        f = fills.get(r.get(status_key))
        if f:
            for c in ws[ws.max_row]:
                c.fill = f
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(cols, 1):
        width = max(len(col), *(len(str(r.get(col) or "")) for r in rows[:400])) if rows else len(col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 48)


def main():
    out = sys.argv[sys.argv.index("--out") + 1] if "--out" in sys.argv else "/app/private/golive/EXPORT_SKU_BOM_DA.xlsx"
    db = MongoClient(os.environ["MONGO_URL"])[os.environ["DB_NAME"]]
    detail, summary, model_rows, stats = build(db)

    wb = Workbook()
    ws = wb.active
    ws.title = "DETAIL_BOM"
    _write_sheet(ws, DETAIL_COLS, detail)
    _write_sheet(wb.create_sheet("RINGKASAN_SKU"), SUMMARY_COLS, summary)
    _write_sheet(wb.create_sheet("RINGKASAN_MODEL"), MODEL_COLS, model_rows, status_key="status_kelengkapan")

    sc = Counter(s["status_bom"] for s in summary)
    ket = wb.create_sheet("KETERANGAN")
    ket.append(["Ekspor SKU + BOM dari DB", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"), f"DB={os.environ['DB_NAME']}"])
    ket.append([])
    ket.append(["Angka", "Jumlah"])
    ket.append(["Model", len(model_rows)])
    ket.append(["Baris RINGKASAN_SKU (SKU + model tanpa SKU)", len(summary)])
    ket.append(["Baris DETAIL_BOM", len(detail)])
    for k in sorted(sc):
        ket.append([f"  status_bom = {k}", sc[k]])
    for k in sorted(stats):
        ket.append([f"  kelengkapan model = {k}", stats[k]])
    ket.append([])
    ket.append(["status_bom", "Arti"])
    ket.append(["BOM_LENGKAP", "SKU punya BOM aktif dengan ≥1 baris aksesoris"])
    ket.append(["BOM_TANPA_AKSESORIS", "BOM aktif ada tapi hanya potongan kain (aksesoris belum diisi) — KUNING"])
    ket.append(["BOM_KOSONG", "BOM aktif ada tapi 0 baris komponen"])
    ket.append(["TANPA_BOM", "SKU ada, belum ada BOM sama sekali — ORANYE"])
    ket.append(["BOM_NONAKTIF", "Hanya ada BOM nonaktif (biasanya SKU dihentikan)"])
    ket.append(["TANPA_SKU", "Model belum punya varian warna+ukuran → BOM tidak mungkin dibuat — ABU"])
    ket.append([])
    ket.append(["Kolom", "Arti"])
    ket.append(["qty / satuan", "Angka & satuan seperti diinput di BOM (mis. 60 cm)"])
    ket.append(["qty_dasar / satuan_dasar", "Hasil konversi ke satuan dasar material (mis. 0.6 m)"])
    ket.append(["harga_per_satuan_dasar", "Harga yang tersimpan di baris BOM saat HPP terakhir dihitung"])
    ket.append(["biaya_baris", "qty_dasar × harga_per_satuan_dasar (dasar HPP)"])
    ket.append(["harga_master_saat_ini", "unit_cost di master material sekarang — bila beda dari harga_per_satuan_dasar, HPP perlu dihitung ulang"])
    ket.append(["jenis_komponen", "potongan_kain (CUT-*, mewakili kain), accessory, fabric"])
    ket.append(["kain_sumber", "Kode kain asal untuk baris potongan_kain"])
    ket.append(["hpp_model_db", "Field hpp di rahaza_models (hasil hitung BOM terakhir)"])
    ket.column_dimensions["A"].width = 46
    ket.column_dimensions["B"].width = 100

    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)
    print(f"✓ {out}")
    print(f"  model={len(model_rows)} ringkasan_sku={len(summary)} detail={len(detail)}")
    for k in sorted(sc):
        print(f"  status_bom {k:22s} {sc[k]}")
    for k in sorted(stats):
        print(f"  model {k:34s} {stats[k]}")


if __name__ == "__main__":
    main()
