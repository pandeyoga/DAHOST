"""Analisis berkas harga owner (revisihargaaksesoris.xlsx) vs master: rencana harga per satuan dasar."""
import sys
from collections import Counter

from openpyxl import load_workbook
from pymongo import MongoClient

sys.path.insert(0, "/app/backend")
from core.bom_uom import PACKAGING_UNITS, global_factor, norm_unit  # noqa: E402

db = MongoClient("mongodb://localhost:27017")["test_database"]
mats = {m["code"]: m for m in db.rahaza_materials.find({}, {"_id": 0})}


def f(v):
    try:
        return float(str(v).replace(",", ""))
    except (TypeError, ValueError):
        return None


wb = load_workbook("/app/private/golive/revisihargaaksesoris.xlsx", read_only=True, data_only=True)


def sheet(name):
    ws = wb[name]
    rows = [list(r) + [None] * 5 for r in ws.iter_rows(values_only=True)]
    ix = {c: i for i, c in enumerate(rows[0]) if c}
    return {r[ix["kode"]]: {k: r[ix[k]] for k in ("satuan_dasar", "satuan_beli", "isi_per_satuan_beli", "harga_per_satuan_beli")} for r in rows[1:] if r[ix["kode"]]}


M, L = sheet("MATERIAL"), sheet("Sheet1lookup")
diff = [(k, M[k], L[k]) for k in L if k in M and (f(M[k]["harga_per_satuan_beli"]) != f(L[k]["harga_per_satuan_beli"]) or f(M[k]["isi_per_satuan_beli"]) != f(L[k]["isi_per_satuan_beli"]))]
print("lookup vs MATERIAL beda:", len(diff))
for d in diff[:10]:
    print("  ", d)


def plan_price(code, r):
    m = mats[code]
    B, SD, SB = norm_unit(m.get("base_uom") or m.get("unit")), norm_unit(r["satuan_dasar"]), norm_unit(r["satuan_beli"])
    I, H = f(r["isi_per_satuan_beli"]) or 1, f(r["harga_per_satuan_beli"]) or 0
    per_sd = H / I
    if SD == B:
        return B, per_sd, "sama", I, H, SB, SD
    gf = global_factor(SD, B)
    if gf:
        return B, per_sd / gf, "konversi", I, H, SB, SD
    if B in PACKAGING_UNITS and SB == B:
        return B, H, "pcs_per_kemasan", I, H, SB, SD
    return SD, per_sd, "ganti_base", I, H, SB, SD


if __name__ == "__main__":
    cnt = Counter()
    plan = []
    for code, r in M.items():
        nb, uc, how, I, H, SB, SD = plan_price(code, r)
        cnt[how] += 1
        old = float(mats[code].get("unit_cost") or 0)
        plan.append((code, mats[code]["name"][:34], norm_unit(mats[code].get("base_uom") or mats[code].get("unit")), SD, SB, I, H, round(uc, 2), old, round(uc / old, 2) if old else None, how, nb))
    print(cnt)
    print("\nPerubahan besar (>5x / <0.2x / dari 0):")
    for p in sorted(plan, key=lambda p: (p[9] is None, -(p[9] or 0))):
        if p[9] is None or p[9] > 5 or p[9] < 0.2:
            print("  ", p)
    print("\nganti_base:")
    for p in plan:
        if p[10] == "ganti_base":
            print("  ", p)
