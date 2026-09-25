"""KONSOLIDASI HARGA MATERIAL · VARIAN · BOM · POTONGAN · R&D — dari 4 berkas owner (2026-09-23).

Jalankan dari /app/backend (env sudah di-load):
  python ../scripts/konsolidasi/konsolidasi.py            # pratinjau (dry-run), tulis laporan JSON
  python ../scripts/konsolidasi/konsolidasi.py --apply    # terapkan ke DB

Sumber & prioritas (keputusan owner):
  harga material  = revisihargaaksesoris.xlsx / MATERIAL  (satuan dasar mengikuti isian owner)
  BOM per SKU     = revisiBOMDA.xlsx / DETAIL_BOM (Hanny DA-1509 disusun ulang dari Hanni DA-1510)
  varian baru     = revisirndda.xlsx / VARIAN_BARU
  BOM SKU tanpa isian = revisirndda BOM_OTOMATIS/BOM_AKSESORIS (kelompok per varian) → varian saudara
  Semua BOM wajib punya POTONGAN (CUT-…); kain sumber baru dibuat mengikuti pola KN-<FAM>-<WARNA>.
  Hangtag A-HTG-0003 + Pin A-PIN-0001 1 pcs di SEMUA SKU aktif. Versi BOM tetap 1.
"""
from __future__ import annotations

import asyncio
import json
import os
import re
import sys
import uuid
from collections import Counter, defaultdict
from datetime import datetime, timezone

sys.path.insert(0, "/app/backend")
from motor.motor_asyncio import AsyncIOMotorClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from core.bom_uom import annotate_line, norm_unit  # noqa: E402
from core.cut_panel_master import bom_line_for_panel, ensure_panel  # noqa: E402

DIR = "/app/private/golive"
F_HARGA, F_BOM, F_RND = f"{DIR}/revisihargaaksesoris.xlsx", f"{DIR}/revisiBOMDA.xlsx", f"{DIR}/revisirndda.xlsx"
USER = {"id": "system", "name": "konsolidasi_2026-09-23"}
CODE_MAP = {"A-K18-0045": "A-K18-0038"}          # keputusan owner #4
MODEL_MAP = {"DA-3306": "DA-3302"}                # keputusan owner #2
SKIP_ALLSIZE = {"DA-3604", "DA-4201"}             # typo ALLSIZE (SKU tetap M)
DISCONTINUED = {"DA-2104", "DA-2107", "DA-2108", "DA-3601", "DA-3602"}
HANGTAG, PIN = "A-HTG-0003", "A-PIN-0001"
TBD_FAM = "TBD"
LABEL_BY_SIZE = {"S": "A-LBL-0007", "M": "A-LBL-0008", "L": "A-LBL-0009", "XL": "A-LBL-0010", "XXL": "A-LBL-0011"}
# isi kemasan yang belum diisi owner (isi=1 padahal BOM memakai pcs): (isi, satuan_pakai, keterangan) — menunggu konfirmasi owner
PACK_OVERRIDE: dict[str, tuple] = {"A-LBL-0004": (1000, "pcs", "jawaban owner 2026-09-23"), "A-PIN-0001": (5000, "pcs", "jawaban owner 2026-09-23")}
LINE_UNIT_FIX = {("A-BAB-0001", "pcs"): (100.0, "cm")}  # 1 pcs babud = 100 cm (jawaban owner)


def now():
    return datetime.now(timezone.utc)


def clean(s):
    return re.sub(r"[\u200b\u00a0]", "", str(s or "")).strip()


def num(v):
    try:
        return float(str(v).replace(",", ""))
    except (TypeError, ValueError):
        return None


def rows_of(path, sheet):
    ws = load_workbook(path, read_only=True, data_only=True)[sheet]
    rows = [list(r) + [None] * 12 for r in ws.iter_rows(values_only=True)]
    ix = {c: i for i, c in enumerate(rows[0]) if c}
    return [{k: r[i] for k, i in ix.items()} for r in rows[1:] if any(x not in (None, "") for x in r[:10])]


# ═══════════════════════════ 1. HARGA MATERIAL ═══════════════════════════
def plan_prices(mats: dict, line_units: dict | None = None) -> tuple[list, list]:
    """line_units: {kode: satuan yang paling sering dipakai baris BOM} — dipakai bila owner menulis satuan_dasar = kemasan (pack/roll) tapi isi > 1."""
    from core.bom_uom import PACKAGING_UNITS
    plan, warn = [], []
    for r in rows_of(F_HARGA, "MATERIAL"):
        code = clean(r["kode"])
        m = mats.get(code)
        if not m:
            warn.append(f"harga: {code} tidak ada di master")
            continue
        sd, sb = norm_unit(r["satuan_dasar"]), norm_unit(r["satuan_beli"])
        isi, harga = num(r["isi_per_satuan_beli"]) or 1.0, num(r["harga_per_satuan_beli"]) or 0.0
        if harga <= 0:
            warn.append(f"harga: {code} harga_per_satuan_beli kosong/0 → dilewati")
            continue
        if sd == sb and sd in PACKAGING_UNITS and isi > 1:  # isi = jumlah satuan pakai per kemasan; satuan pakai = satuan baris BOM
            sd = (line_units or {}).get(code) or "pcs"
            warn.append(f"harga: {code} satuan_dasar '{sb}' = kemasan, isi {isi:g} → dibaca {isi:g} {sd}/{sb}")
        uc = round(harga / isi, 4)
        owner_uc = num(r.get("harga_per_satuan_dasar_sekarang"))
        if owner_uc and abs(owner_uc - uc) / max(owner_uc, 1e-9) > 0.02:
            warn.append(f"harga: {code} harga/isi={uc} ≠ kolom owner {owner_uc}")
        plan.append({"code": code, "name": m["name"], "base_before": norm_unit(m.get("base_uom") or m.get("unit")), "base": sd,
                     "buy_unit": sb, "pack_size": isi, "buy_price": harga, "unit_cost": uc, "unit_cost_before": float(m.get("unit_cost") or 0)})
    return plan, warn


async def apply_prices(db, plan):
    n = 0
    for p in plan:
        uoms = [{"code": p["base"], "name": p["base"].upper(), "factor": 1.0, "is_base": True, "level": 0}]
        if p["buy_unit"] != p["base"]:
            uoms.append({"code": p["buy_unit"], "name": p["buy_unit"].upper(), "factor": p["pack_size"], "is_base": False, "level": 1,
                         "notes": f"1 {p['buy_unit']} = {p['pack_size']:g} {p['base']}"})
        await db.rahaza_materials.update_one({"code": p["code"]}, {"$set": {
            "unit": p["base"], "base_uom": p["base"], "issue_uom": p["base"], "display_uom": p["base"], "purchase_uom": p["buy_unit"],
            "pack_unit": p["buy_unit"], "pack_size": p["pack_size"], "uoms": uoms, "unit_cost": p["unit_cost"], "cost_method": "moving_average",
            "cost_source": "harga_awal_owner", "price_updated_at": now(), "updated_at": now(),
            "konsolidasi_note": f"2026-09-23 owner: {p['pack_size']:g} {p['base']}/{p['buy_unit']} @ Rp {p['buy_price']:,.0f}"}})
        if abs(p["unit_cost"] - p["unit_cost_before"]) > 1e-9 or p["base"] != p["base_before"]:
            await db.rahaza_material_cost_history.insert_one({"id": str(uuid.uuid4()), "material_code": p["code"], "unit_cost": p["unit_cost"], "before": p["unit_cost_before"],
                                                              "unit": p["base"], "unit_before": p["base_before"], "source": "konsolidasi_owner_2026-09-23", "by": USER["name"], "at": now()})
        n += 1
    return n


# ═══════════════════════════ 2. VARIAN BARU ═══════════════════════════
def plan_variants(models_by_code, variants, colors, sizes):
    have = {(v["model_code"], v["color_code"].upper(), v["size_code"].upper()) for v in variants}
    new, skipped, seen = [], [], set()
    for r in rows_of(F_RND, "VARIAN_BARU"):
        mc = MODEL_MAP.get(clean(r["kode_model"]).upper(), clean(r["kode_model"]).upper())
        cc, sz = clean(r["kode_warna"]).upper(), clean(r["ukuran"]).upper()
        if not (mc and cc and sz):
            continue
        key = (mc, cc, sz)
        if key in seen:
            skipped.append(f"{mc}-{cc}-{sz}: duplikat di VARIAN_BARU")
            continue
        seen.add(key)
        m = models_by_code.get(mc)
        if not m:
            skipped.append(f"{mc}-{cc}-{sz}: model tidak ada")
        elif mc in DISCONTINUED:
            skipped.append(f"{mc}-{cc}-{sz}: model dihentikan — tidak diaktifkan (keputusan owner)")
        elif mc in SKIP_ALLSIZE and sz == "ALLSIZE":
            skipped.append(f"{mc}-{cc}-{sz}: ALLSIZE dianggap typo (SKU tetap M)")
        elif cc not in colors:
            skipped.append(f"{mc}-{cc}-{sz}: kode warna tidak ada di master")
        elif sz not in sizes:
            skipped.append(f"{mc}-{cc}-{sz}: ukuran tidak ada di master")
        elif key in have:
            pass
        else:
            new.append({"model": m, "color": colors[cc], "size": sizes[sz], "sku": f"{mc}-{cc}-{sz}"})
    return new, skipped


async def apply_variants(db, new):
    docs = []
    for n in new:
        m, c, s = n["model"], n["color"], n["size"]
        docs.append({"id": str(uuid.uuid4()), "model_id": m["id"], "model_code": m["code"], "model_name": m["name"], "size_id": s["id"], "size_code": s["code"],
                     "color_id": c["id"], "color_code": c["code"], "color_name": c["name"], "color_hex": c.get("hex"), "sku": n["sku"], "barcode": "", "notes": "",
                     "active": True, "created_at": now(), "updated_at": now(), "created_from": "konsolidasi_owner_2026-09-23"})
    if docs:
        await db.rahaza_model_variants.insert_many(docs)
        from utils.variant_ssot import ensure_fg_material
        for d in docs:
            await ensure_fg_material(db, d, user=USER)
    return docs


# ═══════════════════════════ 3. BOM PER SKU ═══════════════════════════
def bom_from_revisi() -> dict[str, list]:
    out = defaultdict(list)
    for r in rows_of(F_BOM, "DETAIL_BOM"):
        sku, code = clean(r["sku"]), clean(r["kode_material"])
        if not sku or not code:
            continue
        code = CODE_MAP.get(code, code)
        qty = num(r["qty"])
        if qty is None or qty <= 0:
            continue
        out[sku].append({"code": code, "name": clean(r["nama_material"]), "qty": qty, "unit": norm_unit(r["satuan"]) or "pcs"})
    return out


def groups_from_rnd() -> dict[str, list]:
    """kelompok BOM per model dari revisirndda (BOM_OTOMATIS + BOM_AKSESORIS): [{varian:set warna-nama lower, lines}]"""
    out = defaultdict(list)
    for sheet in ("BOM_OTOMATIS", "BOM_AKSESORIS"):
        cur = None
        for r in rows_of(F_RND, sheet):
            if r["kode_model"]:
                cur = {"varian": {clean(x).lower() for x in str(r["varian"] or "").split(",") if clean(x)}, "lines": []}
                out[clean(r["kode_model"]).upper()].append(cur)
            if cur is None or not r["kode_material"]:
                continue
            q = clean(r["qty_per_pcs"]).split()
            qty = num(q[0]) if q else None
            unit = norm_unit(q[1]) if len(q) > 1 else norm_unit(r["satuan"]) or "pcs"
            if qty:
                cur["lines"].append({"code": CODE_MAP.get(clean(r["kode_material"]), clean(r["kode_material"])), "name": clean(r["nama_material"]), "qty": qty, "unit": unit})
    return out


def dedupe(lines):
    seen, out = set(), []
    for ln in lines:
        if ln["code"] in seen:
            continue
        seen.add(ln["code"])
        fix = LINE_UNIT_FIX.get((ln["code"], ln["unit"]))
        if fix:
            ln = dict(ln, qty=ln["qty"] * fix[0], unit=fix[1])
        out.append(ln)
    return out


def hanny_from_hanni(hanni_lines, color, size):
    """DA-1509 (M) = pola DA-1510 (XL) warna sama: potongan ganti kode, label XL→M, karet 62→56 cm."""
    out = []
    for ln in hanni_lines:
        ln = dict(ln)
        if ln["code"].startswith("CUT-"):
            ln["code"] = f"CUT-DA-1509-{color}-{size}"
        elif ln["code"] == "A-LBL-0010":
            ln["code"] = "A-LBL-0008"
            ln["name"] = "Label size M putih rajut 1 Pcs"
        elif ln["code"] == "A-KRT-0002" and ln["unit"] == "cm":
            ln["qty"] = 56.0
        out.append(ln)
    return out


class Konsolidasi:
    def __init__(self, db, apply: bool):
        self.db, self.apply = db, apply
        self.report = {"apply": apply, "at": now().isoformat(), "warnings": [], "notes": []}

    async def load(self):
        db = self.db
        self.mats = {m["code"]: m for m in await db.rahaza_materials.find({}, {"_id": 0}).to_list(50000)}
        self.models = {m["id"]: m for m in await db.rahaza_models.find({}, {"_id": 0, "sop_steps": 0}).to_list(5000)}
        self.models_by_code = {m["code"]: m for m in self.models.values()}
        self.variants = await db.rahaza_model_variants.find({}, {"_id": 0}).to_list(50000)
        self.colors = {c["code"].upper(): c for c in await db.rahaza_colors.find({}, {"_id": 0}).to_list(5000)}
        self.sizes = {s["code"].upper(): s for s in await db.rahaza_sizes.find({}, {"_id": 0}).to_list(100)}
        self.boms = await db.rahaza_boms.find({}, {"_id": 0}).to_list(50000)

    # ── kain sumber ──
    def fabric_family_of_model(self, model_code):
        fams = Counter()
        for m in self.mats.values():
            if m.get("is_cut_panel") and m.get("model_code") == model_code and (m.get("source_material_code") or "").startswith("KN-"):
                fams[m["source_material_code"].split("-")[1]] += 1
        if not fams:  # model se-nama (Lunara 1201 ↔ 1209…)
            name = self.models_by_code[model_code]["name"].lower()
            for mm in self.models_by_code.values():
                if mm["name"].lower() == name and mm["code"] != model_code:
                    for m in self.mats.values():
                        if m.get("is_cut_panel") and m.get("model_code") == mm["code"] and (m.get("source_material_code") or "").startswith("KN-"):
                            fams[m["source_material_code"].split("-")[1]] += 1
        if not fams:  # palet warna model hanya ada di SATU keluarga kain → tebakan (ditandai di laporan)
            palette = {v["color_code"].upper() for v in self.active_variants if v["model_code"] == model_code}
            fam_colors = defaultdict(set)
            for m in self.mats.values():
                if m.get("type") == "fabric" and m["code"].startswith("KN-") and not m.get("is_cut_panel") and len(m["code"].split("-")) >= 3:
                    fam_colors[m["code"].split("-")[1]].add(m["code"].split("-")[2])
            cands = [f for f, cs in fam_colors.items() if f != TBD_FAM and palette and palette <= cs]
            if len(cands) == 1:
                self.report.setdefault("fabric_guess", []).append(f"{model_code}: keluarga kain KN-{cands[0]} ditebak dari palet warna {sorted(palette)}")
                return cands[0]
        return fams.most_common(1)[0][0] if fams else None

    async def ensure_fabric(self, fam, color):
        code = f"KN-{fam}-{color['code']}"
        if code in self.mats:
            return self.mats[code]
        tpl = next((m for m in self.mats.values() if m["code"].startswith(f"KN-{fam}-") and not m.get("is_cut_panel")), None)
        doc = {**{k: v for k, v in (tpl or {}).items() if k not in ("id", "code", "color", "color_code", "created_at", "updated_at", "import_batch", "import_source", "konsolidasi_note")},
               "id": str(uuid.uuid4()), "code": code, "type": "fabric", "color": color["name"], "color_code": color["code"], "active": True,
               "created_at": now(), "updated_at": now(), "created_from": "konsolidasi_owner_2026-09-23"}
        if not tpl:
            doc.update({"name": "KAIN BELUM DITENTUKAN", "unit": "kg", "base_uom": "kg", "purchase_uom": "kg", "issue_uom": "kg", "display_uom": "kg",
                        "uoms": [{"code": "kg", "name": "KG", "factor": 1.0, "is_base": True, "level": 0}], "unit_cost": 0.0, "min_stock": 0,
                        "category_code": "FABRIC", "category_name": "Kain/Fabric", "cost_method": "moving_average",
                        "category": self.mats["KN-K24-BRG"].get("category"), "category_id": self.mats["KN-K24-BRG"].get("category_id")})
            doc["konsolidasi_note"] = "Kain sumber belum ditentukan owner — harga 0, ganti kode/harga saat kain diketahui"
        else:
            doc["konsolidasi_note"] = f"Warna baru untuk kain {tpl['name']} (disalin dari {tpl['code']})"
        self.mats[code] = doc
        self.report.setdefault("fabrics_created", []).append(f"{code} ← {doc.get('name')} · Rp {doc.get('unit_cost', 0):,.0f}")
        if self.apply:
            await self.db.rahaza_materials.insert_one(dict(doc))
        return doc

    async def ensure_cut(self, model, variant):
        code = f"CUT-{model['code']}-{variant['color_code']}-{variant['size_code']}"
        if code in self.mats:
            return self.mats[code]
        fam = self.fabric_family_of_model(model["code"]) or TBD_FAM
        fabric = await self.ensure_fabric(fam, self.colors[variant["color_code"].upper()])
        panel, created = await ensure_panel(self.db, model=model, color_code=variant["color_code"], color_name=variant["color_name"], size_code=variant["size_code"],
                                            size_id=variant["size_id"], source_fabric=fabric, created_from="konsolidasi_owner_2026-09-23", apply=self.apply)
        self.mats[panel["code"]] = panel
        self.report.setdefault("panels_created", []).append(f"{panel['code']} ← {fabric['code']}" + (" (KAIN BELUM DITENTUKAN)" if fam == TBD_FAM else ""))
        return panel

    async def ensure_tlk(self):
        if "A-TLK-0001" in self.mats:
            return
        tpl = self.mats["A-KRT-0008"]
        doc = {**{k: v for k, v in tpl.items() if k not in ("id", "code", "name", "created_at", "updated_at", "import_batch", "import_source", "konsolidasi_note", "sub_category")},
               "id": str(uuid.uuid4()), "code": "A-TLK-0001", "name": "Tali Kancing", "unit": "cm", "base_uom": "cm", "issue_uom": "cm", "display_uom": "cm", "purchase_uom": "cm",
               "pack_unit": "cm", "pack_size": 1, "uoms": [{"code": "cm", "name": "CM", "factor": 1.0, "is_base": True, "level": 0}], "unit_cost": 0.0, "active": True,
               "created_at": now(), "updated_at": now(), "created_from": "konsolidasi_owner_2026-09-23", "konsolidasi_note": "Material baru dari revisiBOMDA (Rachel set) — harga belum diisi owner"}
        self.mats["A-TLK-0001"] = doc
        self.report["notes"].append("Material baru A-TLK-0001 Tali Kancing (cm, harga 0) dibuat")
        if self.apply:
            await self.db.rahaza_materials.insert_one(dict(doc))

    # ── target BOM per SKU ──
    def pick_group(self, groups, variant):
        if not groups:
            return None, None
        cname = variant["color_name"].lower()
        exact = [g for g in groups if cname in g["varian"] or variant["color_code"].lower() in g["varian"]]
        if exact:
            return exact[0], "kelompok per warna (revisirndda)"
        if len(groups) == 1:
            return groups[0], "satu kelompok model (revisirndda)"
        colored = [g for g in groups if any("warna" in ln["name"].lower() for ln in g["lines"])]
        if not colored:  # kelompok terpecah tanpa pembeda warna → gabung semua baris
            return {"varian": set(), "lines": [ln for g in groups for ln in g["lines"]]}, "gabungan kelompok model (revisirndda)"
        tok = cname[:4]
        hit = [g for g in colored if any(tok in ln["name"].lower() for ln in g["lines"] if "warna" in ln["name"].lower())]
        if hit:
            return hit[0], "kelompok ditebak dari nama bahan berwarna (revisirndda)"
        common = [ln for ln in groups[0]["lines"] if all(any(x["code"] == ln["code"] for x in g["lines"]) for g in groups[1:])]
        if common:
            return {"varian": set(), "lines": common}, f"baris umum semua kelompok — aksesoris warna {variant['color_name']} BELUM ADA (revisirndda)"
        return None, None

    def sister_lines(self, target_by_sku, model_code, variant):
        for sku, lines in target_by_sku.items():
            if sku.startswith(model_code + "-") and any(not ln["code"].startswith("CUT-") for ln in lines):
                return sku, [ln for ln in lines if not ln["code"].startswith("CUT-")]
        return None, None

    async def build_targets(self, new_variants):
        rev, grp = bom_from_revisi(), groups_from_rnd()
        active = [v for v in self.variants if v.get("active") is not False] + new_variants
        target, source = {}, {}
        # 1) revisiBOMDA (Hanny disusun ulang dari Hanni)
        for v in active:
            sku = v["sku"]
            if v["model_code"] == "DA-1509":
                hanni = rev.get(f"DA-1510-{v['color_code']}-XL")
                if hanni:
                    target[sku], source[sku] = dedupe(hanny_from_hanni(hanni, v["color_code"], v["size_code"])), "pola Hanni DA-1510 (baris Hanny bergeser)"
                continue
            if sku in rev and any(not ln["code"].startswith("CUT-") for ln in rev[sku]):
                target[sku], source[sku] = dedupe(rev[sku]), "revisiBOMDA"
        # 2) sisa: kelompok revisirndda → varian saudara
        for v in active:
            sku = v["sku"]
            if sku in target:
                continue
            g, how = self.pick_group(grp.get(v["model_code"], []), v)
            if g and g["lines"]:
                lines = dedupe(g["lines"])
                # label ukuran ikut ukuran SKU
                lbl = LABEL_BY_SIZE.get(v["size_code"].upper())
                lines = [dict(ln, code=lbl) if lbl and ln["code"] in LABEL_BY_SIZE.values() and ln["code"] != lbl else ln for ln in lines]
                target[sku], source[sku] = lines, how
                continue
            ssku, lines = self.sister_lines(target, v["model_code"], v)
            if lines:
                target[sku], source[sku] = [dict(ln) for ln in lines], f"disalin dari varian saudara {ssku} — CEK aksesoris berwarna"
            else:
                target[sku], source[sku] = [], "TIDAK ADA SUMBER — hanya potongan + hangtag + pin"
        return target, source, active

    async def apply_boms(self, target, source, active):
        by_key = defaultdict(list)
        for b in self.boms:
            by_key[(b["model_id"], (b.get("color_code") or "").upper(), b.get("size_id"))].append(b)
        stats, hpp_lines = Counter(), []
        for v in active:
            sku, model = v["sku"], self.models[v["model_id"]]
            lines = [ln for ln in target.get(sku, []) if not ln["code"].startswith("CUT-")]
            for extra in (HANGTAG, PIN):
                if extra not in {ln["code"] for ln in lines}:
                    lines.append({"code": extra, "name": self.mats[extra]["name"], "qty": 1.0, "unit": "pcs"})
            panel = await self.ensure_cut(model, v)
            key = (v["model_id"], v["color_code"].upper(), v["size_id"])
            cands = [b for b in by_key.get(key, []) if b.get("active") is not False] or by_key.get(key, [])
            existing = max(cands, key=lambda b: b.get("version") or 0) if cands else None
            old_lines = {ln["code"]: ln for ln in (existing or {}).get("materials") or []}
            cut_line = old_lines.get(panel["code"]) or bom_line_for_panel(panel)
            new_materials = [cut_line]
            for ln in lines:
                m = self.mats.get(ln["code"])
                if not m:
                    self.report["warnings"].append(f"{sku}: {ln['code']} tidak ada di master → dilewati")
                    continue
                base = dict(old_lines.get(ln["code"]) or {"notes": "", "unlinked": False})
                base.update({"material_id": m["id"], "code": m["code"], "name": m["name"], "material_type": m.get("type") or "accessory", "qty": ln["qty"], "unit": ln["unit"]})
                new_materials.append(annotate_line(base, m))
            new_materials[0] = annotate_line({**cut_line, "material_type": "fabric", "is_cut_panel": True}, panel)
            new_materials[0]["unit_cost_base"] = float(panel.get("unit_cost") or 0)
            acc = sum(1 for ln in new_materials if ln.get("material_type") == "accessory")
            for ln in new_materials:
                if ln.get("uom_status") == "mismatch":
                    self.report.setdefault("uom_mismatch", []).append(f"{sku}: {ln['code']} {ln['qty']:g} {ln['unit']} → satuan dasar {ln.get('unit_base')} (dihitung 1:1)")
            stats["bom_aksesoris_ok" if acc > 2 else "bom_hanya_hangtag_pin"] += 1
            if existing:
                same = [(ln["code"], ln["qty"], ln["unit"]) for ln in existing.get("materials") or []] == [(ln["code"], ln["qty"], ln["unit"]) for ln in new_materials]
                stats["bom_update" if not same else "bom_sama"] += 1
                if self.apply and (not same or existing.get("active") is False):
                    await self.db.rahaza_boms.update_one({"id": existing["id"]}, {"$set": {"materials": new_materials, "active": True, "is_active": True, "updated_at": now(),
                                                                                           "accessories_source": f"konsolidasi 2026-09-23: {source.get(sku)}"},
                                                                                  "$push": {"change_log": {"at": now(), "by": USER["name"], "note": f"konsolidasi owner ({source.get(sku)})",
                                                                                                           "lines_before": len(existing.get("materials") or []), "lines_after": len(new_materials)}}})
                for b in by_key.get(key, []):  # versi ganda → hanya 1 aktif
                    if b["id"] != existing["id"] and b.get("active") is not False:
                        stats["bom_versi_ganda_dinonaktifkan"] += 1
                        if self.apply:
                            await self.db.rahaza_boms.update_one({"id": b["id"]}, {"$set": {"active": False, "is_active": False, "deactivated_reason": "konsolidasi: versi ganda", "updated_at": now()}})
            else:
                stats["bom_baru"] += 1
                if self.apply:
                    await self.db.rahaza_boms.insert_one({"id": str(uuid.uuid4()), "model_id": v["model_id"], "size_id": v["size_id"], "color": v["color_name"], "color_code": v["color_code"].upper(),
                                                          "version": 1, "is_active": True, "active": True, "materials": new_materials, "notes": "", "created_by": USER["id"],
                                                          "created_at": now(), "updated_at": now(), "accessories_source": f"konsolidasi 2026-09-23: {source.get(sku)}", "change_log": []})
            hpp_lines.append({"sku": sku, "sumber": source.get(sku), "baris": len(new_materials), "aksesoris": acc,
                              "biaya_bom": round(sum(float(ln.get("qty_base") or 0) * float(ln.get("unit_cost_base") or 0) for ln in new_materials), 2)})
        self.report["bom_stats"] = dict(stats)
        self.report["bom_per_sku"] = hpp_lines
        self.report["bom_sumber"] = dict(Counter(s.split(" — ")[0].split(" (")[0] for s in source.values()))

    # ── R&D: tech pack + material R&D ──
    async def rnd_structure(self):
        db = self.db
        from routes.dewi_rnd_hpp import _normalize_techpack_payload
        created = 0
        styles = await db.dewi_rnd_styles.find({}, {"_id": 0}).to_list(5000)
        for st in styles:
            model = self.models_by_code.get(st["style_code"])
            if not model or model.get("active") is False:
                continue
            if await db.dewi_rnd_tech_packs.count_documents({"style_id": st["id"]}):
                continue
            vs = [v for v in await db.rahaza_model_variants.find({"model_id": model["id"], "active": True}, {"_id": 0}).to_list(500)]
            bom = await db.rahaza_boms.find_one({"model_id": model["id"], "active": True}, {"_id": 0})
            items = [{"material_code": ln["code"], "name": ln["name"], "qty": ln["qty"], "unit": ln["unit"], "material_type": ln.get("material_type")}
                     for ln in (bom or {}).get("materials") or [] if not ln.get("is_cut_panel")]
            fabrics, seen = [], set()
            for v in vs:
                p = self.mats.get(f"CUT-{model['code']}-{v['color_code']}-{v['size_code']}") or {}
                fc = p.get("source_material_code")
                if fc and fc not in seen:
                    seen.add(fc)
                    f = self.mats.get(fc) or {}
                    fabrics.append({"material_code": fc, "name": f.get("name", ""), "color_code": v["color_code"], "color_name": v["color_name"], "unit": f.get("unit", ""), "consumption": None})
            body = {"style_id": st["id"], "style_code": st["style_code"], "style_name": st["style_name"], "version": "v1", "title": f"Tech Pack v1 — {st['style_name']} ({st['style_code']})",
                    "description": "Struktur awal dibuat otomatis dari BOM & varian master (konsolidasi 2026-09-23). Detail konstruksi/ukuran menunggu R&D.",
                    "bom_items": items, "fabrics": fabrics, "fabric_consumption": [], "colorways": [{"code": v["color_code"], "name": v["color_name"]} for v in vs],
                    "size_columns": sorted({v["size_code"] for v in vs}), "measurements": [], "status": "draft"}
            if not self.apply:
                created += 1
                continue
            norm = await _normalize_techpack_payload(db, body)
            doc = {"id": str(uuid.uuid4()), **{k: body[k] for k in ("style_id", "style_code", "style_name", "version", "title", "description", "status")}, "doc_url": None, "doc_type": "pdf",
                   "bom_items": norm.get("bom_items", []), "bom_unlinked_count": norm.get("bom_unlinked_count", 0), "fabrics": norm.get("fabrics", []),
                   "fabric_consumption": norm.get("fabric_consumption", []), "fabric_consumption_off_list": norm.get("fabric_consumption_off_list", 0), "colorways": norm.get("colorways", []),
                   "construction_points": [], "construction_notes": "", "stitch_type": "", "seam_allowance_mm": 10, "size_grading_notes": "", "base_size": norm.get("base_size", "M"),
                   "size_range": norm.get("size_range", ""), "style_size_list": norm.get("style_size_list", []), "size_columns": norm.get("size_columns", []), "fit_categories": [],
                   "measurements": norm.get("measurements", []), "measurements_stats": norm.get("measurements_stats", {}), "approved_by": None, "approved_at": None, "is_latest": True,
                   "created_by": USER["id"], "created_by_name": USER["name"], "created_at": now(), "updated_at": now()}
            await db.dewi_rnd_tech_packs.insert_one(doc)
            await db.dewi_rnd_styles.update_one({"id": st["id"]}, {"$set": {"techpack_name": doc["title"], "updated_at": now()}})
            created += 1
        # material R&D per keluarga kain
        fam_created = 0
        fams = defaultdict(list)
        for m in self.mats.values():
            if m.get("type") == "fabric" and m["code"].startswith("KN-") and not m.get("is_cut_panel"):
                fams[m["code"].split("-")[1]].append(m)
        for fam, ms in fams.items():
            code = f"KN-{fam}"
            if await db.dewi_rnd_materials.count_documents({"material_code": code}):
                continue
            fam_created += 1
            if not self.apply:
                continue
            m0 = ms[0]
            colors = [{"color_id": self.colors[m["color_code"].upper()]["id"], "code": m["color_code"].upper(), "name": self.colors[m["color_code"].upper()]["name"], "hex": self.colors[m["color_code"].upper()].get("hex")}
                      for m in ms if m.get("color_code") and m["color_code"].upper() in self.colors]
            unit = m0.get("unit") or "kg"
            await db.dewi_rnd_materials.insert_one({"id": str(uuid.uuid4()), "material_code": code, "material_name": m0["name"], "category": "Kain", "vendor": "", "composition": m0.get("composition", ""),
                                                    "weight": m0.get("gsm") or 0, "colors": colors, "price_unit": unit, "price_per_unit": float(m0.get("unit_cost") or 0),
                                                    "price_per_meter": float(m0.get("unit_cost") or 0) if unit in ("m", "meter", "yard") else 0, "unit": unit, "linked_material_id": m0["id"],
                                                    "min_order_qty": 0, "test_results": "", "notes": f"Dibuat otomatis dari master kain {len(ms)} warna (konsolidasi 2026-09-23)", "status": "active",
                                                    "created_by": USER["id"], "created_by_name": USER["name"], "created_at": now(), "updated_at": now()})
        self.report["rnd"] = {"tech_packs_created": created, "rnd_materials_created": fam_created}

    def line_units(self):
        units = defaultdict(Counter)
        for lines in bom_from_revisi().values():
            for ln in lines:
                units[ln["code"]][ln["unit"]] += 1
        for gs in groups_from_rnd().values():
            for g in gs:
                for ln in g["lines"]:
                    units[ln["code"]][ln["unit"]] += 1
        for b in self.boms:
            for ln in b.get("materials") or []:
                units[ln["code"]][norm_unit(ln["unit"])] += 1
        return {c: u.most_common(1)[0][0] for c, u in units.items()}

    async def run(self):
        await self.load()
        price_plan, warn = plan_prices(self.mats, self.line_units())
        for p in price_plan:  # isi kemasan yang belum diisi owner → override (lihat PACK_OVERRIDE)
            ov = PACK_OVERRIDE.get(p["code"])
            if ov and p["pack_size"] == 1 and p["base"] == p["buy_unit"]:
                p.update({"base": ov[1], "pack_size": ov[0], "unit_cost": round(p["buy_price"] / ov[0], 4)})
                self.report["notes"].append(f"{p['code']}: isi kemasan {ov[0]:g} {ov[1]}/{p['buy_unit']} ({ov[2]})")
        self.report["warnings"] += warn
        big = [p for p in price_plan if p["unit_cost_before"] and (p["unit_cost"] / p["unit_cost_before"] > 5 or p["unit_cost"] / p["unit_cost_before"] < 0.2)]
        self.report["harga"] = {"diterapkan": len(price_plan), "ganti_satuan_dasar": sum(1 for p in price_plan if p["base"] != p["base_before"]),
                                "perubahan_besar": [f"{p['code']} {p['name'][:30]}: {p['unit_cost_before']:g}/{p['base_before']} → {p['unit_cost']:g}/{p['base']}" for p in big]}
        if self.apply:
            await apply_prices(self.db, price_plan)
            for p in price_plan:
                self.mats[p["code"]].update({"unit": p["base"], "base_uom": p["base"], "unit_cost": p["unit_cost"], "uoms": [{"code": p["base"], "factor": 1.0, "is_base": True}]
                                             + ([{"code": p["buy_unit"], "factor": p["pack_size"], "is_base": False}] if p["buy_unit"] != p["base"] else [])})
        else:
            for p in price_plan:
                self.mats[p["code"]].update({"unit": p["base"], "base_uom": p["base"], "unit_cost": p["unit_cost"], "uoms": [{"code": p["base"], "factor": 1.0, "is_base": True}]
                                             + ([{"code": p["buy_unit"], "factor": p["pack_size"], "is_base": False}] if p["buy_unit"] != p["base"] else [])})
        await self.ensure_tlk()
        new, skipped = plan_variants(self.models_by_code, self.variants, self.colors, self.sizes)
        self.report["varian"] = {"baru": len(new), "per_model": dict(Counter(n["model"]["code"] for n in new)), "dilewati": skipped}
        new_docs = await apply_variants(self.db, new) if self.apply else [
            {"id": "preview", "model_id": n["model"]["id"], "model_code": n["model"]["code"], "size_id": n["size"]["id"], "size_code": n["size"]["code"], "color_code": n["color"]["code"],
             "color_name": n["color"]["name"], "sku": n["sku"], "active": True} for n in new]
        target, source, active = await self.build_targets(new_docs)
        self.active_variants = active
        await self.apply_boms(target, source, active)
        if self.apply:
            from core.master_fill import recalc_standard_costs
            from core.bom_fill import flag_hpp_validation
            from core.master_sync import sync_product_masters
            self.report["hpp"] = await recalc_standard_costs(self.db, USER)
            self.report["hpp_flags"] = await flag_hpp_validation(self.db)
            self.report["sync"] = await sync_product_masters(self.db, USER)
            await self.load()
        await self.rnd_structure()
        return self.report


async def main():
    apply = "--apply" in sys.argv
    db = AsyncIOMotorClient(os.environ["MONGO_URL"])[os.environ["DB_NAME"]]
    rep = await Konsolidasi(db, apply).run()
    out = f"{DIR}/konsolidasi_{'apply' if apply else 'preview'}_{now():%Y%m%d_%H%M}.json"
    json.dump(rep, open(out, "w"), indent=1, ensure_ascii=False, default=str)
    slim = {k: v for k, v in rep.items() if k not in ("bom_per_sku", "panels_created", "fabrics_created")}
    slim["panels_created"] = len(rep.get("panels_created", []))
    slim["fabrics_created"] = rep.get("fabrics_created", [])
    print(json.dumps(slim, indent=1, ensure_ascii=False, default=str)[:12000])
    print("→", out)


if __name__ == "__main__":
    asyncio.run(main())
