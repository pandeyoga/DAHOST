#!/usr/bin/env python3
"""Autofix konservatif ke berkas BARU; sumber tidak pernah ditimpa, tidak mengakses DB."""
import argparse
from collections import defaultdict
import hashlib
import json
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'backend'))
sys.path.insert(0, str(ROOT / 'scripts'))
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from import_master_template import REF_COLS, phone, s, is_num, num, _norm_sheet
from master_template_spec import SHEETS
from master_autofix_rules import RULES, arithmetic, compound_qty, packaging, textile_kind, redact


def header(ws):
    return {s(c.value).lower().rstrip('*'): c.column for c in ws[1] if s(c.value)}


def records(ws):
    h = header(ws)
    for i in range(2, ws.max_row + 1):
        values = [c.value for c in ws[i]]
        first = s(values[0])
        if not any(s(v) for v in values) or first == '#' or (first.startswith('#') and first[1:2].isspace()):
            continue
        yield i, {k: ws.cell(i, j).value for k, j in h.items()}


class Autofix:
    def __init__(self, source):
        self.source = Path(source)
        self.wb = load_workbook(source, data_only=False)
        self.changes, self.issues = [], []

    def issue(self, ws, row, reason):
        item = [ws.title, row, reason]
        if item not in self.issues:
            self.issues.append(item)

    def log(self, ws, row, column, old, new, rule, reason):
        self.changes.append([rule, ws.title, row, column, redact(old), redact(new), reason])

    def set(self, ws, row, col, value, rule, reason):
        j = header(ws).get(col) if isinstance(col, str) else col
        if not j:
            return
        cell = ws.cell(row, j)
        if cell.value == value and type(cell.value) is type(value):
            return
        self.log(ws, row, col, cell.value, value, rule, reason)
        cell.value = value
        if isinstance(value, str):
            # Jangan mengubah teks klien yang diawali '=' menjadi formula baru.
            cell.data_type = 's'

    def prepare(self):
        for ws in self.wb:
            matches = [n for n in SHEETS if _norm_sheet(ws.title).startswith(n) or _norm_sheet(ws.title).endswith(n)]
            if len(matches) == 1 and ws.title != matches[0] and matches[0] not in self.wb.sheetnames:
                old = ws.title
                ws.title = matches[0]
                self.log(ws, 1, 'nama_sheet', old, ws.title, 'R01', 'Hiasan nama dihapus; pemetaan unik')
        for name, aliases in [('02_KARYAWAN', {'nik': 'kode_karyawan'}),
                              ('16_LIVEHOST', {'nik_karyawan': 'kode_karyawan'})]:
            if name not in self.wb:
                continue
            ws = self.wb[name]
            if list(header(ws)) == ['done'] and not list(records(ws)):
                self.log(ws, 1, 'header', 'DONE', 'kolom template; data tetap kosong', 'R05', 'DONE bukan data karyawan')
                for j, (k, *_) in enumerate(SHEETS[name]['kolom'], 1):
                    ws.cell(1, j, k)
                self.issue(ws, 1, 'Master karyawan belum diisi; lengkapi kode internal dan data payroll')
            for old, new in aliases.items():
                h = header(ws)
                if old in h and new not in h:
                    self.set(ws, 1, h[old], new, 'R05', 'Alias header lama; bukan konversi nilai identitas')
        if '10_BOM' in self.wb and 'kode_warna' not in header(self.wb['10_BOM']):
            ws = self.wb['10_BOM']
            self.set(ws, 1, ws.max_column + 1, 'kode_warna', 'R05', 'Kolom opsional; dibiarkan kosong, resep warna perlu klien')

    def clean_cells(self):
        for name in SHEETS:
            if name not in self.wb:
                continue
            ws = self.wb[name]
            for i, rec in list(records(ws)):
                for col, value in rec.items():
                    if isinstance(value, str) and not value.startswith('=') and value != value.strip():
                        self.set(ws, i, col, value.strip(), 'R03', 'Spasi tepi dihapus')
                    value = ws.cell(i, header(ws)[col]).value
                    if col in REF_COLS and s(value).startswith('#'):
                        self.set(ws, i, col, s(value).lstrip('#').strip(), 'R02', 'Pagar bukan bagian kode; hex tidak disentuh')
                    value = ws.cell(i, header(ws)[col]).value
                    if (col == 'kode_ukuran' or (name == '04_UKURAN' and col == 'kode')) and s(value).upper() == 'ALL SIZE':
                        self.set(ws, i, col, 'ALLSIZE', 'R04', 'Kode ukuran tanpa spasi')
                    if col == 'telepon' and phone(value) != s(value):
                        self.set(ws, i, col, phone(value), 'R03', 'Pulihkan 0 depan telepon numerik dan simpan teks')
                    if col == 'kode_akun_toko':
                        import re
                        accounts = set(self.index('13_AKUN_TOKO', 'kode_akun'))
                        parts = [v.strip() for v in s(value).split(',')]
                        normalized = [re.sub(r'^(SHP|TTK)\s+(\d+)$', lambda m: f'{m[1]}-{int(m[2]):02d}', p.upper()) for p in parts]
                        if normalized != parts and all(p in accounts for p in normalized):
                            self.set(ws, i, col, ', '.join(normalized), 'R03', 'Format akun cocok persis master toko')

    def index(self, name, key):
        out = defaultdict(list)
        if name in self.wb:
            for _, r in records(self.wb[name]):
                if s(r.get(key)):
                    out[s(r[key]).upper()].append(r)
        return out

    def references(self):
        colors = self.index('03_WARNA', 'kode')
        names = defaultdict(set)
        for code, rows in colors.items():
            for r in rows:
                names[s(r.get('nama')).upper()].add(code)
        models, sizes = self.index('08_MODEL', 'kode'), self.index('04_UKURAN', 'kode')
        if '09_BARANG_JADI' not in self.wb:
            return
        ws = self.wb['09_BARANG_JADI']
        for i, r in list(records(ws)):
            raw_color = s(r.get('kode_warna')).upper()
            if raw_color not in colors and len(names[raw_color]) == 1:
                self.set(ws, i, 'kode_warna', next(iter(names[raw_color])), 'R06', 'Nama warna cocok satu kode master')
            parts = s(r.get('sku')).upper().rsplit('-', 2)
            if len(parts) != 3:
                continue
            mk, ck, sk = parts
            for field, val, master in [('kode_model', mk, models), ('kode_warna', ck, colors), ('kode_ukuran', sk, sizes)]:
                if field not in header(ws):
                    self.issue(ws, i, f'Kolom {field} tidak tersedia; lengkapi struktur template')
                    continue
                current = s(ws.cell(i, header(ws)[field]).value).upper()
                if val in master and current not in master:
                    self.set(ws, i, field, val, 'R07', 'Segmen SKU cocok master; bukan membuat kode baru')
                elif val in master and current in master and current != val:
                    self.issue(ws, i, f'{field} berbeda dengan SKU; keduanya valid, perlu keputusan klien')
            if not s(r.get('nama')) and len(models.get(mk, [])) == 1:
                model_name = s(models[mk][0].get('nama'))
                color_name = s(colors[ck][0].get('nama')) if len(colors.get(ck, [])) == 1 else ck
                self.set(ws, i, 'nama', f'{model_name} [{color_name} · {sk}]', 'R07', 'Nama tampil dari model dan segmen SKU')

    def numeric(self):
        numeric_cols = {'harga_per_satuan', 'harga_jual', 'harga_coret', 'harga_jual_dasar', 'qty_per_pcs',
                        'stok_minimum', 'isi_per_kemasan', 'gramasi_gsm', 'lebar_cm', 'berat_gram'}
        for name in SHEETS:
            if name not in self.wb:
                continue
            ws = self.wb[name]
            for i, r in list(records(ws)):
                if name == '06_MATERIAL_KAIN':
                    kind = textile_kind(r.get('nama'), r.get('jenis'))
                    if kind is not None:
                        self.set(ws, i, 'jenis', kind, 'R08', 'Istilah tekstil eksplisit pada jenis/nama')
                for col in numeric_cols.intersection(r):
                    value = r[col]
                    if isinstance(value, str) and value.startswith('='):
                        result = arithmetic(value)
                        if result is not None:
                            self.set(ws, i, col, result, 'R10', 'Evaluasi aritmetika literal saja, tanpa referensi sel')
                        else:
                            self.issue(ws, i, f'{col}: formula bukan aritmetika literal; isi nilai final terverifikasi')
                if name == '07_AKSESORIS' and not is_num(r.get('isi_per_kemasan')) and s(r.get('isi_per_kemasan')):
                    val = packaging(r.get('isi_per_kemasan'), r.get('satuan_dasar'))
                    if val is not None and val > 0:
                        self.set(ws, i, 'isi_per_kemasan', val, 'R09', 'Jumlah kemasan dalam satuan dasar yang sama/sedimensi')
                    else:
                        self.issue(ws, i, 'Isi kemasan ambigu/berbeda dimensi; pilih jumlah dan satuan stok yang benar')
                if name == '10_BOM' and not is_num(r.get('qty_per_pcs')):
                    val = compound_qty(r.get('qty_per_pcs'), r.get('kode_material'), r.get('satuan'))
                    if val is not None:
                        self.set(ws, i, 'qty_per_pcs', val, 'R10', 'Jumlah daftar komponen positif pada satu bahan+satuan; bukan rentang dua angka')
                    elif s(r.get('qty_per_pcs')):
                        self.issue(ws, i, 'Qty majemuk/rentang atau satuan campuran ambigu; konfirmasi komponen per pcs')
                if name == '14_KATALOG_JUAL':
                    h = header(ws)
                    if not {'harga_jual', 'harga_coret'} <= set(h):
                        self.issue(ws, i, 'Kolom harga_jual/harga_coret tidak lengkap; lengkapi struktur template')
                        continue
                    jual, coret = (ws.cell(i, h[k]).value for k in ('harga_jual', 'harga_coret'))
                    if is_num(jual) and is_num(coret) and 0 < num(coret) < num(jual):
                        self.set(ws, i, 'harga_jual', num(coret), 'R11', 'Harga tayang lebih rendah daripada harga sebelum diskon')
                        self.set(ws, i, 'harga_coret', num(jual), 'R11', 'Harga sebelum diskon lebih besar')

    def catalog(self):
        if '14_KATALOG_JUAL' not in self.wb:
            return
        ws = self.wb['14_KATALOG_JUAL']
        accounts = set(self.index('13_AKUN_TOKO', 'kode_akun'))
        blanks = [c.column for c in ws[1] if not s(c.value)]
        for j in blanks:
            values = {s(ws.cell(i, j).value).upper() for i, _ in records(ws) if s(ws.cell(i, j).value)}
            if values and values <= accounts and 'kode_akun_kedua' not in header(ws):
                self.set(ws, 1, j, 'kode_akun_kedua', 'R12', 'Kolom tanpa judul cocok unik master akun kedua')
        urls = [c.column for c in ws[1] if s(c.value) == 'tautan_produk']
        if len(urls) == 2:
            self.set(ws, 1, urls[1], 'tautan_produk_kedua', 'R12', 'Pertahankan kedua kolom tautan; tidak saling menimpa')
        h = header(ws)
        if 'kode_akun_kedua' not in h:
            return
        if not {'kode_akun', 'tautan_produk'} <= set(h):
            self.issue(ws, 1, 'Kolom kode_akun/tautan_produk tidak lengkap; pemisahan katalog ditahan')
            return
        for i, r in list(records(ws)):
            second = s(r.get('kode_akun_kedua')).upper()
            if not second:
                continue
            if second not in accounts or second == s(r.get('kode_akun')).upper():
                self.issue(ws, i, 'Akun kedua tidak dikenal/sama dengan akun pertama; tidak dipecah otomatis')
                continue
            vals = [ws.cell(i, j).value for j in range(1, ws.max_column + 1)]
            vals[h['kode_akun'] - 1] = second
            vals[h['kode_akun_kedua'] - 1] = None
            vals[h['tautan_produk'] - 1] = r.get('tautan_produk_kedua')
            if 'tautan_produk_kedua' in h:
                vals[h['tautan_produk_kedua'] - 1] = None
            ws.append(vals)
            new_row = ws.max_row
            self.log(ws, i, 'baris', second, f'baris {new_row}', 'R12', 'Satu baris per akun+SKU; tautan pertama tidak disalin ke toko kedua')
            self.set(ws, i, 'kode_akun_kedua', None, 'R12', 'Akun sudah dipindah ke baris tersendiri')
            if 'tautan_produk_kedua' in h:
                self.set(ws, i, 'tautan_produk_kedua', None, 'R12', 'Tautan kedua ikut akun kedua')

    def deduplicate(self):
        for name in ['03_WARNA', '04_UKURAN', '06_MATERIAL_KAIN', '07_AKSESORIS', '08_MODEL', '09_BARANG_JADI']:
            if name not in self.wb:
                continue
            ws, seen = self.wb[name], {}
            key = 'sku' if name == '09_BARANG_JADI' else 'kode'
            for i, r in list(records(ws)):
                code = s(r.get(key)).upper()
                if not code:
                    continue
                values = [s(ws.cell(i, j).value) for j in range(1, ws.max_column + 1)]
                if code not in seen:
                    seen[code] = (i, values)
                elif values == seen[code][1]:
                    self.log(ws, i, 'seluruh_baris', json.dumps(values, ensure_ascii=False), '', 'R13', f'Identik dengan baris {seen[code][0]}; nomor baris tidak digeser')
                    for c in ws[i]:
                        c.value = None
                else:
                    self.issue(ws, i, f'{key} {code} duplikat berbeda isi dengan baris {seen[code][0]}; tidak digabung')

    def save(self, dest):
        dest = Path(dest)
        if self.source.resolve() == dest.resolve():
            raise ValueError('Tujuan harus berbeda dari sumber; sumber tidak boleh ditimpa')
        if dest.exists():
            raise FileExistsError(f'Tujuan sudah ada: {dest}; gunakan nama baru')
        tables = [('LOG_PERUBAHAN', ['Aturan', 'Sheet', 'Baris sumber', 'Kolom', 'Sebelum', 'Sesudah', 'Alasan'], self.changes),
                  ('BUTUH_KLIEN', ['Sheet', 'Baris', 'Pertanyaan/konflik'], self.issues)]
        for name, titles, rows in tables:
            ws = self.wb[name] if name in self.wb else self.wb.create_sheet(name)
            existing = {tuple(s(v) for v in r) for r in ws.values}
            if ws.max_row == 1 and ws.cell(1, 1).value is None:
                ws.append(titles) if ws.max_row > 1 else None
                for j, title in enumerate(titles, 1):
                    ws.cell(1, j, title)
            for r in rows:
                if tuple(s(v) for v in r) not in existing:
                    ws.append(r)
                    for cell in ws[ws.max_row]:
                        if isinstance(cell.value, str):
                            cell.data_type = 's'
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions
            for c in ws[1]:
                c.font = Font(bold=True, color='FFFFFF')
                c.fill = PatternFill('solid', fgColor='1F3A5F')
            for row in ws:
                for c in row:
                    c.alignment = Alignment(vertical='top', wrap_text=True)
                    ws.column_dimensions[c.column_letter].width = 26 if c.column < 4 else 55
        if 'META_AUTOFIX' not in self.wb:
            meta = self.wb.create_sheet('META_AUTOFIX')
            meta.append(['Sumber', self.source.name])
            meta.append(['SHA256 sumber', hashlib.sha256(self.source.read_bytes()).hexdigest()])
            for k, v in RULES.items():
                meta.append([k, v])
        dest.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(dest)
        self.wb.close()
        return {'perubahan': len(self.changes), 'butuh_klien': len(self.issues), 'output': str(dest)}


def build(source, dest):
    fix = Autofix(source)
    try:
        fix.prepare()
        fix.clean_cells()
        fix.references()
        fix.numeric()
        fix.catalog()
        fix.deduplicate()
        return fix.save(dest)
    finally:
        fix.wb.close()


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source', type=Path)
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    print(json.dumps(build(args.source, args.output), ensure_ascii=False, indent=2))