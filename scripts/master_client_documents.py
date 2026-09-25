"""Artefak review untuk manusia; workbook jawaban tidak pernah diimpor otomatis."""
import json
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from master_autofix_rules import RULES, redact
from master_client_questions import evidence_examples


def write_json(path, content):
    path.write_text(json.dumps(content, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def write_questions(dest, summary, questions):
    lines = ['# Pertanyaan untuk klien — master DAHOST', '',
             f"Sumber: `{redact(summary['source'])}` · SHA256 `{summary['source_sha256']}`.",
             f"Hasil: **{summary['errors_after']} pesan kesalahan**, **{summary['warnings_after']} pesan peringatan**.",
             'Baris mengacu ke **MASTER_CLIENT_AUTOFIX.xlsx**. Nomor baris tidak ditebak dari paket lama.',
             'Data belum diimpor. Gunakan kode karyawan internal; jangan kirim NIK KTP.',
             'Jawaban dapat diisi di **JAWABAN_KLIEN.xlsx**. Seluruh temuan tersedia pada sheet TEMUAN.', '']
    for q in questions:
        lines += [f"## {q['id'][1:]}. {redact(q['judul'])}", redact(q['pertanyaan']), '',
                  f"Status: **{q['status']}** · {len(q['temuan'])} pesan terkait."]
        for item in evidence_examples(q['temuan']):
            lines.append(f"- `{item['sheet']}` baris **{item['baris']}** ({item['jenis']}): {redact(item['pesan'])}")
        if not q['temuan']:
            lines.append('Tidak ada temuan validator pada kelompok ini; ini bukan bukti kelengkapan bisnis/HPP.')
        lines += ['', '**Jawaban klien:** …', '**Penanggung jawab / tanggal:** …', '']
    lines += ['## Setelah jawaban diterima',
              'Perbaiki salinan master → dry-run 0 kesalahan → tinjau seluruh peringatan BOM/HPP → '
              'persetujuan owner → backup → Fase 7 terpisah.',
              '**Fase 7 DITAHAN. Tidak menjalankan --apply, termasuk bila validator sudah 0 kesalahan.**', '']
    (dest / 'PERTANYAAN_KLIEN.md').write_text('\n'.join(lines), encoding='utf-8')


def write_answers(dest, questions, catalog):
    wb = Workbook()
    ws = wb.active
    ws.title = 'PERTANYAAN'
    ws.append(['ID', 'Topik', 'Pertanyaan', 'Jumlah pesan', 'Status review',
               'Jawaban klien', 'Kode/nilai final & baris rujukan', 'Penanggung jawab', 'Tanggal'])
    for q in questions:
        ws.append([q['id'], q['judul'], q['pertanyaan'], len(q['temuan']), q['status'], '', '', '', ''])
    validation = DataValidation(type='list', formula1='"PERLU JAWABAN,TINJAU BISNIS,DIJAWAB,PERLU KLARIFIKASI"')
    ws.add_data_validation(validation)
    validation.add('E2:E7')
    evidence = wb.create_sheet('TEMUAN')
    evidence.append(['ID pertanyaan', 'Jenis', 'Sheet master', 'Baris master', 'Temuan',
                     'Keputusan/koreksi klien', 'Penanggung jawab'])
    for q in questions:
        for f in q['temuan']:
            evidence.append([q['id'], f['jenis'], f['sheet'], f['baris'], f['pesan'], '', ''])
    conflicts = wb.create_sheet('KATALOG_DUPLIKAT')
    conflicts.append(['Kode akun', 'SKU', 'Baris master', 'Kolom berbeda', 'Harga jual',
                      'Harga coret', 'Tautan produk', 'Aktif', 'Nama tampil', 'Keputusan klien'])
    for d in catalog['duplicate_groups']:
        for row, values in zip(d['baris'], d['nilai']):
            conflicts.append([d['kode_akun'], d['sku'], row,
                              ', '.join(d['kolom_berbeda']) or 'identik',
                              *[values[k] for k in ('harga_jual', 'harga_coret', 'tautan_produk', 'aktif', 'nama_tampil')], ''])
    note = wb.create_sheet('CATATAN')
    note.append(['Catatan'])
    for text in ['Fase 7 DITAHAN; jawaban tidak langsung menyimpan ke database.',
                 'Baris master mengacu MASTER_CLIENT_AUTOFIX.xlsx; LOG_PERUBAHAN mencatat asal perubahan.',
                 'Satu baris bisa memiliki beberapa pesan. Jumlah pesan bukan jumlah baris unik.',
                 'Simpan secara privat; workbook dapat berisi data kontak klien.',
                 'Status DIJAWAB bukan persetujuan apply. Koreksi master harus melalui dry-run baru.']:
        note.append([text])
    for sheet in wb:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.value = redact(cell.value)
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                sheet.column_dimensions[cell.column_letter].width = 25
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='1F3A5F')
    ws.column_dimensions['C'].width = 85
    ws.column_dimensions['F'].width = 60
    ws.column_dimensions['G'].width = 55
    evidence.column_dimensions['E'].width = 100
    evidence.column_dimensions['F'].width = 60
    conflicts.column_dimensions['G'].width = 70
    note.column_dimensions['A'].width = 110
    for row in range(2, 8):
        ws.row_dimensions[row].height = 150
    wb.save(dest / 'JAWABAN_KLIEN.xlsx')
    wb.close()


def write_readme(dest, summary):
    history = (f"Angka historis yang diberikan: {summary['historical_errors']}. "
               if summary['historical_errors'] is not None else '')
    lines = ['# Paket review master klien — Fase 1–6', '',
             f"Sumber: `{summary['source']}` · SHA256 `{summary['source_sha256']}`.",
             history + f"Validator saat ini: {summary['errors_before_current_validator']} pesan kesalahan sebelum "
             f"→ {summary['errors_after']} sesudah; {summary['warnings_after']} peringatan sesudah.",
             f"Autofix {summary['changes']} perubahan; putaran kedua {summary['changes_second_run']} perubahan. "
             f"Nilai dan tipe sel master identik pada putaran kedua: {summary['data_idempotent']}.",
             'Satu baris dapat memiliki beberapa pesan. Angka historis tidak disamakan dengan hasil validator saat ini.',
             'Hasil bergantung pada master database yang tersedia saat dry-run; backup database lama tidak dipulihkan.',
             '', '## Isi paket',
             '- MASTER_CLIENT_AUTOFIX.xlsx: salinan koreksi dengan LOG_PERUBAHAN, BUTUH_KLIEN, META_AUTOFIX.',
             '- REVIEW_SEBELUM_AUTOFIX.xlsx / REVIEW_SETELAH_AUTOFIX.xlsx: temuan per sheet/baris.',
             '- PERTANYAAN_KLIEN.md: enam kelompok keputusan dengan bukti aktual.',
             '- JAWABAN_KLIEN.xlsx: enam pertanyaan, seluruh temuan, dan semua grup duplikat katalog.',
             '- HASIL.json / KESALAHAN.json / PERTANYAAN_KLIEN.json / KATALOG_DUPLIKAT.json: data review terstruktur.',
             '- MASTER_CLIENT_AUTOFIX_ULANG.xlsx: bukti idempotensi nilai/tipe sel data, bukan kesamaan byte ZIP.',
             '- MANIFEST.json: checksum semua artefak di atas; tidak mencakup dirinya sendiri.',
             '', '## Pengaman',
             '**Fase 7 DITAHAN. Tidak ada apply data klien. HPP/margin/neraca produksi belum diverifikasi.**',
             'Database dibungkus hanya-baca pada dry-run. Folder lama tidak ditimpa. Sumber tidak diubah.',
             'HASIL.json memuat jumlah dan hash isi 15 koleksi master sebelum/sesudah; paket dibatalkan bila kondisinya berubah.',
             'Jawaban klien bukan instruksi otomatis: koreksi master, validasi ulang, review peringatan, persetujuan owner dan backup tetap wajib.',
             '**Privasi:** workbook masih dapat mengandung kontak klien. Jangan unggah paket ke unduhan publik atau repo publik.',
             'BUTUH_KLIEN/LOG_PERUBAHAN memuat jejak historis autofix; TEMUAN dan laporan sesudah memuat hasil validator saat ini.',
             '', '## Aturan konservatif']
    lines += [f'- **{k}:** {v}' for k, v in RULES.items()]
    (dest / 'README.md').write_text('\n'.join(lines) + '\n', encoding='utf-8')