"""Iter226: verify scripts/export_sku_bom.py output vs MongoDB & vs laporan_kekosongan_bom.

Read-only tests (no DB writes).
"""
from __future__ import annotations

import io
import os
import subprocess
from collections import Counter, defaultdict
from pathlib import Path

import pytest
import requests
from openpyxl import load_workbook
from pymongo import MongoClient

BACKEND_DIR = Path("/app/backend")
SCRIPT = Path("/app/scripts/export_sku_bom.py")
OUT_XLSX = Path("/tmp/test_export_sku_bom_iter226.xlsx")
PUBLIC_XLSX = "/downloads/EXPORT_SKU_BOM_DA.xlsx"


def _env():
    e = os.environ.copy()
    for line in (BACKEND_DIR / ".env").read_text().splitlines():
        if "=" in line and not line.strip().startswith("#"):
            k, v = line.split("=", 1)
            e[k.strip()] = v.strip().strip('"').strip("'")
    return e


ENV = _env()


@pytest.fixture(scope="module")
def db():
    return MongoClient(ENV["MONGO_URL"])[ENV["DB_NAME"]]


@pytest.fixture(scope="module")
def wb():
    r = subprocess.run(
        ["python", str(SCRIPT), "--out", str(OUT_XLSX)],
        cwd=BACKEND_DIR, env=ENV, capture_output=True, text=True, timeout=180,
    )
    assert r.returncode == 0, f"script failed:\nSTDOUT={r.stdout}\nSTDERR={r.stderr}"
    assert OUT_XLSX.exists()
    return load_workbook(OUT_XLSX, read_only=False, data_only=False)


@pytest.fixture(scope="module")
def sheets(wb):
    def rows(name):
        ws = wb[name]
        hdr = [c.value for c in ws[1]]
        out = []
        for r in ws.iter_rows(min_row=2, values_only=True):
            out.append(dict(zip(hdr, r)))
        return hdr, out
    return {n: rows(n) for n in ("DETAIL_BOM", "RINGKASAN_SKU", "RINGKASAN_MODEL", "KETERANGAN")}


# 1. Sheets present
def test_sheets_present(wb):
    assert set(wb.sheetnames) >= {"DETAIL_BOM", "RINGKASAN_SKU", "RINGKASAN_MODEL", "KETERANGAN"}


# 2. Row counts match expected & DB
def test_row_counts_match_db(db, sheets):
    n_models = db.rahaza_models.count_documents({})
    n_variants = db.rahaza_model_variants.count_documents({})
    model_ids_with_variants = set(db.rahaza_model_variants.distinct("model_id"))
    n_models_no_variant = db.rahaza_models.count_documents({"id": {"$nin": list(model_ids_with_variants)}})
    _, model_rows = sheets["RINGKASAN_MODEL"]
    _, sku_rows = sheets["RINGKASAN_SKU"]
    _, detail_rows = sheets["DETAIL_BOM"]
    assert len(model_rows) == n_models == 104
    assert len(sku_rows) == n_variants + n_models_no_variant
    assert len(sku_rows) == 665
    assert n_variants == 645
    assert n_models_no_variant == 20
    assert len(detail_rows) == 2174


# 3. Every variant SKU appears once in RINGKASAN_SKU and >=1 in DETAIL_BOM
def test_each_sku_present(db, sheets):
    _, sku_rows = sheets["RINGKASAN_SKU"]
    _, detail_rows = sheets["DETAIL_BOM"]
    db_skus = {v["sku"] for v in db.rahaza_model_variants.find({}, {"_id": 0, "sku": 1})}
    sku_counter = Counter(r["sku"] for r in sku_rows if r.get("sku"))
    detail_skus = {r["sku"] for r in detail_rows if r.get("sku")}
    missing_sum = db_skus - set(sku_counter)
    dupes = {k: v for k, v in sku_counter.items() if v > 1}
    missing_det = db_skus - detail_skus
    assert not missing_sum, f"missing in RINGKASAN_SKU: {list(missing_sum)[:5]}"
    assert not dupes, f"duplicated in RINGKASAN_SKU: {list(dupes.items())[:5]}"
    assert not missing_det, f"missing in DETAIL_BOM: {list(missing_det)[:5]}"


# 4. Per-SKU status assertions
def test_specific_sku_statuses(db, sheets):
    _, sku_rows = sheets["RINGKASAN_SKU"]
    by_sku = {r["sku"]: r for r in sku_rows if r.get("sku")}

    # DA-2201-* (Ona) all TANPA_BOM
    ona = [s for s in by_sku if s.startswith("DA-2201-")]
    assert ona, "no DA-2201-* SKUs found"
    for s in ona:
        assert by_sku[s]["status_bom"] == "TANPA_BOM", f"{s} status={by_sku[s]['status_bom']}"

    # DA-1101-BRG-ALLSIZE = BOM_TANPA_AKSESORIS
    assert by_sku["DA-1101-BRG-ALLSIZE"]["status_bom"] == "BOM_TANPA_AKSESORIS"

    # DA-1101-CRL-ALLSIZE = BOM_LENGKAP with baris_aksesoris == 4
    r = by_sku["DA-1101-CRL-ALLSIZE"]
    assert r["status_bom"] == "BOM_LENGKAP"
    assert r["baris_aksesoris"] == 4, f"baris_aksesoris={r['baris_aksesoris']}"

    # 30 SKU nonaktif → BOM_NONAKTIF
    inactive = [r for r in sku_rows if r.get("sku_aktif") == "TIDAK"]
    assert len(inactive) == 30, f"got {len(inactive)}"
    assert all(r["status_bom"] == "BOM_NONAKTIF" for r in inactive)


# 5. 20 model tanpa varian: TANPA_SKU, kolom sku kosong
def test_no_variant_models(db, sheets):
    _, sku_rows = sheets["RINGKASAN_SKU"]
    no_sku_rows = [r for r in sku_rows if r["status_bom"] == "TANPA_SKU"]
    assert len(no_sku_rows) == 20
    assert all((r.get("sku") in (None, "")) for r in no_sku_rows)


# 6. Sample 5 BOM_LENGKAP DETAIL_BOM rows: numbers match rahaza_boms.materials & master unit_cost
def test_detail_matches_db_sample(db, sheets):
    _, detail_rows = sheets["DETAIL_BOM"]
    complete = [r for r in detail_rows if r["status_bom"] == "BOM_LENGKAP" and r.get("kode_material")]
    assert len(complete) >= 5
    # Deterministic sample: first 5 by (sku, no_baris)
    sample = sorted(complete, key=lambda r: (r["sku"], r["no_baris"]))[:5]
    materials_master = {m["code"]: m for m in db.rahaza_materials.find({}, {"_id": 0})}
    variants = {v["sku"]: v for v in db.rahaza_model_variants.find({}, {"_id": 0})}
    boms_by_model = defaultdict(list)
    for b in db.rahaza_boms.find({}, {"_id": 0}):
        boms_by_model[b["model_id"]].append(b)
    for row in sample:
        v = variants[row["sku"]]
        cands = [b for b in boms_by_model[v["model_id"]]
                 if (b.get("color_code") or "").upper() == (v.get("color_code") or "").upper()
                 and b.get("size_id") == v.get("size_id")]
        actives = [b for b in cands if b.get("active") is not False]
        bom = max(actives, key=lambda b: b.get("version") or 0)
        ln = bom["materials"][row["no_baris"] - 1]
        assert row["kode_material"] == ln.get("code")
        assert row["qty"] == pytest.approx(float(ln.get("qty")))
        assert row["satuan"] == ln.get("unit")
        assert row["qty_dasar"] == pytest.approx(float(ln.get("qty_base")))
        assert row["harga_per_satuan_dasar"] == pytest.approx(float(ln.get("unit_cost_base")))
        expected_biaya = round(float(ln["qty_base"]) * float(ln["unit_cost_base"]), 2)
        assert row["biaya_baris"] == pytest.approx(expected_biaya)
        m = materials_master.get(ln.get("code"))
        if m and m.get("unit_cost") is not None:
            assert row["harga_master_saat_ini"] == pytest.approx(float(m["unit_cost"]))


# 7. Model status counts match laporan_kekosongan_bom.py expectations
def test_model_status_counts(sheets):
    _, model_rows = sheets["RINGKASAN_MODEL"]
    c = Counter(r["status_kelengkapan"] for r in model_rows)
    expected = {"A_TANPA_SKU": 20, "B_ADA_SKU_NOL_BOM": 1, "C_SEBAGIAN_VARIAN_TANPA_BOM": 2,
                "D2_SEBAGIAN_BOM_TANPA_AKSESORIS": 1, "D_BOM_TANPA_AKSESORIS": 18,
                "E_LENGKAP": 56, "F_DIHENTIKAN": 6}
    for k, v in expected.items():
        assert c.get(k, 0) == v, f"{k}: got {c.get(k, 0)} expected {v}"
    assert sum(c.values()) == 104


# 8. Cross-check: run laporan_kekosongan_bom.py and reconcile status distribution
def test_matches_laporan_kekosongan(sheets):
    r = subprocess.run(
        ["python", "/app/scripts/laporan_kekosongan_bom.py"],
        cwd=BACKEND_DIR, env=ENV, capture_output=True, text=True, timeout=60,
    )
    assert r.returncode == 0, r.stderr
    _, model_rows = sheets["RINGKASAN_MODEL"]
    c = Counter(r["status_kelengkapan"] for r in model_rows)
    # laporan only counts active models (excludes MODEL_NONAKTIF)
    active_c = {k: v for k, v in c.items() if k != "MODEL_NONAKTIF"}
    for k in ("A_TANPA_SKU", "B_ADA_SKU_NOL_BOM", "C_SEBAGIAN_VARIAN_TANPA_BOM",
              "D_BOM_TANPA_AKSESORIS", "D2_SEBAGIAN_BOM_TANPA_AKSESORIS",
              "E_LENGKAP", "F_DIHENTIKAN"):
        line = f"{active_c.get(k, 0):3d}  {k}"
        assert line in r.stdout, f"missing line: {line!r}\n{r.stdout}"


# 9. Public download at REACT_APP_BACKEND_URL/downloads/EXPORT_SKU_BOM_DA.xlsx
def test_public_download():
    fe_env = {k.split("=", 1)[0].strip(): k.split("=", 1)[1].strip()
              for k in Path("/app/frontend/.env").read_text().splitlines() if "=" in k}
    base = fe_env["REACT_APP_BACKEND_URL"].rstrip("/")
    resp = requests.get(base + PUBLIC_XLSX, timeout=60)
    assert resp.status_code == 200, f"status={resp.status_code}"
    assert len(resp.content) > 10_000
    wb2 = load_workbook(io.BytesIO(resp.content), read_only=True)
    assert set(wb2.sheetnames) >= {"DETAIL_BOM", "RINGKASAN_SKU", "RINGKASAN_MODEL", "KETERANGAN"}
