"""Iteration 225 — REVIEW_HARGA_MATERIAL_DA.xlsx: GET /harga-review + roundtrip via fill-preview.
NOTE: DILARANG memanggil fill-apply / seed / endpoint tulis lain (data klien nyata).
"""
from __future__ import annotations

import io
import os
import re

import openpyxl
import pytest
import requests
from motor.motor_asyncio import AsyncIOMotorClient

BASE = "http://localhost:8001"
LOGIN = {"email": "admin@garment.com", "password": "Admin@123"}
MONGO_URL = os.environ.get("MONGO_URL", "mongodb://localhost:27017")
DB_NAME = os.environ.get("DB_NAME", "test_database")

RED = "F8CBAD"
YELLOW = "FFF2CC"

MAT_COLS = ["kode", "nama", "tipe", "kategori", "satuan_dasar", "satuan_beli",
            "isi_per_satuan_beli", "harga_per_satuan_beli",
            "harga_per_satuan_dasar_sekarang", "min_stok", "keterangan"]


# ── shared fixtures ──────────────────────────────────────────────────────────
@pytest.fixture(scope="module")
def token() -> str:
    r = requests.post(f"{BASE}/api/auth/login", json=LOGIN, timeout=30)
    assert r.status_code == 200, f"login failed: {r.status_code} {r.text[:200]}"
    tok = r.json().get("access_token") or r.json().get("token")
    assert tok
    return tok


@pytest.fixture(scope="module")
def auth_headers(token) -> dict:
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture(scope="module")
def review_response(auth_headers):
    r = requests.get(f"{BASE}/api/rahaza/master/harga-review", headers=auth_headers, timeout=120)
    assert r.status_code == 200, f"harga-review status={r.status_code} body={r.text[:300]}"
    return r


@pytest.fixture(scope="module")
def wb(review_response):
    return openpyxl.load_workbook(io.BytesIO(review_response.content), data_only=True)


# ── 1. endpoint basics ───────────────────────────────────────────────────────
def test_harga_review_no_token():
    r = requests.get(f"{BASE}/api/rahaza/master/harga-review", timeout=30)
    assert r.status_code in (401, 403), f"unauth got {r.status_code}"


def test_harga_review_headers(review_response):
    ct = review_response.headers.get("content-type", "")
    assert "spreadsheetml" in ct or "xlsx" in ct, f"ct={ct}"
    cd = review_response.headers.get("content-disposition", "")
    assert "REVIEW_HARGA_MATERIAL_DA.xlsx" in cd, f"cd={cd}"


def test_sheet_order(wb):
    assert wb.sheetnames == ["PETUNJUK", "MATERIAL", "HPP_MODEL", "BOM_BARIS_MAHAL", "KAIN_PER_MODEL"], wb.sheetnames


# ── 2. MATERIAL sheet ────────────────────────────────────────────────────────
def test_material_headers(wb):
    ws = wb["MATERIAL"]
    first11 = [ws.cell(row=1, column=i + 1).value for i in range(11)]
    assert first11 == MAT_COLS, first11
    assert ws.cell(row=1, column=19).value == "tingkat"
    assert ws.cell(row=1, column=20).value == "kenapa_dicurigai"


def _material_rows(ws):
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=False):
        if not r[0].value:
            continue
        rows.append(r)
    return rows


def test_material_rows_valid_and_price_empty(wb):
    ws = wb["MATERIAL"]
    rows = _material_rows(ws)
    for r in rows:
        code = r[0].value
        tipe = r[2].value
        assert tipe in ("accessory", "fabric"), f"bad tipe {tipe} @ {code}"
        assert not str(code).startswith("CUT-"), f"CUT- leaked: {code}"
        assert r[7].value in (None, "", 0), f"harga_per_satuan_beli should be empty at {code}: {r[7].value}"


def test_material_count_matches_mongo(wb):
    import asyncio

    async def _count():
        client = AsyncIOMotorClient(MONGO_URL)
        try:
            db = client[DB_NAME]
            q = {"type": {"$in": ["accessory", "fabric"]},
                 "active": {"$ne": False},
                 "code": {"$not": re.compile(r"^CUT-")}}
            return await db.rahaza_materials.count_documents(q)
        finally:
            client.close()

    n_db = asyncio.get_event_loop().run_until_complete(_count()) if False else asyncio.new_event_loop().run_until_complete(_count())
    n_sheet = len(_material_rows(wb["MATERIAL"]))
    assert n_sheet == n_db, f"MATERIAL rows {n_sheet} != mongo {n_db}"


def _row_level(row):
    return row[18].value or ""  # column 19


def _row_fill(row):
    fg = row[0].fill.fgColor
    return (fg.rgb or "")[-6:].upper() if fg else ""


def test_material_ordering_red_then_yellow(wb):
    rows = _material_rows(wb["MATERIAL"])
    seen_yellow = False
    seen_plain = False
    reds, yellows = 0, 0
    for r in rows:
        lvl = _row_level(r)
        if lvl == "MERAH":
            assert not seen_yellow and not seen_plain, f"MERAH after non-MERAH at {r[0].value}"
            reds += 1
        elif lvl == "KUNING":
            assert not seen_plain, f"KUNING after plain at {r[0].value}"
            seen_yellow = True
            yellows += 1
        else:
            seen_plain = True
    assert reds >= 4, f"expected ≥4 MERAH, got {reds}"


def test_material_fill_colors(wb):
    for r in _material_rows(wb["MATERIAL"]):
        lvl = _row_level(r)
        if lvl == "MERAH":
            for c in r:
                fg = (c.fill.fgColor.rgb or "")[-6:].upper() if c.fill.fgColor else ""
                # harga_per_satuan_beli col is yellow always; harga_per_satuan_dasar & pcs col are blue
                if c.column in (8, 9, 12):
                    continue
                assert fg == RED, f"MERAH row {r[0].value} col {c.column} fill={fg}"
        elif lvl == "KUNING":
            fg0 = (r[0].fill.fgColor.rgb or "")[-6:].upper()
            assert fg0 == YELLOW, f"KUNING row {r[0].value} first-cell fill={fg0}"
        # kenapa_dicurigai
        why = r[19].value
        if lvl:
            assert why, f"missing reason for {lvl} {r[0].value}"
        else:
            assert not why, f"reason present w/o level: {r[0].value}"


def test_material_expected_red_items(wb):
    expected = {"A-REN-0004", "A-KRT-0007", "A-BIS-0001", "A-BIS-0002"}
    red_codes = {r[0].value for r in _material_rows(wb["MATERIAL"]) if _row_level(r) == "MERAH"}
    missing = expected - red_codes
    assert not missing, f"missing red codes: {missing}"

    # verify A-REN-0004 specifics
    for r in _material_rows(wb["MATERIAL"]):
        if r[0].value == "A-REN-0004":
            # kolom I = harga_per_satuan_dasar_sekarang (index 8)
            assert abs(float(r[8].value) - 328900) < 1, f"A-REN-0004 price={r[8].value}"
            # model_biaya_maks kolom 17 (index 16)
            assert r[16].value == "DA-4104", f"A-REN-0004 model={r[16].value}"
            # porsi_hpp_maks_% kolom 18 (index 17)
            assert float(r[17].value) >= 90, f"A-REN-0004 porsi={r[17].value}"
        if r[0].value == "A-KRT-0007":
            assert abs(float(r[8].value) - 33000) < 1, r[8].value
        if r[0].value in ("A-BIS-0001", "A-BIS-0002"):
            assert abs(float(r[8].value) - 15400) < 1, f"{r[0].value} price={r[8].value}"


# ── 3. HPP_MODEL sheet ───────────────────────────────────────────────────────
def test_hpp_model_sheet(wb):
    ws = wb["HPP_MODEL"]
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
    codes = [r[0] for r in rows]
    assert len(codes) == len(set(codes)), "duplicate model codes in HPP_MODEL"
    # sorted desc by hpp_per_pcs_bom_ini (col 4 idx 3)
    hpp_list = [float(r[3] or 0) for r in rows]
    assert hpp_list == sorted(hpp_list, reverse=True), "HPP_MODEL not sorted desc"

    top = rows[0]
    assert top[0] == "DA-4104", f"top model={top[0]}"
    assert "Ochi" in (top[1] or ""), f"top name={top[1]}"
    assert abs(float(top[3]) - 515151) < 500, f"hpp={top[3]}"
    assert abs(float(top[5]) - 506506) < 500, f"biaya_aks={top[5]}"
    assert abs(float(top[6]) - 98.3) < 0.5, f"porsi_aks={top[6]}"
    assert top[10] == "MERAH", f"tingkat={top[10]}"
    assert "A-REN-0004" in (top[8] or ""), f"baris_termahal={top[8]}"

    # hpp == biaya_kain + biaya_aksesoris ±1
    for r in rows:
        hpp, fab, acc = float(r[3] or 0), float(r[4] or 0), float(r[5] or 0)
        assert abs(hpp - (fab + acc)) <= 1, f"{r[0]}: {hpp} != {fab}+{acc}"


# ── 4. BOM_BARIS_MAHAL + KAIN_PER_MODEL ──────────────────────────────────────
def test_bom_baris_mahal(wb):
    ws = wb["BOM_BARIS_MAHAL"]
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
    assert len(rows) <= 200
    costs = [float(r[10] or 0) for r in rows]
    assert costs == sorted(costs, reverse=True), "not sorted by biaya_per_pcs desc"
    top = rows[0]
    assert top[0] == "DA-4104" and top[3] == "A-REN-0004", f"top row={top[:5]}"
    assert abs(float(top[10]) - 506506) < 500, f"biaya={top[10]}"
    # biaya_per_pcs ≈ qty_satuan_dasar × harga_per_satuan_dasar
    for r in rows:
        q, h, c = float(r[7] or 0), float(r[9] or 0), float(r[10] or 0)
        assert abs(c - q * h) <= 1, f"{r[0]}/{r[3]}: {c} != {q}*{h}"


def test_kain_per_model(wb):
    ws = wb["KAIN_PER_MODEL"]
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
    assert len(rows) > 0
    for r in rows:
        code = r[3] or ""
        assert code == "" or code.startswith("KN-"), f"bad fabric code: {code}"
        assert isinstance(r[5], (int, float)), f"qty not numeric: {r[5]}"


# ── 5. ROUNDTRIP fill-preview (no writes) ────────────────────────────────────
def test_fill_preview_roundtrip_no_changes(review_response, auth_headers):
    files = {"file": ("REVIEW_HARGA_MATERIAL_DA.xlsx", review_response.content,
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = requests.post(f"{BASE}/api/rahaza/master/fill-preview", headers=auth_headers, files=files, timeout=180)
    assert r.status_code == 200, r.text[:400]
    js = r.json()
    assert js["ok"] is True, js.get("errors")
    assert js["errors"] == []
    assert js["totals"]["materials"] == 0, f"materials should be 0, got {js['totals']['materials']}"


def test_fill_preview_edit_A_REN_0004(review_response, auth_headers):
    wb = openpyxl.load_workbook(io.BytesIO(review_response.content))
    ws = wb["MATERIAL"]
    found = False
    for r in ws.iter_rows(min_row=2):
        if r[0].value == "A-REN-0004":
            r[5].value = "roll"   # satuan_beli
            r[6].value = 20        # isi_per_satuan_beli
            r[7].value = 214500    # harga_per_satuan_beli
            found = True
            break
    assert found, "A-REN-0004 not found in MATERIAL"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    files = {"file": ("REVIEW_EDIT.xlsx", buf.read(),
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = requests.post(f"{BASE}/api/rahaza/master/fill-preview", headers=auth_headers, files=files, timeout=180)
    assert r.status_code == 200, r.text[:400]
    js = r.json()
    assert js["ok"] is True, js.get("errors")
    mats = [m for m in js["materials"] if m["code"] == "A-REN-0004"]
    assert len(mats) == 1, f"expected 1 change for A-REN-0004, got {len(mats)} / totals={js['totals']['materials']}"
    unit_cost = float(mats[0]["unit_cost"])
    assert abs(unit_cost - 10725) < 0.01, f"unit_cost={unit_cost} (expected 10725=214500/20)"
