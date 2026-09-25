"""QA iter208 — jalan pintas pengisian master (BOM copy/add + template harga·satuan·rekening + HPP standar).
Cleanup: kembalikan A-BIS-0001, 1-1211, SHP-04 ke keadaan semula karena DB = master NYATA."""
import io
import os
import pytest
import requests
import openpyxl
from pymongo import MongoClient

BASE = os.environ["REACT_APP_BACKEND_URL"].rstrip("/")
MONGO_URL = os.environ.get("MONGO_URL") or "mongodb://localhost:27017"
DB_NAME = os.environ.get("DB_NAME") or "test_database"

CREATOR_EMAIL = "iori.oliviara@creator.id"
MODEL_WITH_MISSING = "5068919a-16c1-4373-a549-c0e3c45dbcff"  # 2 varian tanpa BOM
ACC_MATERIAL_ID = "bd52d6ce-ecd4-454d-88bc-630cefe91576"  # A-BIS-0001


# ────────────────────────────────────────────────────────────────────
@pytest.fixture(scope="module")
def token():
    r = requests.post(f"{BASE}/api/auth/login", json={"email": "admin@garment.com", "password": "Admin@123"})
    assert r.status_code == 200, r.text
    return r.json()["token"]


@pytest.fixture(scope="module")
def auth(token):
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture(scope="module")
def loop():
    return None


@pytest.fixture(scope="module")
def db():
    c = MongoClient(MONGO_URL)
    return c[DB_NAME]


def _run(loop, coro):
    return coro


# ═══════════════ BACKEND ═══════════════
class TestFillTemplate:
    def test_download_template(self, auth):
        r = requests.get(f"{BASE}/api/rahaza/master/fill-template", headers=auth)
        assert r.status_code == 200
        assert "spreadsheet" in r.headers.get("content-type", "")
        wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
        assert set(wb.sheetnames) >= {"PETUNJUK", "MATERIAL", "REKENING", "TOKO"}
        # MATERIAL tidak memuat CUT-*
        rows = list(wb["MATERIAL"].iter_rows(values_only=True))
        cut = [r for r in rows[1:] if r and r[0] and str(r[0]).startswith("CUT-")]
        assert not cut, f"CUT-* seharusnya tak ada di template, ada {len(cut)}"

    def test_preview_valid(self, auth):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        m = wb.create_sheet("MATERIAL")
        m.append(["kode", "nama", "tipe", "kategori", "satuan_dasar", "satuan_beli", "isi_per_satuan_beli", "harga_per_satuan_beli", "harga_per_satuan_dasar_sekarang", "min_stok", "keterangan"])
        m.append(["A-BIS-0001", "", "", "", "m", "roll", 21.919, 337553, "", 0, ""])
        r = wb.create_sheet("REKENING")
        r.append(["kode_akun", "nama_akun", "bank", "no_rekening", "atas_nama"])
        r.append(["1-1211", "", "Bank BCA", "999", "QA"])
        t = wb.create_sheet("TOKO")
        t.append(["kode_toko", "nama_toko", "platform", "rekening_pencairan_kode_akun"])
        t.append(["SHP-04", "", "", "1-1213"])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        rp = requests.post(f"{BASE}/api/rahaza/master/fill-preview", headers=auth, files={"file": ("t.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert rp.status_code == 200, rp.text
        j = rp.json()
        assert j["ok"] is True, j
        assert j["totals"] == {"materials": 1, "accounts": 1, "stores": 1}
        assert abs(j["materials"][0]["unit_cost"] - 15400.02) < 0.05, j["materials"][0]

    def test_preview_errors(self, auth):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        m = wb.create_sheet("MATERIAL")
        m.append(["kode", "nama", "tipe", "kategori", "satuan_dasar", "satuan_beli", "isi_per_satuan_beli", "harga_per_satuan_beli", "harga_per_satuan_dasar_sekarang", "min_stok", "keterangan"])
        # harga terisi tapi isi kosong, satuan beli beda dari dasar
        m.append(["A-BIS-0001", "", "", "", "m", "roll", 0, 100000, "", 0, ""])
        t = wb.create_sheet("TOKO")
        t.append(["kode_toko", "nama_toko", "platform", "rekening_pencairan_kode_akun"])
        t.append(["ZZZ-999", "", "", "1-1213"])
        buf = io.BytesIO()
        wb.save(buf)
        rp = requests.post(f"{BASE}/api/rahaza/master/fill-preview", headers=auth, files={"file": ("t.xlsx", buf.getvalue())})
        assert rp.status_code == 200
        j = rp.json()
        assert j["ok"] is False
        joined = " | ".join(j["errors"])
        assert "isi_per_satuan_beli" in joined, joined
        assert "tidak ada" in joined, joined


class TestFillApply:
    def test_apply_and_cleanup(self, auth, db, loop):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        m = wb.create_sheet("MATERIAL")
        m.append(["kode", "nama", "tipe", "kategori", "satuan_dasar", "satuan_beli", "isi_per_satuan_beli", "harga_per_satuan_beli", "harga_per_satuan_dasar_sekarang", "min_stok", "keterangan"])
        m.append(["A-BIS-0001", "", "", "", "m", "roll", 21.919, 337553, "", 0, ""])
        r = wb.create_sheet("REKENING")
        r.append(["kode_akun", "nama_akun", "bank", "no_rekening", "atas_nama"])
        r.append(["1-1211", "", "Bank BCA", "999", "QA"])
        t = wb.create_sheet("TOKO")
        t.append(["kode_toko", "nama_toko", "platform", "rekening_pencairan_kode_akun"])
        t.append(["SHP-04", "", "", "1-1213"])
        buf = io.BytesIO()
        wb.save(buf)
        try:
            rp = requests.post(f"{BASE}/api/rahaza/master/fill-apply", headers=auth, files={"file": ("t.xlsx", buf.getvalue())})
            assert rp.status_code == 200, rp.text
            j = rp.json()
            assert j["materials_updated"] == 1
            assert j["accounts_updated"] == 1
            assert j["stores_updated"] == 1
            assert "panels_standard_costed" in j
            assert "models_hpp_applied" in j

            mat = _run(loop, db.rahaza_materials.find_one({"code": "A-BIS-0001"}, {"_id": 0}))
            assert abs(mat["unit_cost"] - 15400.02) < 0.05, mat["unit_cost"]
            assert mat.get("purchase_uom") == "roll"
            assert abs(float(mat.get("pack_size")) - 21.919) < 1e-3
            assert len(mat.get("uoms") or []) == 2
            acc = _run(loop, db.rahaza_cash_accounts.find_one({"gl_account_code": "1-1211"}, {"_id": 0}))
            assert acc.get("account_number") == "999"
            assert acc.get("holder_name") == "QA"
            store = _run(loop, db.marketing_platform_accounts.find_one({"account_code": "SHP-04"}, {"_id": 0}))
            assert store.get("coa_cash_code") == "1-1213"
            hist = _run(loop, db.rahaza_material_cost_history.count_documents({"material_code": "A-BIS-0001", "source": "template_harga_owner"}))
            assert hist >= 1
        finally:
            # RESTORE - master NYATA
            _run(loop, db.rahaza_materials.update_one({"code": "A-BIS-0001"}, {
                "$set": {"unit_cost": 15400.0, "pack_size": 1.0,
                         "uoms": [{"code": "m", "name": "M", "factor": 1.0, "is_base": True, "level": 0}]},
                "$unset": {"purchase_uom": "", "pack_unit": "", "cost_source": "", "price_updated_at": "", "standard_cost_basis": "", "value_status": "", "value_note": ""}
            }))
            _run(loop, db.rahaza_cash_accounts.update_one({"gl_account_code": "1-1211"}, {"$set": {"account_number": "", "holder_name": ""}}))
            _run(loop, db.marketing_platform_accounts.update_one({"account_code": "SHP-04"}, {"$unset": {"coa_cash_code": ""}}))
            _run(loop, db.rahaza_material_cost_history.delete_many({"material_code": "A-BIS-0001", "source": "template_harga_owner"}))


class TestBomShortcuts:
    def test_copy_missing_idempotent(self, auth, db, loop):
        r1 = requests.post(f"{BASE}/api/rahaza/master/bom/copy-missing", headers=auth, json={"model_id": MODEL_WITH_MISSING})
        assert r1.status_code == 200, r1.text
        j1 = r1.json()
        # created >=0; may be 0 if a prior run already copied. Just check idempoten on second call.
        r2 = requests.post(f"{BASE}/api/rahaza/master/bom/copy-missing", headers=auth, json={"model_id": MODEL_WITH_MISSING})
        assert r2.status_code == 200
        assert r2.json()["created"] == 0, r2.json()
        # verify copied BOM exists for a variant with copied_from_bom_id
        copied = _run(loop, db.rahaza_boms.find_one({"model_id": MODEL_WITH_MISSING, "copied_from_bom_id": {"$exists": True}}, {"_id": 0, "materials": 1, "color_code": 1, "copied_from_bom_id": 1}))
        # if j1.created > 0, must have copied BOM
        if j1["created"] > 0:
            assert copied is not None
            # CUT-* line should contain color_code of target
            cutlines = [ln for ln in copied["materials"] if ln.get("is_cut_panel")]
            for ln in cutlines:
                assert copied["color_code"].upper() in (ln.get("code") or "").upper()

    def test_add_lines_idempotent(self, auth, db, loop):
        payload = {"model_id": MODEL_WITH_MISSING, "lines": [{"material_id": ACC_MATERIAL_ID, "code": "A-BIS-0001", "name": "Bisban QA", "material_type": "accessory", "qty": 2, "unit": "pcs"}]}
        try:
            r1 = requests.post(f"{BASE}/api/rahaza/master/bom/add-lines", headers=auth, json=payload)
            assert r1.status_code == 200, r1.text
            j1 = r1.json()
            assert j1["boms_touched"] > 0
            assert j1["lines_appended"] > 0
            r2 = requests.post(f"{BASE}/api/rahaza/master/bom/add-lines", headers=auth, json=payload)
            assert r2.status_code == 200
            assert r2.json()["lines_appended"] == 0, r2.json()
        finally:
            # cleanup - pull the line we added
            _run(loop, db.rahaza_boms.update_many({"model_id": MODEL_WITH_MISSING}, {"$pull": {"materials": {"material_id": ACC_MATERIAL_ID}}}))


class TestRecalcHPP:
    def test_recalc(self, auth, db, loop):
        rp = requests.post(f"{BASE}/api/rahaza/master/recalc-hpp", headers=auth)
        assert rp.status_code == 200, rp.text
        j = rp.json()
        assert "panels_standard_costed" in j
        assert j["models_hpp_applied"] >= 70, j
        # models with hpp>0
        cnt = _run(loop, db.rahaza_models.count_documents({"hpp": {"$gt": 0}, "hpp_source": "bom"}))
        assert cnt >= 70, cnt
        # CUT-* standard cost
        std_cnt = _run(loop, db.rahaza_materials.count_documents({"code": {"$regex": "^CUT-"}, "value_status": "standard", "unit_cost": {"$gt": 0}}))
        assert std_cnt > 0


class TestCreatorPortalLogin:
    def test_login_success(self):
        r = requests.post(f"{BASE}/api/marketing/creator-portal/auth/login", json={"email": CREATOR_EMAIL, "password": "Dewi@123"})
        assert r.status_code == 200, r.text
        j = r.json()
        assert j.get("token") or j.get("access_token"), j

    def test_login_wrong_pw(self):
        r = requests.post(f"{BASE}/api/marketing/creator-portal/auth/login", json={"email": CREATOR_EMAIL, "password": "SALAH"})
        assert r.status_code == 401, r.status_code


class TestShiftEmployees:
    def test_shift_default(self, db, loop):
        s = _run(loop, db.rahaza_shifts.find_one({"start_time": "08:00", "end_time": "16:00", "active": True}, {"_id": 0}))
        assert s is not None
        # all active employees must have shift_id
        no_shift = _run(loop, db.rahaza_employees.count_documents({"active": True, "$or": [{"shift_id": None}, {"shift_id": ""}, {"shift_id": {"$exists": False}}]}))
        assert no_shift == 0
        # all point to this shift
        others = _run(loop, db.rahaza_employees.count_documents({"active": True, "shift_id": {"$ne": s["id"]}}))
        assert others == 0, others
