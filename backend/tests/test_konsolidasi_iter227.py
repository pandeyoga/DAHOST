"""
Iteration 227 backend tests — Konsolidasi master data CV. Dewi Aditya.
Read-only verification of DB invariants, prices, HPP, R&D and public API/downloads.
"""
import os
import io
import subprocess
import pytest
import requests
from pymongo import MongoClient
from openpyxl import load_workbook

# ---- Setup ----
BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/") or "https://sku-inventory-pull.preview.emergentagent.com"
MONGO_URL = "mongodb://localhost:27017"
DB_NAME = "test_database"

# Load real values from backend/.env
_env = "/app/backend/.env"
if os.path.exists(_env):
    for line in open(_env):
        line = line.strip()
        if line.startswith("MONGO_URL"):
            MONGO_URL = line.split("=", 1)[1].strip().strip('"')
        elif line.startswith("DB_NAME"):
            DB_NAME = line.split("=", 1)[1].strip().strip('"')

_client = MongoClient(MONGO_URL)
db = _client[DB_NAME]


@pytest.fixture(scope="session")
def api_token():
    r = requests.post(f"{BASE_URL}/api/auth/login",
                      json={"email": "admin@garment.com", "password": "Admin@123"},
                      timeout=30)
    assert r.status_code == 200, f"Login failed: {r.status_code} {r.text}"
    data = r.json()
    token = data.get("access_token") or data.get("token")
    assert token, f"No token in login response: {data}"
    return token


@pytest.fixture(scope="session")
def auth_headers(api_token):
    return {"Authorization": f"Bearer {api_token}"}


# ---- Module: DB invariants ----
class TestDBInvariants:
    def test_variant_and_bom_counts(self):
        assert db.rahaza_model_variants.count_documents({}) == 807
        assert db.rahaza_model_variants.count_documents({"active": True}) == 777
        assert db.rahaza_boms.count_documents({}) == 807
        assert db.rahaza_boms.count_documents({"active": True}) == 777

    def test_all_boms_version_1(self):
        assert db.rahaza_boms.distinct("version") == [1]

    def test_active_variant_has_exactly_one_active_bom(self):
        pipeline = [
            {"$match": {"active": True}},
            {"$group": {"_id": {"model_id": "$model_id", "cc": {"$toUpper": "$color_code"}, "size_id": "$size_id"},
                        "n": {"$sum": 1}}},
            {"$match": {"n": {"$ne": 1}}},
        ]
        dups = list(db.rahaza_boms.aggregate(pipeline))
        assert dups == [], f"Duplicate active BOM keys: {dups[:5]}"

        # each active variant has exactly one active BOM
        missing = 0
        checked = 0
        for v in db.rahaza_model_variants.find({"active": True}, {"model_id": 1, "color_code": 1, "size_id": 1}):
            checked += 1
            n = db.rahaza_boms.count_documents({
                "model_id": v["model_id"],
                "color_code": (v["color_code"] or "").upper(),
                "size_id": v["size_id"],
                "active": True,
            })
            if n != 1:
                missing += 1
        assert missing == 0, f"{missing}/{checked} active variants without exactly one active BOM"


# ---- Module: BOM content rules ----
class TestBOMContent:
    def test_each_active_bom_has_cut_hangtag_pin_accessory(self):
        problems = []
        models_by_id = {m["id"]: m for m in db.rahaza_models.find({}, {"id": 1, "code": 1})}
        sizes_by_id = {s["id"]: s for s in db.rahaza_sizes.find({}, {"id": 1, "code": 1})}
        for bom in db.rahaza_boms.find({"active": True}):
            items = bom.get("materials", [])
            cuts = [i for i in items if i.get("is_cut_panel")]
            model = models_by_id.get(bom["model_id"], {})
            size = sizes_by_id.get(bom["size_id"], {})
            expected_prefix = f"CUT-{model.get('code','')}-{(bom.get('color_code') or '').upper()}-{size.get('code','')}"
            if not cuts:
                problems.append(f"{bom['id']} no CUT panel")
                continue
            if not any((i.get("code") or "").startswith(expected_prefix) for i in cuts):
                problems.append(f"{bom['id']} cut prefix mismatch expected {expected_prefix}")
            htg = [i for i in items if i.get("code") == "A-HTG-0003"]
            pin = [i for i in items if i.get("code") == "A-PIN-0001"]
            if not htg or htg[0].get("qty") != 1 or htg[0].get("unit") != "pcs":
                problems.append(f"{bom['id']} bad hangtag")
            if not pin or pin[0].get("qty") != 1 or pin[0].get("unit") != "pcs":
                problems.append(f"{bom['id']} bad pin")
            other_acc = [i for i in items
                         if i.get("material_type") == "accessory"
                         and i.get("code") not in ("A-HTG-0003", "A-PIN-0001")]
            if not other_acc:
                problems.append(f"{bom['id']} no extra accessory")
        assert not problems, f"{len(problems)} BOM issues, sample: {problems[:5]}"

    def test_material_links_and_costs(self):
        mats = {m["id"]: m for m in db.rahaza_materials.find({}, {"id": 1, "unit_cost": 1, "code": 1})}
        problems = []
        for bom in db.rahaza_boms.find({"active": True}):
            for it in bom.get("materials", []):
                mid = it.get("material_id")
                if not mid or mid not in mats:
                    problems.append(f"{bom['id']} {it.get('code')} unlinked material_id")
                    continue
                if it.get("uom_status") in ("mismatch", "unlinked"):
                    problems.append(f"{bom['id']} {it.get('code')} uom_status={it.get('uom_status')}")
                exp_cost = mats[mid].get("unit_cost")
                got = it.get("unit_cost_base")
                if exp_cost is None or got is None:
                    continue
                if abs(float(exp_cost) - float(got)) > 0.01:
                    problems.append(f"{bom['id']} {it.get('code')} cost {got} != material {exp_cost}")
        assert not problems, f"{len(problems)} link/cost issues, sample: {problems[:5]}"


# ---- Module: Material prices ----
class TestMaterialPrices:
    @pytest.mark.parametrize("code,unit,cost", [
        ("A-K22-0014", "pcs", 46.0),
        ("A-REN-0004", "cm", 11.9167),
        ("A-KRT-0007", "cm", 0.4701),
        ("A-BIS-0001", "cm", 3.4899),
        ("A-LBL-0004", "pcs", 61.8),
        ("A-PIN-0001", "pcs", 3.96),
        ("A-HTG-0003", "pcs", 88.0),
    ])
    def test_price(self, code, unit, cost):
        m = db.rahaza_materials.find_one({"code": code})
        assert m, f"missing {code}"
        assert m["unit"] == unit
        assert abs(float(m["unit_cost"]) - cost) < 0.001

    def test_a_k22_uoms_gross_factor(self):
        m = db.rahaza_materials.find_one({"code": "A-K22-0014"})
        gross = [u for u in m["uoms"] if u["code"] == "gross"]
        assert gross and gross[0]["factor"] == 1440.0

    def test_a_lbl_roll_1000(self):
        m = db.rahaza_materials.find_one({"code": "A-LBL-0004"})
        roll = [u for u in m["uoms"] if u["code"] == "roll"]
        assert roll and int(roll[0]["factor"]) == 1000

    def test_a_pin_pack_5000(self):
        m = db.rahaza_materials.find_one({"code": "A-PIN-0001"})
        pack = [u for u in m["uoms"] if u["code"] == "pack"]
        assert pack and int(pack[0]["factor"]) == 5000

    def test_a_tlk_0001_exists(self):
        m = db.rahaza_materials.find_one({"code": "A-TLK-0001"})
        assert m and m["unit"] == "cm" and float(m["unit_cost"]) == 0.0


# ---- Module: laporan_kekosongan_bom script ----
class TestLaporanKekosongan:
    def test_run_script(self):
        cmd = "cd /app/backend && set -a && . .env && set +a && python ../scripts/laporan_kekosongan_bom.py"
        r = subprocess.run(["bash", "-c", cmd], capture_output=True, text=True, timeout=180)
        out = (r.stdout or "") + (r.stderr or "")
        assert r.returncode == 0, f"script failed: {out[-2000:]}"
        # expected: 98 E_LENGKAP, 6 F_DIHENTIKAN, 0 harus diisi
        import re
        e_lengkap = re.search(r"(\d+)\s+E_LENGKAP", out)
        f_diht = re.search(r"(\d+)\s+F_DIHENTIKAN", out)
        assert e_lengkap and int(e_lengkap.group(1)) == 98, out[-1500:]
        assert f_diht and int(f_diht.group(1)) == 6, out[-1500:]
        must = re.search(r"HARUS DIISI:\s*(\d+)", out)
        assert must and int(must.group(1)) == 0, out[-500:]


# ---- Module: HPP model values ----
class TestHPP:
    @pytest.mark.parametrize("code,expected", [
        ("DA-4104", 10572.13),
        ("DA-1101", 37227.58),
        ("DA-1508", 22051.62),
        ("DA-1509", 21546.14),
    ])
    def test_hpp(self, code, expected):
        m = db.rahaza_models.find_one({"code": code})
        assert m, f"model {code} missing"
        hpp = float(m.get("hpp") or 0)
        assert abs(hpp - expected) < 1.0, f"{code} hpp={hpp} expected {expected}"
        validation = m.get("hpp_validation") or {}
        assert validation.get("status") == "tervalidasi", f"{code} status={validation.get('status')}"


# ---- Module: Variant creation rules ----
class TestVariantsCreated:
    @pytest.mark.parametrize("sku", [
        "DA-1209-HTM-S", "DA-1210-HTM-M", "DA-1211-HTM-L", "DA-3302-HTM-XL",
    ])
    def test_new_variants_exist_active_with_bom(self, sku):
        v = db.rahaza_model_variants.find_one({"sku": sku})
        assert v and v.get("active") is True, f"{sku} missing/inactive"
        bom = db.rahaza_boms.find_one({
            "model_id": v["model_id"], "size_id": v["size_id"],
            "color_code": v["color_code"].upper(), "active": True,
        })
        assert bom, f"{sku} no active BOM"
        # FG in materials
        fg = db.rahaza_materials.find_one({"code": sku, "type": "fg"})
        assert fg, f"{sku} no FG material"

    def test_da_2201_has_accessory(self):
        # at least one DA-2201-* active variant with accessory besides hangtag/pin
        v = db.rahaza_model_variants.find_one({"model_code": "DA-2201", "active": True})
        assert v, "no active DA-2201 variant"
        bom = db.rahaza_boms.find_one({"model_id": v["model_id"], "color_code": v["color_code"].upper(),
                                       "size_id": v["size_id"], "active": True})
        acc = [i for i in bom["materials"] if i.get("material_type") == "accessory"
               and i.get("code") not in ("A-HTG-0003", "A-PIN-0001")]
        assert acc, "DA-2201 has no extra accessory"

    def test_forbidden_variants_absent(self):
        for pattern in ["DA-3604-", "DA-4201-"]:
            n = db.rahaza_model_variants.count_documents(
                {"sku": {"$regex": f"^{pattern}.*-ALLSIZE$"}, "active": True})
            assert n == 0, f"{pattern}*-ALLSIZE still active: {n}"
        assert db.rahaza_model_variants.count_documents(
            {"sku": {"$regex": "^DA-3306-"}, "active": True}) == 0

    def test_discontinued_models_inactive(self):
        for code in ["DA-2104", "DA-2107", "DA-2108", "DA-3601", "DA-3602"]:
            m = db.rahaza_models.find_one({"code": code})
            if not m:
                continue
            n_active = db.rahaza_model_variants.count_documents({"model_id": m["id"], "active": True})
            assert n_active == 0, f"{code} still has {n_active} active variants"


# ---- Module: R&D collections ----
class TestRnD:
    def test_tech_packs(self):
        assert db.dewi_rnd_tech_packs.count_documents({}) == 104
        # 98 active styles must have populated bom_items; 6 discontinued may be empty
        discontinued = {"DA-3601", "DA-2501", "DA-2104", "DA-3602", "DA-2107", "DA-2108"}
        problems = []
        for tp in db.dewi_rnd_tech_packs.find({}):
            if tp.get("version") != "v1":
                problems.append(f"{tp.get('style_code')} version={tp.get('version')}")
            if tp.get("status") != "draft":
                problems.append(f"{tp.get('style_code')} status={tp.get('status')}")
            if tp.get("is_latest") is not True:
                problems.append(f"{tp.get('style_code')} is_latest={tp.get('is_latest')}")
            if not tp.get("bom_items") and tp.get("style_code") not in discontinued:
                problems.append(f"{tp.get('style_code')} empty bom_items (active style)")
            if tp.get("bom_unlinked_count", 0) != 0:
                problems.append(f"{tp.get('style_code')} unlinked={tp.get('bom_unlinked_count')}")
        assert not problems, problems[:5]

    def test_rnd_materials_count(self):
        assert db.dewi_rnd_materials.count_documents({}) == 12
        # material_code prefix KN-
        n_kn = db.dewi_rnd_materials.count_documents({"material_code": {"$regex": "^KN-"}})
        assert n_kn == 12

    def test_rnd_styles(self):
        assert db.dewi_rnd_styles.count_documents({}) == 104
        n_promoted = db.dewi_rnd_styles.count_documents({"promoted_to_model_id": {"$exists": True, "$nin": [None, ""]}})
        assert n_promoted == 104


# ---- Module: API ----
class TestAPI:
    def test_completeness(self, auth_headers):
        r = requests.get(f"{BASE_URL}/api/dewi/rnd/completeness", headers=auth_headers, timeout=30)
        assert r.status_code == 200, r.text
        d = r.json()
        mc = d.get("missing_counts", {})
        assert mc.get("bom") == 0, d
        assert mc.get("accessories") == 0, d
        assert mc.get("techpack") == 0, d
        assert d.get("total_models") == 104, d
        assert d.get("discontinued_models") == 6, d

    def test_tech_packs_search_lyora(self, auth_headers):
        r = requests.get(f"{BASE_URL}/api/dewi/rnd/tech-packs",
                         params={"search": "DA-1101"}, headers=auth_headers, timeout=30)
        assert r.status_code == 200, r.text
        d = r.json()
        if isinstance(d, list):
            items = d
        else:
            items = d.get("items") or d.get("results") or d.get("data") or d.get("tech_packs") or []
        assert len(items) == 1, f"expected 1 tech pack, got {len(items)} → {str(d)[:500]}"
        tp = items[0]
        assert tp.get("version") == "v1"
        assert tp.get("status") == "draft"
        assert len(tp.get("bom_items") or []) == 6, tp
        assert len(tp.get("colorways") or []) == 6, tp

    def test_harga_review_xlsx(self, auth_headers):
        r = requests.get(f"{BASE_URL}/api/rahaza/master/harga-review",
                         headers=auth_headers, timeout=60)
        assert r.status_code == 200, r.text[:500]
        wb = load_workbook(io.BytesIO(r.content))
        assert "MATERIAL" in wb.sheetnames, wb.sheetnames
        ws = wb["MATERIAL"]
        # count MERAH
        rows = list(ws.iter_rows(values_only=True))
        header = [str(x or "").strip().lower() for x in rows[0]]
        # find tingkat column
        idx = None
        for i, h in enumerate(header):
            if "tingkat" in h or "level" in h:
                idx = i
                break
        assert idx is not None, header
        merah = sum(1 for r in rows[1:] if r[idx] and "merah" in str(r[idx]).lower())
        assert merah <= 4, f"MERAH={merah}"


# ---- Module: Public downloads ----
class TestDownloads:
    def test_export_sku_bom(self):
        r = requests.get(f"{BASE_URL}/downloads/EXPORT_SKU_BOM_DA.xlsx", timeout=60)
        assert r.status_code == 200, r.status_code
        wb = load_workbook(io.BytesIO(r.content), read_only=True)
        assert "RINGKASAN_SKU" in wb.sheetnames, wb.sheetnames
        ws = wb["RINGKASAN_SKU"]
        rows = list(ws.iter_rows(values_only=True))
        header = [str(x or "").strip() for x in rows[0]]
        data = rows[1:]
        # skip trailing empty rows
        data = [r for r in data if any(c is not None and str(c).strip() for c in r)]
        assert len(data) == 807, f"data rows={len(data)}"
        # find status_bom column
        try:
            idx = [h.lower() for h in header].index("status_bom")
        except ValueError:
            # try other names
            idx = None
            for i, h in enumerate(header):
                if "status" in h.lower() and "bom" in h.lower():
                    idx = i
                    break
        assert idx is not None, header
        statuses = [str(r[idx]) for r in data]
        lengkap = sum(1 for s in statuses if s == "BOM_LENGKAP")
        nonaktif = sum(1 for s in statuses if s == "BOM_NONAKTIF")
        assert lengkap == 777, lengkap
        assert nonaktif == 30, nonaktif

    def test_laporan_md(self):
        r = requests.get(f"{BASE_URL}/downloads/LAPORAN_KONSOLIDASI_2026-09-23.md", timeout=30)
        assert r.status_code == 200, r.status_code


# ---- Module: Konsolidasi dry-run idempotency ----
class TestKonsolidasiDryRun:
    def test_dry_run(self):
        cmd = "cd /app/backend && set -a && . .env && set +a && python ../scripts/konsolidasi/konsolidasi.py"
        r = subprocess.run(["bash", "-c", cmd], capture_output=True, text=True, timeout=600)
        out = (r.stdout or "") + (r.stderr or "")
        assert r.returncode == 0, f"dry-run failed rc={r.returncode}\n{out[-3000:]}"
        # Must not include bom_update or bom_baru with non-zero counts, or varian.baru>0
        import re, json as _json
        # try parse json output
        # simple textual checks
        low = out.lower()
        # bom_update / bom_baru appearances with value != 0
        for key in ["bom_update", "bom_baru"]:
            for m in re.finditer(rf"{key}[^0-9\-]*(\d+)", low):
                assert int(m.group(1)) == 0, f"{key}={m.group(1)} not idempotent\n{out[-1500:]}"
        for m in re.finditer(r"varian[^a-z]*baru[^0-9\-]*(\d+)", low):
            assert int(m.group(1)) == 0, f"varian.baru={m.group(1)}\n{out[-1500:]}"
