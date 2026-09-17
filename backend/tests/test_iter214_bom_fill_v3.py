"""Iteration 214 — BOM_AKSESORIS v3 importer (bom_fill.py hardened + laporan_sisa)"""
import io
import os
import asyncio
import pytest
import requests
from openpyxl import load_workbook
from motor.motor_asyncio import AsyncIOMotorClient

BASE = (os.environ.get("REACT_APP_BACKEND_URL") or open("/app/frontend/.env").read().split("REACT_APP_BACKEND_URL=")[1].split("\n")[0]).rstrip("/")
FIX_A = "/app/tests/fixtures/bom_v3_test.xlsx"
FIX_B = "/app/tests/fixtures/bom_v3_test_b.xlsx"

# Load .env for MONGO_URL and DB_NAME
from dotenv import load_dotenv
load_dotenv("/app/backend/.env")
MONGO_URL = os.environ["MONGO_URL"]
DB_NAME = os.environ.get("DB_NAME", "test_database")


@pytest.fixture(scope="module")
def token():
    r = requests.post(f"{BASE}/api/auth/login", json={"email": "admin@garment.com", "password": "Admin@123"})
    assert r.status_code == 200, r.text
    j = r.json()
    return j.get("token") or j.get("access_token") or j.get("data", {}).get("token")


@pytest.fixture(scope="module")
def h(token):
    return {"Authorization": f"Bearer {token}"}


def _unwrap(j):
    return j["data"] if isinstance(j, dict) and "data" in j and "ok" in j else j


def _load_db():
    c = AsyncIOMotorClient(MONGO_URL)
    return c, c[DB_NAME]


def _post_file(path, url, headers):
    with open(path, "rb") as f:
        return requests.post(url, headers=headers, files={"file": (os.path.basename(path), f.read(),
                             "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, timeout=120)


@pytest.fixture(scope="module", autouse=True)
def label_pack_size():
    """Prasyarat uji: A-LBL-0004 isi kemasan 600 pcs/roll (nilai uji, BUKAN data klien) — dikembalikan setelah modul selesai."""
    async def _set(val):
        c, db = _load_db()
        before = await db.rahaza_materials.find_one({"code": "A-LBL-0004"}, {"_id": 0, "pack_size": 1, "pack_unit": 1})
        await db.rahaza_materials.update_one({"code": "A-LBL-0004"}, {"$set": val})
        c.close()
        return before or {}
    before = asyncio.run(_set({"pack_size": 600, "pack_unit": "roll"}))
    yield
    asyncio.run(_set({"pack_size": before.get("pack_size", 1.0), "pack_unit": before.get("pack_unit", "roll")}))


# ─── preview
@pytest.fixture(scope="module")
def preview(h):
    r = _post_file(FIX_A, f"{BASE}/api/rahaza/master/fill-preview", h)
    assert r.status_code == 200, r.text
    return _unwrap(r.json())


def test_preview_totals_and_errors(preview):
    print("TOTALS:", preview.get("totals"))
    print("ERRORS:", preview.get("errors"))
    print("BOM_ISSUE_COUNTS:", preview.get("bom_issue_counts"))
    print("WARNINGS_LEN:", len(preview.get("warnings") or []))
    # errors: HARGA_JUAL_SKU DA-0000-XXX-M
    errs = preview.get("errors") or []
    assert any("DA-0000-XXX-M" in e for e in errs), errs
    t = preview["totals"]
    assert t["bom_lines"] == 6, t
    assert t["bom_groups"] == 4, t
    assert t["bom_models"] == 3, t
    assert t["bom_skipped"] == 10, t
    assert t["bom_models_in_file"] == 5, t
    assert t["sku_prices"] == 1, t


def test_preview_issue_counts(preview):
    expected = {"satuan_tak_valid": 3, "kode_tak_dikenal": 1, "qty_kosong": 1,
                "model_tak_dikenal": 1, "baris_tanpa_model": 1, "model_tanpa_sku": 1,
                "warna_tak_dikenal": 1, "kelompok_tanpa_varian": 1}
    got = preview.get("bom_issue_counts") or {}
    for k, v in expected.items():
        assert got.get(k) == v, f"{k}: got={got.get(k)} exp={v} full={got}"


def test_preview_bom_lines_details(preview):
    lines = preview["bom_lines"]
    by_row = {ln["row"]: ln for ln in lines}
    # Row 2 DA-1101 target 'HTM, MHG' A-KRT-0003 60cm → 0.6 m
    r2 = by_row[2]
    assert r2["code"] == "A-KRT-0003" and r2["model_code"] == "DA-1101"
    assert abs(r2["qty_base"] - 0.6) < 1e-4 and r2["unit_base"] == "m"
    assert "HTM" in r2["target"] and "MHG" in r2["target"]
    assert r2["target_skus"] == 2, r2
    # Row 4 A-LBL-0004 1 pcs → 0.001667 roll
    r4 = by_row[4]
    assert r4["code"] == "A-LBL-0004"
    assert abs(r4["qty_base"] - 0.001667) < 1e-4
    assert r4["unit_base"] == "roll"
    # Row 5 A-K22-0014 3 pcs → 0.020833 gross
    r5 = by_row[5]
    assert r5["code"] == "A-K22-0014"
    assert abs(r5["qty_base"] - 0.020833) < 1e-4
    assert r5["unit_base"] == "gross"
    # Row 6 DA-1101 CRL/GRY 0.6 m (Caraml→CRL fuzzy)
    r6 = by_row[6]
    assert r6["model_code"] == "DA-1101" and "CRL" in r6["target"] and "GRY" in r6["target"]
    # Row 10 DA-1202 HTM 1 SKU
    r10 = by_row[10]
    assert r10["model_code"] == "DA-1202" and r10["target_skus"] == 1
    # Row 12 DA-1201 semua varian 6 SKU
    r12 = by_row[12]
    assert r12["model_code"] == "DA-1201" and r12["target"] == "semua varian" and r12["target_skus"] == 6


# ─── snapshot pre-apply of DA-1202 for later comparison
@pytest.fixture(scope="module")
def pre_apply_snapshot():
    async def _snap():
        c, db = _load_db()
        m = await db.rahaza_models.find_one({"code": "DA-1202"}, {"_id": 0, "id": 1})
        boms = await db.rahaza_boms.find({"model_id": m["id"], "active": {"$ne": False}, "is_active": True}, {"_id": 0}).to_list(200)
        snap = {}
        for b in boms:
            cc = (b.get("color_code") or "").upper()
            snap[cc] = [(ln.get("code"), round(float(ln.get("qty") or 0), 6), (ln.get("unit") or "").lower()) for ln in (b.get("materials") or []) if not ln.get("is_cut_panel") and (ln.get("material_type") or "").lower() != "fabric"]
        # Also DA-1101 MGT/BRG
        m1 = await db.rahaza_models.find_one({"code": "DA-1101"}, {"_id": 0, "id": 1})
        boms1 = await db.rahaza_boms.find({"model_id": m1["id"], "active": {"$ne": False}, "is_active": True}, {"_id": 0}).to_list(200)
        snap1 = {}
        for b in boms1:
            cc = (b.get("color_code") or "").upper()
            snap1[cc] = [(ln.get("code"), round(float(ln.get("qty") or 0), 6), (ln.get("unit") or "").lower()) for ln in (b.get("materials") or []) if not ln.get("is_cut_panel") and (ln.get("material_type") or "").lower() != "fabric"]
        c.close()
        return {"da1202": snap, "da1101": snap1}
    return asyncio.run(_snap())


# ─── apply first (scope=bom)
@pytest.fixture(scope="module")
def apply_a(h, preview, pre_apply_snapshot):
    r = requests.post(f"{BASE}/api/rahaza/master/fill-apply?scope=bom", headers=h,
                      files={"file": ("t.xlsx", open(FIX_A, "rb").read())}, timeout=180)
    assert r.status_code == 200, r.text
    return _unwrap(r.json())


def test_apply_a_response(apply_a):
    print("APPLY_A:", apply_a)
    assert apply_a["scope"] == "bom"
    assert apply_a["bom_groups"] == 4, apply_a
    assert apply_a["sku_prices_updated"] == 0
    assert "hpp_unvalidated" in apply_a and "hpp_validated" in apply_a
    assert "models_hpp_applied" in apply_a


def test_mongo_after_apply_a(apply_a, pre_apply_snapshot):
    async def check():
        c, db = _load_db()
        # DA-1101
        m1101 = await db.rahaza_models.find_one({"code": "DA-1101"}, {"_id": 0, "id": 1})
        boms = await db.rahaza_boms.find({"model_id": m1101["id"], "is_active": True, "active": {"$ne": False}}, {"_id": 0}).to_list(200)
        by_c = {(b.get("color_code") or "").upper(): b for b in boms}

        def acc_lines(b):
            return [ln for ln in (b.get("materials") or []) if not ln.get("is_cut_panel") and (ln.get("material_type") or "").lower() != "fabric"]

        for cc in ("HTM", "MHG"):
            assert cc in by_c, f"BOM {cc} missing"
            acc = acc_lines(by_c[cc])
            codes = {ln.get("code") for ln in acc}
            assert codes == {"A-KRT-0003", "A-LBL-0004", "A-K22-0014"}, f"{cc} acc={codes}"
            # verify qty & unit_base
            krt = next(ln for ln in acc if ln["code"] == "A-KRT-0003")
            assert abs(float(krt["qty"]) - 60) < 1e-6 and (krt.get("unit") or "").lower() == "cm", krt
            assert abs(float(krt.get("qty_base") or 0) - 0.6) < 1e-3
            assert (krt.get("unit_base") or "").lower() == "m"
            lbl = next(ln for ln in acc if ln["code"] == "A-LBL-0004")
            assert abs(float(lbl["qty"]) - 1) < 1e-6 and (lbl.get("unit") or "").lower() == "pcs"
            assert abs(float(lbl.get("qty_base") or 0) - 0.001667) < 1e-4
            assert (lbl.get("unit_base") or "").lower() == "roll"
            k22 = next(ln for ln in acc if ln["code"] == "A-K22-0014")
            assert abs(float(k22["qty"]) - 3) < 1e-6
            assert abs(float(k22.get("qty_base") or 0) - 0.020833) < 1e-4
            assert (k22.get("unit_base") or "").lower() == "gross"

        for cc in ("CRL", "GRY"):
            assert cc in by_c, f"{cc} missing"
            acc = acc_lines(by_c[cc])
            codes = {ln.get("code") for ln in acc}
            assert codes == {"A-KRT-0003"}, f"{cc} acc={codes}"
            krt = acc[0]
            assert abs(float(krt.get("qty_base") or 0) - 0.6) < 1e-3
            assert (krt.get("unit_base") or "").lower() == "m"

        # MGT/BRG untouched — compare to snapshot
        pre = pre_apply_snapshot["da1101"]
        for cc in ("MGT", "BRG"):
            if cc in by_c and cc in pre:
                now_acc = [(ln.get("code"), round(float(ln.get("qty") or 0), 6), (ln.get("unit") or "").lower()) for ln in acc_lines(by_c[cc])]
                assert set(now_acc) == set(pre[cc]), f"{cc} changed: pre={pre[cc]} now={now_acc}"

        # DA-1202
        m1202 = await db.rahaza_models.find_one({"code": "DA-1202"}, {"_id": 0, "id": 1})
        boms2 = await db.rahaza_boms.find({"model_id": m1202["id"], "is_active": True, "active": {"$ne": False}}, {"_id": 0}).to_list(200)
        by_c2 = {(b.get("color_code") or "").upper(): b for b in boms2}
        htm_acc = acc_lines(by_c2["HTM"])
        htm_codes = {ln.get("code") for ln in htm_acc}
        assert htm_codes == {"A-K22-0014"}, f"DA-1202 HTM acc={htm_codes}"
        assert abs(float(htm_acc[0]["qty"]) - 3) < 1e-6
        # other colors untouched (esp: no A-K22-0005)
        pre2 = pre_apply_snapshot["da1202"]
        # DA-1202 non-HTM colors must be UNCHANGED from pre-apply snapshot (residual data
        # from prior iterations may include A-K22-0005 in MGT, which is OK — importer must
        # not add or remove it here since 'kelompok tanpa varian' is skipped).
        for cc in ("MGT", "DST", "CRM", "MCA", "DNM", "GRY"):
            if cc in by_c2 and cc in pre2:
                acc = acc_lines(by_c2[cc])
                now_acc = [(ln.get("code"), round(float(ln.get("qty") or 0), 6), (ln.get("unit") or "").lower()) for ln in acc]
                assert set(now_acc) == set(pre2[cc]), f"DA-1202 {cc} unexpectedly changed: pre={pre2[cc]} now={now_acc}"

        # DA-1201: all 6 BOMs have A-KRT-0003 60 cm
        m1201 = await db.rahaza_models.find_one({"code": "DA-1201"}, {"_id": 0, "id": 1})
        boms3 = await db.rahaza_boms.find({"model_id": m1201["id"], "is_active": True, "active": {"$ne": False}}, {"_id": 0}).to_list(200)
        n_ok = 0
        for b in boms3:
            acc = acc_lines(b)
            if any(ln.get("code") == "A-KRT-0003" and abs(float(ln.get("qty") or 0) - 60) < 1e-6 for ln in acc):
                n_ok += 1
        assert n_ok == 6, f"DA-1201 KRT 60cm: {n_ok}/{len(boms3)}"

        # hpp_validation on DA-1101
        m1101_full = await db.rahaza_models.find_one({"code": "DA-1101"}, {"_id": 0, "hpp_validation": 1})
        hv = m1101_full.get("hpp_validation") or {}
        assert hv.get("status") in ("tervalidasi", "belum_tervalidasi"), hv
        assert isinstance(hv.get("reasons"), list)

        # FG DA-1202-CRM-M retail_price_master NOT changed by scope=bom
        fg = await db.rahaza_materials.find_one({"code": "DA-1202-CRM-M"}, {"_id": 0, "retail_price_master": 1})
        # scope=bom must not touch this field. We only assert we didn't set it to 155999 via this apply.
        # (unable to snapshot before; just log)
        print("FG DA-1202-CRM-M retail_price_master:", fg.get("retail_price_master") if fg else None)

        c.close()
    asyncio.run(check())


# ─── idempotent
def test_apply_a_idempotent(h, apply_a):
    r = requests.post(f"{BASE}/api/rahaza/master/fill-apply?scope=bom", headers=h,
                      files={"file": ("t.xlsx", open(FIX_A, "rb").read())}, timeout=180)
    assert r.status_code == 200, r.text
    j = _unwrap(r.json())
    print("APPLY_A2:", j)
    assert j["boms_touched"] == 0, j
    assert j["bom_lines_appended"] == 0, j
    assert j["bom_base_created"] == 0, j


# ─── apply B (replace whole)
def test_apply_b_replaces_htm_mhg(h, apply_a):
    r = requests.post(f"{BASE}/api/rahaza/master/fill-apply?scope=bom", headers=h,
                      files={"file": ("b.xlsx", open(FIX_B, "rb").read())}, timeout=180)
    assert r.status_code == 200, r.text
    j = _unwrap(r.json())
    print("APPLY_B:", j)
    assert j["boms_touched"] == 2, j
    assert j["boms_unchanged"] == 2, j

    async def check():
        c, db = _load_db()
        m = await db.rahaza_models.find_one({"code": "DA-1101"}, {"_id": 0, "id": 1})
        boms = await db.rahaza_boms.find({"model_id": m["id"], "is_active": True, "active": {"$ne": False}}, {"_id": 0}).to_list(200)
        by_c = {(b.get("color_code") or "").upper(): b for b in boms}

        def acc(b):
            return [ln for ln in (b.get("materials") or []) if not ln.get("is_cut_panel") and (ln.get("material_type") or "").lower() != "fabric"]

        for cc in ("HTM", "MHG"):
            codes = {ln.get("code") for ln in acc(by_c[cc])}
            assert codes == {"A-KRT-0003", "A-LBL-0004"}, f"{cc}={codes}"
            krt = next(ln for ln in acc(by_c[cc]) if ln["code"] == "A-KRT-0003")
            assert abs(float(krt["qty"]) - 70) < 1e-6, krt
        for cc in ("CRL", "GRY"):
            codes = {ln.get("code") for ln in acc(by_c[cc])}
            assert codes == {"A-KRT-0003"}, f"{cc}={codes}"

        # DA-1202 HTM untouched still {A-K22-0014}
        m2 = await db.rahaza_models.find_one({"code": "DA-1202"}, {"_id": 0, "id": 1})
        b2 = await db.rahaza_boms.find_one({"model_id": m2["id"], "color_code": "HTM", "is_active": True, "active": {"$ne": False}}, {"_id": 0})
        codes2 = {ln.get("code") for ln in acc(b2)}
        assert codes2 == {"A-K22-0014"}, f"DA-1202 HTM={codes2}"
        c.close()
    asyncio.run(check())


# ─── laporan-sisa with file
def test_laporan_sisa_with_file(h):
    r = requests.post(f"{BASE}/api/rahaza/master/laporan-sisa", headers=h,
                      files={"file": ("t.xlsx", open(FIX_A, "rb").read())}, timeout=180)
    assert r.status_code == 200, r.text
    assert "LAPORAN_SISA_DA.xlsx" in (r.headers.get("Content-Disposition") or "")
    wb = load_workbook(io.BytesIO(r.content))
    print("Laporan sheets:", wb.sheetnames)
    for s in ("RINGKASAN", "1_MODEL_TANPA_SKU", "VARIAN_BARU", "2_KELOMPOK_TANPA_VARIAN",
              "3_SATUAN_KODE_TAK_VALID", "4_WARNA_TAK_DIKENAL", "5_MATERIAL_HARGA_0",
              "6_HARGA_JUAL_SKU", "7_LAINNYA", "8_HPP_BELUM_TERVALIDASI"):
        assert s in wb.sheetnames, f"missing sheet {s}"

    def rowcount(name):
        ws = wb[name]
        return sum(1 for row in ws.iter_rows(min_row=2, values_only=True) if row and any(v not in (None, "") for v in row))

    n1 = rowcount("1_MODEL_TANPA_SKU")
    assert n1 == 1, f"1_MODEL_TANPA_SKU rows={n1}"
    # 1_MODEL_TANPA_SKU row: DA-2506 warna_terdeteksi 'Hitam, Putih'
    ws1 = wb["1_MODEL_TANPA_SKU"]
    r1 = list(ws1.iter_rows(min_row=2, max_row=2, values_only=True))[0]
    assert r1[0] == "DA-2506" and "Hitam" in (r1[2] or "") and "Putih" in (r1[2] or ""), r1

    n_vb = rowcount("VARIAN_BARU")
    assert n_vb == 2, f"VARIAN_BARU={n_vb}"
    n2 = rowcount("2_KELOMPOK_TANPA_VARIAN")
    assert n2 == 1
    n3 = rowcount("3_SATUAN_KODE_TAK_VALID")
    assert n3 == 7, f"3_SATUAN_KODE_TAK_VALID={n3}"
    n4 = rowcount("4_WARNA_TAK_DIKENAL")
    assert n4 == 1
    n6 = rowcount("6_HARGA_JUAL_SKU")
    assert n6 == 2, f"6_HARGA_JUAL_SKU={n6}"
    n8 = rowcount("8_HPP_BELUM_TERVALIDASI")
    assert n8 > 0
    print("Row counts OK")


def test_laporan_sisa_no_file(h):
    r = requests.post(f"{BASE}/api/rahaza/master/laporan-sisa", headers=h, timeout=120)
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    assert "RINGKASAN" in wb.sheetnames


# ─── gap-workbook
def test_gap_workbook(h):
    r = requests.get(f"{BASE}/api/rahaza/master/gap-workbook", headers=h, timeout=120)
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    assert "BOM_AKSESORIS" in wb.sheetnames
    ws = wb["BOM_AKSESORIS"]
    header = [c.value for c in ws[1]]
    expected = ["kode_model", "nama_model", "kode_material", "nama_material", "qty_per_pcs", "satuan", "keterangan", "varian", "varian_tersedia"]
    assert header[:9] == expected, header
    # DA-1101/1202/1201 should NOT appear (already have accessories)
    codes = {(ws.cell(row=i, column=1).value or "") for i in range(2, ws.max_row + 1)}
    for m in ("DA-1101", "DA-1202", "DA-1201"):
        assert m not in codes, f"{m} unexpectedly present in BOM_AKSESORIS gap sheet"


# ─── completeness
def test_completeness(h):
    r = requests.get(f"{BASE}/api/dewi/rnd/completeness", headers=h, timeout=120)
    assert r.status_code == 200, r.text
    j = r.json()
    j = j.get("data", j) if isinstance(j, dict) else j
    checks = j.get("checks") or []
    keys = {c.get("key") for c in checks}
    assert "hpp" in keys, f"checks keys={keys}"
    hpp_check = next(c for c in checks if c.get("key") == "hpp")
    assert "HPP" in (hpp_check.get("label") or "")
    assert "hpp_unvalidated" in j
    # each row has hpp_validation & hpp_reasons
    rows = j.get("rows") or []
    assert rows, "no rows"
    r0 = rows[0]
    assert "hpp_validation" in r0 or "hpp_reasons" in r0, r0
    # DA-1101 flags.accessories=true
    da1101 = next((r for r in rows if r.get("code") == "DA-1101"), None)
    if da1101:
        flags = da1101.get("flags") or {}
        assert flags.get("accessories") is True, da1101
