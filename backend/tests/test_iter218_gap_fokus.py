"""Iter218 — Berkas FOKUS (DATA_YANG_PERLU_DIISI_DA_FOKUS.xlsx): hanya BOM_AKSESORIS · VARIAN_BARU · MATERIAL,
varian disarankan dari nama bahan pembeda (sel biru), klien mengisi sel kuning; MATERIAL isi_per_satuan_beli =
pcs per kemasan dasar (harga tetap per kemasan) dan baris BOM 'pcs' material itu ikut masuk pada unggahan yang sama.
DB dipulihkan ke seed go-live di akhir modul (scripts/seed_golive_restore.sh --force)."""
import io
import os
import subprocess
import time

import openpyxl
import pytest
import requests
from pymongo import MongoClient

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "https://dahost-dev-1.preview.emergentagent.com").rstrip("/")
SISA_XLSX = "/app/private/golive/DATA_YANG_PERLU_DIISI_DA_SISA_upload.xlsx"
MONGO_URL = os.environ.get("MONGO_URL", "mongodb://localhost:27017")
DB_NAME = os.environ.get("DB_NAME", "test_database")
YELLOW, BLUE, GREY = "00FFF2CC", "00DDEBF7", "00EDEDED"


def _restore_seed():
    subprocess.run(["bash", "/app/scripts/seed_golive_restore.sh", "--force"], check=True, capture_output=True, timeout=300)
    for _ in range(40):
        try:
            if requests.get(f"{BASE_URL}/api/health", timeout=5).status_code == 200:
                return
        except requests.RequestException:
            pass
        time.sleep(2)


@pytest.fixture(scope="module", autouse=True)
def _clean_db_after():
    yield
    _restore_seed()


@pytest.fixture(scope="module")
def headers():
    for _ in range(3):
        r = requests.post(f"{BASE_URL}/api/auth/login", json={"email": "admin@garment.com", "password": "Admin@123"}, timeout=30)
        if r.status_code == 200:
            return {"Authorization": f"Bearer {r.json()['token']}"}
        time.sleep(3)
    pytest.fail(f"login gagal: {r.status_code} {r.text}")


@pytest.fixture(scope="module")
def db():
    c = MongoClient(MONGO_URL)
    yield c[DB_NAME]
    c.close()


def _post_xlsx(headers, path, file_path=None):
    files = {"file": (os.path.basename(file_path), open(file_path, "rb"))} if file_path else None
    return requests.post(f"{BASE_URL}{path}", headers=headers, files=files, timeout=180)


@pytest.fixture(scope="module")
def fokus_wb(headers):
    r = _post_xlsx(headers, "/api/rahaza/master/gap-fokus", SISA_XLSX)
    assert r.status_code == 200, r.text
    assert "DATA_YANG_PERLU_DIISI_DA_FOKUS.xlsx" in r.headers.get("content-disposition", "")
    return openpyxl.load_workbook(io.BytesIO(r.content))


def _fill_rgb(cell):
    return cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type == "solid" else None


# ---------- 1) Struktur berkas FOKUS ----------
def test_fokus_sheets_only_needed(fokus_wb):
    assert fokus_wb.sheetnames == ["PETUNJUK", "VARIAN_BARU", "BOM_AKSESORIS", "MATERIAL", "REF_AKSESORIS"]
    for absent in ("STOK_AWAL_FG", "SALDO_AWAL", "HARGA_JUAL_SKU", "GAJI", "TOKO", "REKENING", "MODEL"):
        assert absent not in fokus_wb.sheetnames


def test_fokus_bom_columns_and_hints(fokus_wb):
    ws = fokus_wb["BOM_AKSESORIS"]
    hdr = [c.value for c in ws[1]]
    assert hdr[:8] == ["kode_model", "nama_model", "kode_material", "nama_material", "qty_per_pcs", "satuan", "keterangan", "varian"]
    assert hdr[-1] == "yang_perlu_diisi"
    rows = list(ws.iter_rows(min_row=2))
    assert len(rows) >= 400
    heads = [r for r in rows if r[0].value]
    assert all(r[9].value for r in heads), "setiap baris kepala kelompok harus punya petunjuk"
    auto = [r for r in heads if _fill_rgb(r[7]) == BLUE]
    manual = [r for r in heads if _fill_rgb(r[7]) == YELLOW]
    assert len(auto) >= 50, f"saran varian otomatis terlalu sedikit: {len(auto)}"
    assert all(r[7].value for r in auto), "sel biru harus terisi"
    assert all("otomatis" in r[9].value for r in auto)
    assert len(manual) >= 40 and all("varian" in r[9].value for r in manual)
    # bahan pembeda 'Kancing … Butter Yellow' → varian BUTTER YELLOW pada DA-1401
    da1401 = {r[7].value for r in auto if r[0].value == "DA-1401"}
    assert "BUTTER YELLOW" in da1401 and "DUSTY" in da1401 and "MAROON" in da1401, da1401
    # bahan pembeda yang warnanya bukan varian model (Kancing Hitam di model tanpa warna Hitam) TIDAK ditebak
    assert not any(r[7].value == "HITAM" for r in auto if r[0].value in ("DA-1401", "DA-3401"))


def test_fokus_varian_baru_and_material(fokus_wb):
    ws = fokus_wb["VARIAN_BARU"]
    rows = list(ws.iter_rows(min_row=2))
    assert len(rows) == 38
    hdr = [c.value for c in ws[1]]
    uk = hdr.index("ukuran")
    assert all(_fill_rgb(r[uk]) == YELLOW for r in rows)
    lunara = [r for r in rows if r[0].value == "DA-1209"]
    assert len(lunara) == 7 and all(r[3].value and _fill_rgb(r[3]) == BLUE for r in lunara)
    ws = fokus_wb["MATERIAL"]
    hdr = [c.value for c in ws[1]]
    by = {r[0].value: r for r in ws.iter_rows(min_row=2)}
    assert 80 <= len(by) <= 120
    bab = by["A-BAB-0001"]
    assert bab[hdr.index("satuan_beli")].value == "roll" and _fill_rgb(bab[hdr.index("isi_per_satuan_beli")]) == YELLOW
    assert bab[hdr.index("harga_per_satuan_beli")].value == 120000 and _fill_rgb(bab[hdr.index("harga_per_satuan_beli")]) == BLUE
    b20 = by["A-B20-0002"]
    assert b20[hdr.index("isi_per_satuan_beli")].value == 1 and _fill_rgb(b20[hdr.index("harga_per_satuan_beli")]) == YELLOW


# ---------- 2) Unggah balik apa adanya → saran varian langsung terpakai, tanpa error ----------
def test_fokus_reupload_preview(headers, fokus_wb, tmp_path):
    p = tmp_path / "fokus.xlsx"
    fokus_wb.save(p)
    r = _post_xlsx(headers, "/api/rahaza/master/fill-preview", str(p))
    assert r.status_code == 200, r.text
    j = r.json()
    assert j["ok"] is True and j["errors"] == []
    assert j["totals"]["bom_groups"] >= 60, j["totals"]
    assert j["bom_issue_counts"].get("kelompok_tanpa_varian", 0) < 50, j["bom_issue_counts"]


# ---------- 3) Klien mengisi sel kuning → terapkan (idempoten) ----------
@pytest.fixture(scope="module")
def filled_path(fokus_wb, tmp_path_factory):
    wb = fokus_wb
    ws = wb["MATERIAL"]
    hdr = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2):
        if r[0].value == "A-BAB-0001":
            r[hdr.index("isi_per_satuan_beli")].value = 50
        if r[0].value == "A-B20-0002":
            r[hdr.index("harga_per_satuan_beli")].value = 15000
    ws = wb["VARIAN_BARU"]
    hdr = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2):
        if r[0].value == "DA-1209":
            r[hdr.index("ukuran")].value = "ALLSIZE"
    ws = wb["BOM_AKSESORIS"]
    for r in ws.iter_rows(min_row=2):
        if r[2].value == "A-K34-0008" and not r[4].value:
            r[4].value = "1 pcs"
    p = tmp_path_factory.mktemp("fokus") / "FOKUS_terisi.xlsx"
    wb.save(p)
    return str(p)


def test_apply_filled_then_idempotent(headers, db, filled_path):
    r1 = _post_xlsx(headers, "/api/rahaza/master/fill-apply?scope=all", filled_path)
    assert r1.status_code == 200, r1.text
    j = r1.json()
    assert j["variants_created"] == 7 and "DA-1209-HTM-ALLSIZE" in j["variants_created_skus"]
    assert j["bom_lines_appended"] > 200 and j["boms_touched"] > 50
    m = db.rahaza_materials.find_one({"code": "A-BAB-0001"}, {"_id": 0, "unit_cost": 1, "pack_size": 1, "uoms": 1})
    assert m["pack_size"] == 50 and m["unit_cost"] == 120000, m  # harga tetap per roll, isi = pcs per roll
    assert any(u["code"] == "pcs" and abs(u["factor"] - 0.02) < 1e-9 for u in m["uoms"])
    assert db.rahaza_materials.find_one({"code": "A-B20-0002"})["unit_cost"] == 15000
    # baris BOM 'pcs' Babud ikut masuk pada unggahan yang sama
    mid = db.rahaza_models.find_one({"code": "DA-2105"})["id"]
    assert db.rahaza_boms.count_documents({"model_id": mid, "is_active": True, "active": {"$ne": False}, "materials.code": "A-BAB-0001"}) > 0
    # saran varian → Kancing Butter Yellow hanya di BOM BTY, bukan di DST
    mid = db.rahaza_models.find_one({"code": "DA-1401"})["id"]
    bty = db.rahaza_boms.find_one({"model_id": mid, "is_active": True, "color_code": "BTY"})
    dst = db.rahaza_boms.find_one({"model_id": mid, "is_active": True, "color_code": "DST"})
    assert any("Butter Yellow" in x["name"] for x in bty["materials"]) and not any("Butter Yellow" in x["name"] for x in dst["materials"])
    r2 = _post_xlsx(headers, "/api/rahaza/master/fill-apply?scope=all", filled_path)
    j2 = r2.json()
    assert j2["variants_created"] == 0 and j2["bom_lines_appended"] == 0 and j2["boms_touched"] == 0, j2


def test_fokus_after_apply_shrinks(headers, filled_path):
    r = _post_xlsx(headers, "/api/rahaza/master/gap-fokus", filled_path)
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb["VARIAN_BARU"].max_row - 1 == 31  # 38 − 7 Lunara
    assert wb["BOM_AKSESORIS"].max_row - 1 < 300
    assert wb["MATERIAL"].max_row - 1 == 86


def test_fokus_without_file(headers):
    r = _post_xlsx(headers, "/api/rahaza/master/gap-fokus")
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["PETUNJUK", "VARIAN_BARU", "BOM_AKSESORIS", "BOM_OTOMATIS", "MATERIAL", "REF_AKSESORIS"]
    assert wb["VARIAN_BARU"].max_row - 1 >= 1  # model tanpa SKU dari DB


def test_fokus_requires_auth():
    assert requests.post(f"{BASE_URL}/api/rahaza/master/gap-fokus", timeout=30).status_code in (401, 403)
