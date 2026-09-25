"""Iter216 — POST gap-workbook (dgn berkas klien) + GET gap-workbook (tanpa berkas) +
Importir VARIAN_BARU end-to-end (fill-preview → fill-apply?scope=bom idempoten) + DB verifikasi
+ regression preview turunnya model_tanpa_sku setelah DA-1209 punya SKU.

Order-sensitive: gunakan pytest -x agar berhenti pada kegagalan.
"""
import io
import os
import pytest
import requests
from openpyxl import load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")
GAP_FILE = "/app/private/golive/DATA_YANG_PERLU_DIISI_DA_2.xlsx"
VARIAN_FILE = "/app/private/golive/uji_varian_baru_DA-1209.xlsx"
ADMIN_EMAIL = "admin@garment.com"
ADMIN_PASS = "Admin@123"

# module-scope shared state
_state = {}


@pytest.fixture(scope="module")
def hdr():
    r = requests.post(f"{BASE_URL}/api/auth/login",
                      json={"email": ADMIN_EMAIL, "password": ADMIN_PASS}, timeout=30)
    assert r.status_code == 200, f"login failed {r.status_code} {r.text[:300]}"
    tok = r.json().get("access_token") or r.json().get("token")
    assert tok, f"no token in {r.json()}"
    return {"Authorization": f"Bearer {tok}"}


def _upload(url, hdr, path, extra=None):
    with open(path, "rb") as f:
        data = f.read()
    files = {"file": (os.path.basename(path), io.BytesIO(data),
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    return requests.post(url, headers=hdr, files=files, data=extra or {}, timeout=180)


# ---------- (A) POST /gap-workbook dengan berkas klien ----------
def test_post_gap_workbook_with_client_file(hdr):
    r = _upload(f"{BASE_URL}/api/rahaza/master/gap-workbook", hdr, GAP_FILE)
    assert r.status_code == 200, f"status {r.status_code} {r.text[:300]}"
    cd = r.headers.get("Content-Disposition", "")
    assert "DATA_YANG_PERLU_DIISI_DA_SISA.xlsx" in cd, f"content-disposition={cd}"
    wb = load_workbook(io.BytesIO(r.content), read_only=True)
    names = wb.sheetnames
    for s in ("PETUNJUK", "BOM_AKSESORIS", "VARIAN_BARU", "MATERIAL"):
        assert s in names, f"missing sheet {s}; got {names}"

    # PETUNJUK: contains DAMPAK + DITERAPKAN row
    ws = wb["PETUNJUK"]
    text_all = "\n".join(
        " | ".join(str(c) if c is not None else "" for c in row)
        for row in ws.iter_rows(values_only=True)
    )
    assert "DAMPAK BERKAS SEBELUMNYA" in text_all, "no DAMPAK BERKAS SEBELUMNYA in PETUNJUK"
    assert "DITERAPKAN" in text_all and "340 baris bahan" in text_all and \
           "89 kelompok" in text_all and "62 model" in text_all, \
        f"expected DITERAPKAN 340/89/62 not found. Excerpt:\n{text_all[:2000]}"

    # BOM_AKSESORIS: has 'masalah' & 'tindakan' cols; first data row DA-1101 with 'maroon' problem
    ws = wb["BOM_AKSESORIS"]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(c).strip().lower() if c else "" for c in rows[0]]
    assert "masalah" in header, f"masalah col missing in {header}"
    assert "tindakan" in header, f"tindakan col missing in {header}"
    idx_masalah = header.index("masalah")
    # find model col
    idx_model = next((i for i, h in enumerate(header) if "model" in h and "kode" in h), None)
    if idx_model is None:
        idx_model = next((i for i, h in enumerate(header) if "model" in h), 0)
    # first data row
    first_data = rows[1]
    assert str(first_data[idx_model]).strip() == "DA-1101", \
        f"first row model={first_data[idx_model]}"
    assert "maroon" in str(first_data[idx_masalah]).lower(), \
        f"first row masalah={first_data[idx_masalah]}"

    # VARIAN_BARU: 38 data rows, cols check + DA-1209 Hitam→HTM
    ws = wb["VARIAN_BARU"]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(c).strip().lower() if c else "" for c in rows[0]]
    expected_cols = ["kode_model", "nama_model", "warna", "kode_warna", "ukuran",
                     "harga_jual", "keterangan", "kode_warna_tersedia", "ukuran_tersedia"]
    for c in expected_cols:
        assert c in header, f"VARIAN_BARU missing col {c}; got {header}"
    data_rows = rows[1:]
    # filter empties
    data_rows = [r for r in data_rows if any(v not in (None, "") for v in r)]
    assert len(data_rows) == 38, f"VARIAN_BARU data_rows={len(data_rows)}"
    idx_km = header.index("kode_model")
    idx_w = header.index("warna")
    idx_kw = header.index("kode_warna")
    da1209_hitam = next((r for r in data_rows
                         if str(r[idx_km]).strip() == "DA-1209"
                         and str(r[idx_w]).strip().lower() == "hitam"), None)
    assert da1209_hitam, "DA-1209 Hitam row missing in VARIAN_BARU"
    assert str(da1209_hitam[idx_kw]).strip().upper() == "HTM", \
        f"DA-1209 Hitam kode_warna={da1209_hitam[idx_kw]}"

    # MATERIAL: A-BAB-0001 with keterangan containing 'isi kemasan'
    ws = wb["MATERIAL"]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(c).strip().lower() if c else "" for c in rows[0]]
    idx_code = next((i for i, h in enumerate(header) if "kode" in h), 0)
    idx_ket = next((i for i, h in enumerate(header) if "keterangan" in h), -1)
    assert idx_ket >= 0, f"no keterangan col in MATERIAL header {header}"
    babrow = next((r for r in rows[1:]
                   if r and str(r[idx_code]).strip() == "A-BAB-0001"), None)
    assert babrow, "A-BAB-0001 missing in MATERIAL"
    assert "isi kemasan" in str(babrow[idx_ket]).lower(), \
        f"A-BAB-0001 keterangan={babrow[idx_ket]}"


# ---------- (B) GET /gap-workbook (tanpa file) ----------
def test_get_gap_workbook_no_file(hdr):
    r = requests.get(f"{BASE_URL}/api/rahaza/master/gap-workbook", headers=hdr, timeout=60)
    assert r.status_code == 200, f"status {r.status_code}"
    cd = r.headers.get("Content-Disposition", "")
    assert "DATA_YANG_PERLU_DIISI_DA.xlsx" in cd, f"cd={cd}"
    # must NOT be DA_SISA filename
    assert "DA_SISA" not in cd, f"unexpected DA_SISA in {cd}"
    wb = load_workbook(io.BytesIO(r.content), read_only=True)
    assert "VARIAN_BARU" not in wb.sheetnames, f"VARIAN_BARU should be absent; got {wb.sheetnames}"


# ---------- (C) Importir VARIAN_BARU end-to-end ----------
def test_variant_import_preview(hdr):
    r = _upload(f"{BASE_URL}/api/rahaza/master/fill-preview", hdr, VARIAN_FILE)
    assert r.status_code == 200, f"preview status {r.status_code} {r.text[:400]}"
    j = r.json()
    totals = j.get("totals", {})
    assert totals.get("new_variants") == 7, f"new_variants={totals.get('new_variants')} full={totals}"
    assert totals.get("bom_lines") == 0, f"bom_lines={totals.get('bom_lines')} full={totals}"
    assert j.get("errors") in ([], None), f"errors={j.get('errors')}"


def test_variant_import_apply_bom_first_run(hdr):
    r = _upload(f"{BASE_URL}/api/rahaza/master/fill-apply?scope=bom", hdr, VARIAN_FILE)
    assert r.status_code == 200, f"apply status {r.status_code} {r.text[:400]}"
    j = r.json()
    assert j.get("variants_created") == 7, f"variants_created={j.get('variants_created')} full={j}"
    skus = j.get("variants_created_skus") or []
    assert "DA-1209-HTM-ALLSIZE" in skus, f"HTM-ALLSIZE not in {skus}"
    assert j.get("bom_models") == 1, f"bom_models={j.get('bom_models')}"
    assert j.get("bom_base_created") == 7, f"bom_base_created={j.get('bom_base_created')}"
    assert j.get("bom_lines_appended") == 28, f"bom_lines_appended={j.get('bom_lines_appended')} full={j}"


def test_variant_import_apply_idempotent(hdr):
    r = _upload(f"{BASE_URL}/api/rahaza/master/fill-apply?scope=bom", hdr, VARIAN_FILE)
    assert r.status_code == 200, f"apply2 status {r.status_code} {r.text[:400]}"
    j = r.json()
    assert j.get("variants_created") == 0, f"variants_created={j.get('variants_created')} full={j}"
    assert j.get("boms_unchanged") == 7, f"boms_unchanged={j.get('boms_unchanged')} full={j}"
    assert j.get("bom_lines_appended") == 0, f"bom_lines_appended={j.get('bom_lines_appended')} full={j}"


# ---------- (D) DB verification via mongo ----------
def test_db_variants_and_materials_and_boms():
    from pymongo import MongoClient
    mongo_url = os.environ.get("MONGO_URL")
    db_name = os.environ.get("DB_NAME", "test_database")
    # Load backend .env if needed
    if not mongo_url:
        try:
            with open("/app/backend/.env") as f:
                for ln in f:
                    if ln.startswith("MONGO_URL="):
                        mongo_url = ln.strip().split("=", 1)[1].strip('"').strip("'")
                    if ln.startswith("DB_NAME="):
                        db_name = ln.strip().split("=", 1)[1].strip('"').strip("'")
        except FileNotFoundError:
            pass
    assert mongo_url, "MONGO_URL missing"
    cli = MongoClient(mongo_url)
    db = cli[db_name]

    # variants: 7 for DA-1209
    n_var = db.rahaza_model_variants.count_documents({"model_code": "DA-1209"})
    assert n_var == 7, f"DA-1209 variants={n_var}"

    # materials: FG with code DA-1209-HTM-ALLSIZE and price 99000
    fg = db.rahaza_materials.find_one({"type": "fg", "code": "DA-1209-HTM-ALLSIZE"})
    assert fg, "FG DA-1209-HTM-ALLSIZE missing"
    assert fg.get("retail_price_master") == 99000, f"retail_price_master={fg.get('retail_price_master')}"

    # BOM DA-1209 HTM contains 4 accessories (BOMs use model_id, not model_code)
    mdoc = db.rahaza_models.find_one({"code": "DA-1209"})
    assert mdoc, "model DA-1209 missing"
    mid = mdoc.get("id") or mdoc.get("_id")
    boms = list(db.rahaza_boms.find({"model_id": mid, "color_code": "HTM"}))
    assert boms, f"no BOM for DA-1209 HTM (model_id={mid})"
    txt = str(boms)
    for code in ("A-KRT-0003", "A-LBL-0004", "A-R15-0001", "A-LBL-0007"):
        assert code in txt, f"{code} missing in DA-1209 HTM BOM"


# ---------- (E) Preview ulang berkas klien: model_tanpa_sku turun ----------
def test_reprocess_client_file_model_tanpa_sku_drops(hdr):
    r = _upload(f"{BASE_URL}/api/rahaza/master/fill-preview", hdr, GAP_FILE)
    assert r.status_code == 200, f"preview status {r.status_code} {r.text[:400]}"
    j = r.json()
    ic = j.get("bom_issue_counts", {})
    mts = ic.get("model_tanpa_sku")
    assert mts is not None, f"no model_tanpa_sku in {ic}"
    assert mts < 47, f"model_tanpa_sku={mts} expected <47"
    totals = j.get("totals", {})
    bm = totals.get("bom_models")
    assert bm is not None and bm >= 63, f"bom_models={bm} expected >=63"
