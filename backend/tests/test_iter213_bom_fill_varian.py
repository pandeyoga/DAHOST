"""Iteration 213 — BOM_AKSESORIS varian importer (bom_fill.py)"""
import pytest as _pytest
_pytest.skip("Uji importir BOM v2 (tebakan warna, pcs→roll tanpa isi kemasan) — digantikan v3: test_iter214_bom_fill_v3.py", allow_module_level=True)

import io
import os
import pytest
import requests
from openpyxl import load_workbook
from motor.motor_asyncio import AsyncIOMotorClient
import asyncio

BASE = (os.environ.get("REACT_APP_BACKEND_URL") or open("/app/frontend/.env").read().split("REACT_APP_BACKEND_URL=")[1].split("\n")[0]).rstrip("/")
FIXTURE = "/app/tests/fixtures/bom_aksesoris_varian_test.xlsx"
FIXTURE_REAL = "/app/private/golive/DATA_YANG_PERLU_DIISI_DA_isi_bom.xlsx"

# ─── auth
@pytest.fixture(scope="module")
def token():
    r = requests.post(f"{BASE}/api/auth/login", json={"email": "admin@garment.com", "password": "Admin@123"})
    assert r.status_code == 200, r.text
    return r.json().get("token") or r.json().get("access_token") or r.json()["data"]["token"]

@pytest.fixture(scope="module")
def h(token):
    return {"Authorization": f"Bearer {token}"}


def _unwrap(j):
    return j["data"] if isinstance(j, dict) and "data" in j and "ok" in j else j


# ─── preview
@pytest.fixture(scope="module")
def preview(h):
    with open(FIXTURE, "rb") as f:
        r = requests.post(f"{BASE}/api/rahaza/master/fill-preview", headers=h,
                          files={"file": ("test.xlsx", f.read(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    j = r.json()
    assert j.get("ok") is True, j
    return _unwrap(j)


def test_preview_totals(preview):
    t = preview["totals"]
    print("TOTALS:", t)
    print("WARNINGS:", preview.get("warnings"))
    assert preview.get("errors") == []
    assert t["bom_lines"] == 9, t
    assert t["bom_groups"] == 4, t
    assert t["bom_models"] == 2, t
    assert t["sku_prices"] == 1, t
    assert t["skipped"] == 7, t
    assert len(preview.get("warnings", [])) == 7, preview.get("warnings")


def test_preview_bom_lines_details(preview):
    lines = preview["bom_lines"]
    assert len(lines) == 9
    # Row 2 — A-KRT-0003 60 cm → 0.6 m
    row2 = next(ln for ln in lines if ln["row"] == 2)
    assert row2["code"] == "A-KRT-0003"
    assert abs(row2["qty_base"] - 0.6) < 1e-4
    assert row2["unit_base"] == "m"
    # Row 3 — A-R45-0001 1 pcs → 0.016667 pack
    row3 = next(ln for ln in lines if ln["row"] == 3)
    assert row3["code"] == "A-R45-0001"
    assert abs(row3["qty_base"] - 0.016667) < 1e-4
    assert row3["unit_base"] == "pack"
    # Row 4 — A-LBL-0004 zero-width in code
    row4 = next(ln for ln in lines if ln["row"] == 4)
    assert row4["code"] == "A-LBL-0004"
    assert abs(row4["qty_base"] - 0.001667) < 1e-4
    assert row4["unit_base"] == "roll"


# ─── apply (first time)
@pytest.fixture(scope="module")
def apply_first(h):
    with open(FIXTURE, "rb") as f:
        r = requests.post(f"{BASE}/api/rahaza/master/fill-apply", headers=h,
                          files={"file": ("test.xlsx", f.read())})
    assert r.status_code == 200, r.text
    return _unwrap(r.json())


def test_apply_first(apply_first):
    print("APPLY1:", apply_first)
    assert apply_first["bom_groups"] == 4
    assert apply_first["bom_lines_appended"] > 0
    assert apply_first["boms_touched"] > 0
    assert apply_first["sku_prices_updated"] == 1
    assert apply_first["bom_pcs_uoms_added"] >= 1


def test_mongo_state(apply_first):
    async def check():
        c = AsyncIOMotorClient("mongodb://localhost:27017")
        db = c["test_database"]
        # DA-1101
        m1101 = await db.rahaza_models.find_one({"code": "DA-1101"}, {"_id": 0, "id": 1})
        assert m1101
        boms = await db.rahaza_boms.find({"model_id": m1101["id"], "is_active": True}, {"_id": 0}).to_list(100)
        by_color = {(b.get("color_code") or "").upper(): b for b in boms}
        for cc in ("HTM", "MHG"):
            assert cc in by_color, f"BOM {cc} missing"
            codes = {m.get("code") for m in by_color[cc]["materials"]}
            assert "A-R45-0001" in codes, f"{cc} missing A-R45-0001; got {codes}"
            assert "A-LBL-0004" in codes, f"{cc} missing A-LBL-0004"
            assert "A-KRT-0003" in codes, f"{cc} missing A-KRT-0003"
        for cc in ("CRL", "GRY", "MGT"):
            assert cc in by_color, f"BOM {cc} missing"
            codes = {m.get("code") for m in by_color[cc]["materials"]}
            assert "A-BIS-0001" in codes, f"{cc} missing A-BIS-0001"
            assert "A-KRT-0003" in codes, f"{cc} missing A-KRT-0003"
            assert "A-R45-0001" not in codes, f"{cc} unexpectedly has A-R45-0001"
        # BRG untouched (should not have any of our new mats OR should be same as pre)
        if "BRG" in by_color:
            codes = {m.get("code") for m in by_color["BRG"]["materials"]}
            # BRG shouldn't have accessories only added to CRL/GRY/MGT
            assert "A-BIS-0001" not in codes, "BRG should not have A-BIS-0001"

        # DA-1202
        m1202 = await db.rahaza_models.find_one({"code": "DA-1202"}, {"_id": 0, "id": 1})
        assert m1202
        boms2 = await db.rahaza_boms.find({"model_id": m1202["id"], "is_active": True}, {"_id": 0}).to_list(100)
        by_color2 = {(b.get("color_code") or "").upper(): b for b in boms2}
        htm_codes = {m.get("code") for m in by_color2.get("HTM", {}).get("materials", [])}
        mgt_codes = {m.get("code") for m in by_color2.get("MGT", {}).get("materials", [])}
        assert "A-K22-0014" in htm_codes, f"HTM missing A-K22-0014; got {htm_codes}"
        assert "A-K22-0005" in mgt_codes, f"MGT missing A-K22-0005; got {mgt_codes}"
        assert "A-K22-0014" not in mgt_codes
        assert "A-K22-0005" not in htm_codes
        # A-K22-0001 kancing bening — should not be in any BOM
        for cc, b in by_color2.items():
            codes = {m.get("code") for m in b["materials"]}
            assert "A-K22-0001" not in codes, f"BOM {cc} unexpectedly has A-K22-0001"

        # A-KRT-0003 qty_base/unit_base
        for cc in ("HTM", "MHG"):
            mat_line = next((m for m in by_color[cc]["materials"] if m.get("code") == "A-KRT-0003"), None)
            assert mat_line, f"A-KRT-0003 missing in {cc}"
            assert abs(float(mat_line.get("qty_base") or 0) - 0.6) < 1e-3, mat_line
            assert (mat_line.get("unit_base") or "").lower() == "m"

        # A-LBL-0004 uoms entry
        lbl = await db.rahaza_materials.find_one({"code": "A-LBL-0004"}, {"_id": 0})
        assert lbl
        pcs = [u for u in (lbl.get("uoms") or []) if (u.get("code") or "").lower() == "pcs"]
        assert pcs, f"A-LBL-0004 missing pcs uom; uoms={lbl.get('uoms')}"
        assert abs(float(pcs[0]["factor"]) - 0.00166667) < 1e-4

        # FG retail price
        fg = await db.rahaza_materials.find_one({"code": "DA-1202-CRM-M"}, {"_id": 0})
        assert fg, "FG DA-1202-CRM-M not found"
        assert float(fg.get("retail_price_master") or 0) == 155999, fg.get("retail_price_master")

        c.close()
    asyncio.run(check())


# ─── apply idempotency
def test_apply_idempotent(h, apply_first):
    with open(FIXTURE, "rb") as f:
        r = requests.post(f"{BASE}/api/rahaza/master/fill-apply", headers=h, files={"file": ("t.xlsx", f.read())})
    assert r.status_code == 200
    j = _unwrap(r.json())
    print("APPLY2:", j)
    assert j["bom_lines_appended"] == 0
    assert j["bom_lines_updated"] == 0
    assert j["boms_touched"] == 0
    assert j["sku_prices_updated"] == 0
    assert j["bom_pcs_uoms_added"] == 0


# ─── gap workbook + template headers
def _check_bom_header(wb):
    ws = wb["BOM_AKSESORIS"]
    header = [c.value for c in ws[1]]
    expected = ["kode_model", "nama_model", "kode_material", "nama_material", "qty_per_pcs", "satuan", "keterangan", "varian"]
    assert header[:8] == expected, f"header={header}"
    # varian_tersedia can be 9th
    assert len(header) >= 9 and header[8] == "varian_tersedia", f"header={header}"


def test_gap_workbook(h):
    r = requests.get(f"{BASE}/api/rahaza/master/gap-workbook", headers=h)
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    assert "BOM_AKSESORIS" in wb.sheetnames
    _check_bom_header(wb)
    # varian_tersedia rows contain 'Warna:' text
    ws = wb["BOM_AKSESORIS"]
    sample = [ws.cell(row=i, column=9).value for i in range(2, min(ws.max_row + 1, 20))]
    print("gap varian_tersedia sample:", sample[:5])
    assert any(v and "Warna:" in str(v) for v in sample if v), sample
    # MODEL sheet
    if "MODEL" in wb.sheetnames:
        mh = [c.value for c in wb["MODEL"][1]]
        print("MODEL header:", mh)
        assert mh[-1] == "varian_tersedia", mh


def test_fill_template(h):
    r = requests.get(f"{BASE}/api/rahaza/master/fill-template", headers=h)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    _check_bom_header(wb)


# ─── regression on real client file (preview only)
def test_preview_real_file(h):
    if not os.path.exists(FIXTURE_REAL):
        pytest.skip("real file missing")
    with open(FIXTURE_REAL, "rb") as f:
        r = requests.post(f"{BASE}/api/rahaza/master/fill-preview", headers=h,
                          files={"file": ("real.xlsx", f.read())}, timeout=120)
    assert r.status_code == 200, r.text
    j = r.json()
    assert j.get("ok") is True, j
    t = _unwrap(j)["totals"]
    print("REAL totals:", t)
    assert t["bom_lines"] == 515, t
    assert t["bom_groups"] == 148, t
    assert t["sku_prices"] == 229, t
    assert t["skipped"] == 150, t
