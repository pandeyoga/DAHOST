"""QA — Portal Keuangan: laporan 12 bulan · neraca lajur · saldo awal go-live · arus kas per akun · COA cash_flow_group.

Semua request memakai satu token superadmin. Jurnal uji dihapus di akhir (tearDown per test).
"""
from __future__ import annotations

import io
import os
import time
from datetime import date

import openpyxl
import pytest
import requests


def _load_frontend_env():
    try:
        with open("/app/frontend/.env") as f:
            for line in f:
                if line.startswith("REACT_APP_BACKEND_URL="):
                    return line.split("=", 1)[1].strip().strip('"').strip("'")
    except FileNotFoundError:
        pass
    return None


_URL = os.environ.get("REACT_APP_BACKEND_URL") or _load_frontend_env() or ""
BASE = _URL.rstrip("/") + "/api"
assert _URL, "REACT_APP_BACKEND_URL kosong"


def _login() -> str:
    for _ in range(3):
        r = requests.post(f"{BASE}/auth/login", json={"email": "admin@garment.com", "password": "Admin@123"}, timeout=30)
        if r.status_code == 200:
            return r.json()["token"]
        time.sleep(2)
    pytest.skip(f"login gagal {r.status_code} {r.text[:200]}")


@pytest.fixture(scope="module")
def token() -> str:
    return _login()


@pytest.fixture(scope="module")
def H(token) -> dict:
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture(scope="module")
def db():
    from pymongo import MongoClient
    cli = MongoClient(os.environ.get("MONGO_URL") or "mongodb://localhost:27017")
    return cli[os.environ.get("DB_NAME") or "test_database"]


YEAR = 2026


# ═══════════ 12 BULAN LAPORAN ═══════════
class TestReports12M:
    def test_pnl_monthly_ok(self, H):
        r = requests.get(f"{BASE}/rahaza/finance/reports/profit-loss-monthly", params={"year": YEAR}, headers=H, timeout=30)
        assert r.status_code == 200, r.text
        j = r.json()
        assert len(j["meta"]["months"]) == 12
        for k in ("revenue", "cogs", "expense", "other_income", "other_expense"):
            assert k in j["groups"]
        assert len(j["lines"]["net_income"]["months"]) == 12
        assert "net_income_cumulative" in j["lines"]

    def test_pnl_monthly_bad_year(self, H):
        assert requests.get(f"{BASE}/rahaza/finance/reports/profit-loss-monthly", params={"year": "abc"}, headers=H, timeout=15).status_code == 422
        assert requests.get(f"{BASE}/rahaza/finance/reports/profit-loss-monthly", params={"year": 1900}, headers=H, timeout=15).status_code == 422

    def test_bs_monthly_ok(self, H):
        r = requests.get(f"{BASE}/rahaza/finance/reports/balance-sheet-monthly", params={"year": YEAR}, headers=H, timeout=30)
        assert r.status_code == 200, r.text
        j = r.json()
        for k in ("assets", "liabilities", "equity"):
            assert k in j["sections"]
        assert len(j["totals"]["diff"]) == 12
        assert isinstance(j["balanced"], bool)


# ═══════════ NERACA LAJUR ═══════════
class TestWorksheet:
    def test_ok(self, H):
        r = requests.get(f"{BASE}/rahaza/finance/reports/worksheet", params={"from": f"{YEAR}-01-01", "to": f"{YEAR}-12-31"}, headers=H, timeout=30)
        assert r.status_code == 200, r.text
        j = r.json()
        assert "rows" in j and "totals" in j and "grand_totals" in j and "balanced" in j
        assert set(j["balancing"].keys()) >= {"label", "net_income"}
        for k in ("pl_debit", "pl_credit", "bs_debit", "bs_credit"):
            assert k in j["totals"]


# ═══════════ EXPORT XLSX ═══════════
class TestExportXlsx:
    @pytest.mark.parametrize("params,fname", [
        ({"report": "profit-loss-monthly", "year": YEAR}, "laba"),
        ({"report": "balance-sheet-monthly", "year": YEAR}, "neraca"),
        ({"report": "worksheet", "from": f"{YEAR}-01-01", "to": f"{YEAR}-12-31"}, "lajur"),
        ({"report": "cash-flow", "from": f"{YEAR}-01-01", "to": f"{YEAR}-12-31"}, "arus"),
    ])
    def test_export_ok(self, H, params, fname):
        r = requests.get(f"{BASE}/rahaza/finance/reports/export-xlsx", params=params, headers=H, timeout=60)
        assert r.status_code == 200, r.text[:300]
        assert "spreadsheet" in r.headers.get("content-type", "")
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        assert len(wb.sheetnames) >= 1

    def test_export_bad_report(self, H):
        r = requests.get(f"{BASE}/rahaza/finance/reports/export-xlsx", params={"report": "xyz"}, headers=H, timeout=15)
        assert r.status_code == 422


# ═══════════ CASH FLOW ═══════════
class TestCashFlow:
    def test_ok(self, H):
        r = requests.get(f"{BASE}/rahaza/finance/reports/cash-flow", params={"from": f"{YEAR}-01-01", "to": f"{YEAR}-12-31"}, headers=H, timeout=30)
        assert r.status_code == 200, r.text
        j = r.json()
        assert "activities" in j
        ba = j["by_account"]
        assert "1-1101" in ba["cash_account_codes"]
        for g in ("OPR", "INV", "PND"):
            assert g in ba["groups"]
        for k in ("opening_cash_gl", "closing_cash_gl", "net_change_in_cash", "cash_gl_net_change", "unexplained"):
            assert k in ba["totals"]


# ═══════════ COA cash_flow_group ═══════════
class TestCOA:
    def _get_acc(self, H, code):
        r = requests.get(f"{BASE}/rahaza/coa/accounts", headers=H, timeout=30)
        assert r.status_code == 200, r.text
        data = r.json()
        accs = data if isinstance(data, list) else data.get("accounts", [])
        for a in accs:
            if a.get("code") == code:
                return a
        return None

    def test_all_active_have_cfg(self, H):
        r = requests.get(f"{BASE}/rahaza/coa/accounts", headers=H, timeout=30)
        assert r.status_code == 200
        data = r.json()
        accs = data.get("accounts") if isinstance(data, dict) else data
        missing = [a["code"] for a in accs if a.get("active", True) and not a.get("is_group") and a.get("cash_flow_group") not in ("OPR", "INV", "PND")]
        assert not missing, f"akun aktif tanpa cash_flow_group: {missing[:10]}"

    def test_invalid_cfg(self, H):
        acc = self._get_acc(H, "4-1100")
        if not acc:
            pytest.skip("akun 4-1100 tidak ada")
        r = requests.put(f"{BASE}/rahaza/coa/accounts/{acc['id']}", json={"cash_flow_group": "XXX"}, headers=H, timeout=15)
        assert r.status_code == 400, r.text

    def test_update_cfg_and_revert(self, H):
        acc = self._get_acc(H, "4-1100")
        if not acc:
            pytest.skip("akun 4-1100 tidak ada")
        orig = acc.get("cash_flow_group") or "OPR"
        r = requests.put(f"{BASE}/rahaza/coa/accounts/{acc['id']}", json={"cash_flow_group": "INV"}, headers=H, timeout=15)
        assert r.status_code == 200, r.text
        got = self._get_acc(H, "4-1100")
        assert got["cash_flow_group"] == "INV"
        # revert
        rv = requests.put(f"{BASE}/rahaza/coa/accounts/{acc['id']}", json={"cash_flow_group": orig}, headers=H, timeout=15)
        assert rv.status_code == 200


# ═══════════ E2E JOURNAL + REPORTS ═══════════
class TestE2EJournal:
    JE_ID = None

    def test_flow(self, H, db):
        # buat jurnal Apr 2026
        payload = {"date": "2026-04-10", "memo": "QA-F48 test", "post": True,
                   "lines": [{"account_code": "1-1101", "debit": 500000, "credit": 0},
                             {"account_code": "4-1100", "debit": 0, "credit": 500000}]}
        r = requests.post(f"{BASE}/rahaza/journals", json=payload, headers=H, timeout=30)
        assert r.status_code in (200, 201), r.text
        je = r.json()
        je_id = je.get("id") or je.get("journal", {}).get("id")
        assert je_id
        TestE2EJournal.JE_ID = je_id
        try:
            # cek pnl apr
            pnl = requests.get(f"{BASE}/rahaza/finance/reports/profit-loss-monthly", params={"year": YEAR}, headers=H, timeout=30).json()
            rev = pnl["groups"]["revenue"]
            row = next((a for a in rev["accounts"] if a["code"] == "4-1100"), None)
            assert row and row["months"][3] >= 500000, f"apr revenue 4-1100 = {row and row['months'][3]}"
            # worksheet net_income increase
            ws = requests.get(f"{BASE}/rahaza/finance/reports/worksheet", params={"from": f"{YEAR}-01-01", "to": f"{YEAR}-12-31"}, headers=H, timeout=30).json()
            assert ws["balancing"]["net_income"] >= 500000
            # cash flow: 4-1100 di OPR inflow 500000, unexplained 0
            cf = requests.get(f"{BASE}/rahaza/finance/reports/cash-flow", params={"from": f"{YEAR}-01-01", "to": f"{YEAR}-12-31"}, headers=H, timeout=30).json()
            opr_items = cf["by_account"]["groups"]["OPR"]["items"]
            item = next((i for i in opr_items if i["code"] == "4-1100"), None)
            assert item and item["inflow"] >= 500000, f"item={item}"
            assert cf["by_account"]["totals"]["unexplained"] == 0, f"unexplained={cf['by_account']['totals']['unexplained']}"
            # bs balanced
            bs = requests.get(f"{BASE}/rahaza/finance/reports/balance-sheet-monthly", params={"year": YEAR}, headers=H, timeout=30).json()
            assert bs["balanced"] is True, f"diff={bs['totals']['diff']}"
        finally:
            # void
            rv = requests.post(f"{BASE}/rahaza/journals/{je_id}/void", json={"reason": "QA"}, headers=H, timeout=30)
            assert rv.status_code == 200, rv.text
            # hapus dari mongo
            db.rahaza_journal_entries.delete_one({"id": je_id})
            db.rahaza_journal_lines.delete_many({"je_id": je_id})


# ═══════════ OPENING BALANCE ═══════════
class TestOpeningBalance:
    def test_status_empty(self, H):
        r = requests.get(f"{BASE}/rahaza/finance/opening-balance/status", headers=H, timeout=15)
        assert r.status_code == 200, r.text
        j = r.json()
        assert j["balance_accounts"] > 0
        assert len(j["relation_sheets"]) == 4

    def test_template_ok(self, H):
        r = requests.get(f"{BASE}/rahaza/finance/opening-balance/template", headers=H, timeout=30)
        assert r.status_code == 200
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        for s in ("PETUNJUK", "SALDO_AWAL", "PIUTANG_PELANGGAN", "PIUTANG_MAKLON", "HUTANG_SUPPLIER", "HUTANG_VENDOR_CMT"):
            assert s in wb.sheetnames, f"sheet {s} hilang"

    def _make_file(self, H, entries: list, header_code: str | None = None) -> bytes:
        """Unduh template, isi baris debit/kredit untuk kode akun tertentu."""
        r = requests.get(f"{BASE}/rahaza/finance/opening-balance/template", headers=H, timeout=30)
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        ws = wb["SALDO_AWAL"]
        # peta code → row
        code_row = {}
        for i, row in enumerate(ws.iter_rows(min_row=2, values_only=False), start=2):
            if row[0].value:
                code_row[str(row[0].value).strip()] = i
        for code, d, c in entries:
            i = code_row.get(code)
            assert i, f"kode {code} tidak ada di template"
            ws.cell(row=i, column=6, value=d)
            ws.cell(row=i, column=7, value=c)
        if header_code:
            i = code_row.get(header_code)
            if i:
                ws.cell(row=i, column=6, value=1000)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_preview_ok_balanced(self, H):
        data = self._make_file(H, [("1-1101", 5000000, 0), ("3-1000", 0, 5000000)])
        r = requests.post(f"{BASE}/rahaza/finance/opening-balance/preview", headers=H,
                         files={"file": ("t.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, timeout=30)
        assert r.status_code == 200, r.text
        j = r.json()
        assert j["ok"] is True, j.get("errors")
        assert j["totals"]["debit"] == 5000000
        assert j["totals"]["credit"] == 5000000

    def test_preview_header_error(self, H):
        data = self._make_file(H, [("1-1101", 5000000, 0), ("3-1000", 0, 5000000)], header_code="1-1000")
        r = requests.post(f"{BASE}/rahaza/finance/opening-balance/preview", headers=H,
                         files={"file": ("t.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, timeout=30)
        assert r.status_code == 200
        j = r.json()
        assert j["ok"] is False
        assert any("HEADER" in e or "1-1000" in e for e in j["errors"]), j["errors"]

    def test_preview_unbalanced_and_retained(self, H):
        # DK tidak balance tanpa balance_to_retained
        data = self._make_file(H, [("1-1101", 5000000, 0), ("3-1000", 0, 4000000)])
        r = requests.post(f"{BASE}/rahaza/finance/opening-balance/preview", headers=H,
                         files={"file": ("t.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, timeout=30)
        j = r.json()
        assert j["ok"] is False
        # dengan balance_to_retained
        r2 = requests.post(f"{BASE}/rahaza/finance/opening-balance/preview", headers=H,
                          data={"balance_to_retained": "true"},
                          files={"file": ("t.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, timeout=30)
        j2 = r2.json()
        assert j2["ok"] is True, j2.get("errors")
        assert any(ln["account_code"] == "3-2000" for ln in j2["lines"]), "penyeimbang ke 3-2000 tidak ada"

    def test_apply_and_lifecycle(self, H, db):
        data = self._make_file(H, [("1-1101", 5000000, 0), ("3-1000", 0, 5000000)])
        r = requests.post(f"{BASE}/rahaza/finance/opening-balance/apply", headers=H,
                         data={"ob_date": f"{YEAR}-01-01"},
                         files={"file": ("t.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, timeout=60)
        assert r.status_code == 200, r.text
        j = r.json()
        je = j["journal"]
        je_id = je["id"]
        try:
            assert je["status"] == "posted"
            assert je.get("flags", {}).get("locked") is True
            # apply kedua → 409
            data2 = self._make_file(H, [("1-1101", 5000000, 0), ("3-1000", 0, 5000000)])
            r2 = requests.post(f"{BASE}/rahaza/finance/opening-balance/apply", headers=H,
                              data={"ob_date": f"{YEAR}-01-01"},
                              files={"file": ("t.xlsx", data2, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, timeout=30)
            assert r2.status_code == 409, r2.text
            # status menampilkan journal
            st = requests.get(f"{BASE}/rahaza/finance/opening-balance/status", headers=H, timeout=15).json()
            assert st["journal"] is not None
            assert len(st["lines"]) >= 2
            # void — needs force=True karena source_module bukan manual
            rv = requests.post(f"{BASE}/rahaza/journals/{je_id}/void", json={"reason": "QA", "force": True}, headers=H, timeout=30)
            assert rv.status_code == 200, rv.text
            st2 = requests.get(f"{BASE}/rahaza/finance/opening-balance/status", headers=H, timeout=15).json()
            assert st2["journal"] is None
        finally:
            db.rahaza_journal_entries.delete_one({"id": je_id})
            db.rahaza_journal_lines.delete_many({"je_id": je_id})


# ═══════════ REGRESI ═══════════
class TestRegression:
    @pytest.mark.parametrize("path,params", [
        ("/rahaza/finance/reports/trial-balance", {"from": f"{YEAR}-01-01", "to": f"{YEAR}-12-31"}),
        ("/rahaza/finance/reports/general-ledger", {"account_code": "1-1101"}),
        ("/rahaza/finance/reports/journal-list", {"limit": 10}),
        ("/rahaza/finance/reports/balance-sheet", {"as_of": f"{YEAR}-12-31"}),
        ("/rahaza/finance/reports/profit-loss", {}),
    ])
    def test_ok(self, H, path, params):
        r = requests.get(f"{BASE}{path}", params=params, headers=H, timeout=30)
        assert r.status_code == 200, f"{path} → {r.status_code} {r.text[:200]}"

    def test_tb_balanced(self, H):
        r = requests.get(f"{BASE}/rahaza/finance/reports/trial-balance", params={"from": f"{YEAR}-01-01", "to": f"{YEAR}-12-31"}, headers=H, timeout=30)
        assert r.json().get("balanced") is True
