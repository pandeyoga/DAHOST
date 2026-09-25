"""core.opening_balance — SSOT template, pemeriksaan & posting SALDO AWAL NERACA GO-LIVE.

Dipakai layar Master Akuntansi → Saldo Awal (routes/rahaza_opening_balance.py) DAN
scripts/saldo_awal_golive.py, supaya aturan pemeriksaannya satu: akun harus aktif & postable &
tipe neraca, satu sisi saja, D = K (selisih boleh ditutup ke 3-2000 Laba Ditahan), rincian per
relasi = saldo akun kontrol, dan hanya SATU jurnal `opening_balance` aktif.
"""
from __future__ import annotations

import io
from datetime import date, datetime, timezone
from uuid import uuid4

import openpyxl
from openpyxl.styles import Font, PatternFill

BALANCE_TYPES = ("ASSET", "LIABILITY", "EQUITY")
RETAINED_EARNINGS = "3-2000"
SOURCE_MODULE = "opening_balance"
RELATION_SHEETS = [
    ("PIUTANG_PELANGGAN", "1-1301", ["kode_pelanggan", "nama_pelanggan", "no_invoice", "tanggal_invoice", "jumlah", "keterangan"]),
    ("PIUTANG_MAKLON", "1-1305", ["kode_klien", "nama_klien", "no_invoice", "tanggal_invoice", "jumlah", "keterangan"]),
    ("HUTANG_SUPPLIER", "2-1100", ["kode_supplier", "nama_supplier", "no_tagihan", "tanggal_tagihan", "jumlah", "keterangan"]),
    ("HUTANG_VENDOR_CMT", "2-1110", ["kode_vendor", "nama_vendor", "no_tagihan", "tanggal_tagihan", "jumlah", "keterangan"]),
]
_HDR = Font(bold=True, color="FFFFFF")
_FILL = PatternFill("solid", fgColor="1F4E78")
PETUNJUK = [
    "SALDO AWAL NERACA GO-LIVE — CV. Dewi Aditya",
    "1. Isi kolom debit / kredit di sheet SALDO_AWAL sesuai neraca penutup pembukuan lama per tanggal go-live (Rp, tanpa titik).",
    "2. Akun HEADER (huruf tebal, kolom is_header = YA) TIDAK boleh diisi — isi di akun anaknya.",
    "3. Total Debit harus = Total Kredit. Selisih boleh ditutup otomatis ke 3-2000 Laba Ditahan (centang opsi saat unggah).",
    "4. Rincian per relasi (PIUTANG_PELANGGAN, HUTANG_SUPPLIER, HUTANG_VENDOR_CMT, PIUTANG_MAKLON) harus SAMA dengan saldo akun kontrolnya.",
    "5. Akumulasi penyusutan (1-2201 dst.) diisi di kolom KREDIT. Prive/rugi berjalan di 3-4000 / 3-3000 kolom DEBIT.",
    "6. Unggah di Portal Keuangan → Master Akuntansi → Saldo Awal: pratinjau dulu, baru Posting.",
]


def _head(ws, cols):
    ws.append(cols)
    for c in ws[1]:
        c.font, c.fill = _HDR, _FILL


async def build_template(db) -> tuple[bytes, int]:
    acc = await db.rahaza_coa_accounts.find({"active": True, "type": {"$in": list(BALANCE_TYPES)}}, {"_id": 0}).sort("code", 1).to_list(5000)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PETUNJUK"
    for row in PETUNJUK:
        ws.append([row])
    ws.column_dimensions["A"].width = 120
    ws = wb.create_sheet("SALDO_AWAL")
    _head(ws, ["kode_akun", "nama_akun", "tipe", "saldo_normal", "is_header", "debit", "kredit", "keterangan"])
    for a in acc:
        hdr = bool(a.get("is_group"))
        ws.append([a["code"], a["name"], a["type"], a.get("normal_balance"), "YA" if hdr else "",
                   None if hdr else 0, None if hdr else 0, ""])
        if hdr:
            for c in ws[ws.max_row]:
                c.font = Font(bold=True)
    for col, w in zip("ABCDEFGH", (12, 46, 12, 12, 10, 16, 16, 30)):
        ws.column_dimensions[col].width = w
    for name, ctrl, cols in RELATION_SHEETS:
        ws = wb.create_sheet(name)
        _head(ws, cols)
        ws.append([f"# akun kontrol {ctrl} — total sheet ini harus = saldo {ctrl} di SALDO_AWAL"])
        for col in "ABCDEF":
            ws.column_dimensions[col].width = 20
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), len(acc)


def _num(v) -> float:
    if v in (None, ""):
        return 0.0
    if isinstance(v, str):
        s = v.strip().replace("Rp", "").replace(" ", "")
        if not s:
            return 0.0
        if "," in s and "." in s:
            s = s.replace(".", "").replace(",", ".")
        elif "," in s:
            s = s.replace(",", ".") if s.count(",") == 1 and len(s.split(",")[1]) <= 2 else s.replace(",", "")
        elif s.count(".") > 1 or (s.count(".") == 1 and len(s.split(".")[1]) == 3):
            s = s.replace(".", "")
        return float(s)
    return float(v)


async def existing_opening(db) -> dict | None:
    je = await db.rahaza_journal_entries.find_one({"source_module": SOURCE_MODULE, "status": {"$ne": "voided"}},
                                                  {"_id": 0, "id": 1, "je_number": 1, "date": 1, "memo": 1, "status": 1,
                                                   "total_debit": 1, "total_credit": 1, "posted_at": 1, "flags": 1, "source_ref": 1,
                                                   "created_by_name": 1})
    return je


async def parse_workbook(db, data: bytes, balance_to_retained: bool = False) -> dict:
    """Baca berkas → {lines, rows, errors, totals, relations, ok}. TIDAK menulis apa pun."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        return {"ok": False, "errors": [f"Berkas bukan Excel .xlsx yang valid: {e}"], "lines": [], "rows": [], "totals": {}, "relations": []}
    if "SALDO_AWAL" not in wb.sheetnames:
        return {"ok": False, "errors": ["Sheet SALDO_AWAL tidak ditemukan — pakai template yang diunduh dari layar ini."],
                "lines": [], "rows": [], "totals": {}, "relations": []}
    acc = {a["code"]: a async for a in db.rahaza_coa_accounts.find({"active": True}, {"_id": 0})}
    errors, lines, rows, td, tc = [], [], [], 0.0, 0.0
    for i, r in enumerate(wb["SALDO_AWAL"].iter_rows(values_only=True, min_row=2), start=2):
        if not r or not r[0]:
            continue
        code = str(r[0]).strip()
        try:
            d, c = _num(r[5] if len(r) > 5 else 0), _num(r[6] if len(r) > 6 else 0)
        except ValueError:
            errors.append(f"baris {i}: {code} nilai debit/kredit bukan angka")
            continue
        if d == 0 and c == 0:
            continue
        a = acc.get(code)
        err = None
        if not a:
            err = f"akun {code} tidak ada / nonaktif"
        elif a.get("is_group"):
            err = f"{code} adalah HEADER — isi di akun anaknya"
        elif a["type"] not in BALANCE_TYPES:
            err = f"{code} bukan akun neraca ({a['type']})"
        elif d and c:
            err = f"{code} debit DAN kredit terisi — pilih satu"
        elif d < 0 or c < 0:
            err = f"{code} nilai negatif — pindahkan ke sisi lawan"
        rows.append({"row": i, "code": code, "name": (a or {}).get("name") or (str(r[1]).strip() if len(r) > 1 and r[1] else ""),
                     "type": (a or {}).get("type"), "debit": round(d, 2), "credit": round(c, 2), "error": err})
        if err:
            errors.append(f"baris {i}: {err}")
            continue
        lines.append({"account_code": code, "account_name": a["name"], "account_type": a["type"], "debit": round(d, 2), "credit": round(c, 2),
                      "description": (str(r[7]).strip() if len(r) > 7 and r[7] else "Saldo awal go-live")})
        td += d
        tc += c
    relations = []
    for name, ctrl, _cols in RELATION_SHEETS:
        if name not in wb.sheetnames:
            continue
        tot, n = 0.0, 0
        for r in wb[name].iter_rows(values_only=True, min_row=2):
            if not r or not r[0] or str(r[0]).startswith("#"):
                continue
            try:
                tot += _num(r[4] if len(r) > 4 else 0)
            except ValueError:
                errors.append(f"{name}: kolom jumlah bukan angka pada baris {r[0]}")
            n += 1
        bal = next((ln["debit"] - ln["credit"] for ln in lines if ln["account_code"] == ctrl), 0.0)
        if acc.get(ctrl, {}).get("type") == "LIABILITY":
            bal = -bal
        match = not tot or abs(tot - bal) <= 0.5
        relations.append({"sheet": name, "control_account": ctrl, "detail_total": round(tot, 2), "control_balance": round(bal, 2), "rows": n, "match": match})
        if not match:
            errors.append(f"{name}: rincian Rp {tot:,.0f} ≠ saldo akun kontrol {ctrl} Rp {bal:,.0f}")
    diff = round(td - tc, 2)
    balancing = None
    if diff and balance_to_retained and RETAINED_EARNINGS in acc:
        balancing = {"account_code": RETAINED_EARNINGS, "account_name": acc[RETAINED_EARNINGS]["name"], "account_type": "EQUITY",
                     "debit": round(-diff, 2) if diff < 0 else 0.0, "credit": round(diff, 2) if diff > 0 else 0.0,
                     "description": "Penyeimbang saldo awal → Laba Ditahan"}
        lines.append(balancing)
        tc += max(0, diff)
        td += -diff if diff < 0 else 0
        diff = 0.0
    if diff:
        errors.append(f"Debit Rp {td:,.0f} ≠ Kredit Rp {tc:,.0f} (selisih Rp {diff:,.0f}) — centang 'tutup selisih ke Laba Ditahan' bila memang begitu")
    return {"ok": not errors, "errors": errors, "lines": lines, "rows": rows, "relations": relations, "balancing": balancing,
            "totals": {"debit": round(td, 2), "credit": round(tc, 2), "diff": diff, "lines": len(lines)}}


async def post_opening(db, user: dict, ob_date: str, lines: list, source_ref: str) -> dict:
    """Posting SATU jurnal pembuka terkunci lewat mesin jurnal yang sama dengan Jurnal Umum."""
    from fastapi import HTTPException

    from routes.rahaza_journals import (
        _check_period_open,
        _gen_je_number,
        _mirror_lines,
        _validate_lines,
    )
    try:
        d = date.fromisoformat(ob_date)
    except ValueError:
        raise HTTPException(400, "Tanggal saldo awal harus YYYY-MM-DD.")
    if await existing_opening(db):
        raise HTTPException(409, "Sudah ada jurnal saldo awal (opening_balance) yang aktif — void dulu bila ingin mengulang.")
    if len(lines) < 2:
        raise HTTPException(400, "Semua saldo 0 — tidak ada jurnal pembuka yang perlu dibuat.")
    total_d, total_c = await _validate_lines(db, lines)
    await _check_period_open(db, d)
    now = datetime.now(timezone.utc)
    doc = {"id": str(uuid4()), "je_number": await _gen_je_number(db, d), "date": d.isoformat(),
           "memo": f"Saldo Awal Go-Live per {d.isoformat()} ({source_ref})", "source_module": SOURCE_MODULE, "source_ref": source_ref,
           "status": "posted", "total_debit": total_d, "total_credit": total_c, "created_at": now, "updated_at": now,
           "created_by": user["id"], "created_by_name": user.get("name", ""), "posted_at": now, "posted_by": user["id"],
           "voided_at": None, "voided_by": None, "flags": {"opening_balance": True, "locked": True},
           "lines": [{"line_id": str(uuid4()), "account_code": ln["account_code"], "account_name": ln["account_name"],
                      "account_type": ln["account_type"], "debit": ln["debit"], "credit": ln["credit"],
                      "description": (ln.get("description") or "").strip(), "cost_center_id": None} for ln in lines]}
    await db.rahaza_journal_entries.insert_one(doc)
    await _mirror_lines(db, doc)
    doc.pop("_id", None)
    return doc
