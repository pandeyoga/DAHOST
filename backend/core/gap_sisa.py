"""gap_sisa — DATA_YANG_PERLU_DIISI_DA_SISA.xlsx: berkas isian BERIKUTNYA yang dibangun dari berkas klien yang
baru diterapkan. Hanya kelompok/baris yang DILEWATI importir yang ditulis ulang (utuh per kelompok, dengan kolom
`masalah` + `tindakan`), ditambah sheet VARIAN_BARU untuk model tanpa SKU dan material yang butuh isi kemasan.
"""
from __future__ import annotations

import difflib
import io

import openpyxl

from core.bom_fill import CATEGORIES, _SPLIT_RE, _variants_label, clean, load_model_variants

VARIAN_BARU_COLS = ["kode_model", "nama_model", "warna", "kode_warna", "ukuran", "harga_jual", "keterangan", "kode_warna_tersedia", "ukuran_tersedia"]
SISA_BOM_EXTRA = ["masalah", "tindakan"]

TINDAKAN = {
    "kelompok_tanpa_varian": "isi kolom 'varian' dengan warna/ukuran pemakai bahan ini (pilih dari varian_tersedia, pisah koma)",
    "warna_tak_dikenal": "ganti nama warna di kolom 'varian' dengan yang ada di varian_tersedia",
    "model_tanpa_sku": "isi sheet VARIAN_BARU untuk model ini (kode_warna + ukuran) — kelompok ini otomatis ikut masuk setelah varian dibuat",
    "model_dihentikan": "tidak perlu diisi — model sudah dihentikan (semua SKU nonaktif); hapus kelompok ini atau aktifkan kembali SKU-nya di RnD bila masih dijual",
    "satuan_tak_valid": "perbaiki satuan/qty (m, cm, roll, pack, pcs, gross) atau isi isi_per_satuan_beli material di sheet MATERIAL",
    "qty_kosong": "isi qty_per_pcs (contoh: 60 cm, 1 pcs, 0,5)",
    "kode_tak_dikenal": "ganti kode_material dengan kode di REF_AKSESORIS",
    "model_tak_dikenal": "kode model tidak ada di master — periksa kode (lihat sheet MODEL)",
    "baris_tanpa_model": "tulis kode_model di baris pertama kelompok",
}


BOM_SHEETS = ("BOM_AKSESORIS", "BOM_OTOMATIS")  # BOM_OTOMATIS = sheet berkas FOKUS (kelompok yang sudah terisi otomatis)


def bom_rows_from_wb(wb) -> list[tuple]:
    """Baris BOM_AKSESORIS lalu BOM_OTOMATIS (satu daftar, indeks baris virtual mulai 2). Elemen ke-10 = label sumber
    'SHEET baris N' untuk pesan ke pengguna; kolom data (0..8) tidak berubah."""
    out: list[tuple] = []
    for sheet in BOM_SHEETS:
        if sheet not in wb.sheetnames:
            continue
        for i, r in enumerate(wb[sheet].iter_rows(values_only=True, min_row=2), start=2):
            r = tuple(r)[:9]
            out.append(r + (None,) * (9 - len(r)) + (f"{sheet} baris {i}",))
    return out


def bom_rows_from_file(data: bytes) -> list[tuple]:
    if not data:
        return []
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    return bom_rows_from_wb(wb)


def group_header_rows(rows: list[tuple]) -> dict[int, int]:
    """row → baris kepala kelompoknya (baris ber-kode_model)."""
    out, head = {}, None
    for i, r in enumerate(rows, start=2):
        if clean(r[0]):
            head = i
        if head is not None and (clean(r[0]) or clean(r[2])):
            out[i] = head
    return out


def sisa_bom_rows(parsed: dict, rows: list[tuple], vlabel: dict[str, str]) -> list[list]:
    """Kelompok yang punya ≥1 masalah ditulis ulang UTUH (agar unggah ulang = ganti utuh yang benar)."""
    heads = group_header_rows(rows)
    by_head: dict[int, list[dict]] = {}
    for i in parsed.get("bom_issues") or []:
        if i["kategori"] == "model_dihentikan":
            continue  # model sudah dihentikan klien → tidak perlu ditulis ulang
        h = heads.get(i["row"], i["row"])
        by_head.setdefault(h, []).append(i)
    out: list[list] = []
    for h in sorted(by_head):
        members = [r for r in sorted(heads) if heads[r] == h] or [h]
        for r in members:
            src = rows[r - 2] if 0 <= r - 2 < len(rows) else (None,) * 9
            here = [i for i in by_head[h] if i["row"] == r]
            masalah = "; ".join(f"{CATEGORIES.get(i['kategori'], i['kategori'])}: {i['detail']}" for i in here)
            tindakan = "; ".join(dict.fromkeys(TINDAKAN.get(i["kategori"], "") for i in here if TINDAKAN.get(i["kategori"])))
            mcode = clean(src[0]).upper()
            out.append([src[0], src[1], src[2], src[3], src[4], src[5], src[6], src[7],
                        vlabel.get(mcode, "") if mcode else "", masalah, tindakan])
    return out


async def varian_baru_rows(db, parsed: dict) -> list[list]:
    colors = await db.rahaza_colors.find({"active": {"$ne": False}}, {"_id": 0, "code": 1, "name": 1}).to_list(2000)
    by_name = {clean(c.get("name")).upper(): c["code"] for c in colors}
    sizes = [s["code"] for s in await db.rahaza_sizes.find({"active": {"$ne": False}}, {"_id": 0, "code": 1, "order_seq": 1}).sort("order_seq", 1).to_list(100)]
    color_list = ", ".join(f"{c['code']}={clean(c.get('name')).title()}" for c in sorted(colors, key=lambda c: c["code"]))
    tanpa_sku: dict[str, dict] = {}
    for i in parsed.get("bom_issues") or []:
        if i["kategori"] != "model_tanpa_sku":
            continue
        d = tanpa_sku.setdefault(i["model_code"], {"name": i.get("model_name"), "colors": []})
        for tok in _SPLIT_RE.split(i.get("varian_raw") or ""):
            t = tok.strip()
            if t and t.title() not in d["colors"]:
                d["colors"].append(t.title())
    out = []
    for k, v in sorted(tanpa_sku.items()):
        for warna in (v["colors"] or [""]):
            key = warna.upper()
            code = by_name.get(key)
            if not code:
                close = difflib.get_close_matches(key, list(by_name), n=1, cutoff=0.8)
                code = by_name[close[0]] if close else ""
            ket = "kode_warna disarankan dari master (periksa)" if code else "warna tidak ada di master — pilih kode dari kode_warna_tersedia atau buat warna baru di RnD"
            out.append([k, v["name"], warna, code, "", None, ket, color_list, ", ".join(sizes)])
    return out


def materials_needing_pack(parsed: dict) -> dict[str, str]:
    """kode material → catatan, untuk baris BOM yang gagal karena isi kemasan (pcs→roll/pack) belum ada."""
    out: dict[str, str] = {}
    for i in parsed.get("bom_issues") or []:
        if i["kategori"] == "satuan_tak_valid" and "isi_per_satuan_beli" in i.get("detail", "") and i.get("code"):
            out.setdefault(i["code"], f"isi kemasan (pcs per {i.get('unit_raw') or 'kemasan'}) belum ada — dipakai BOM {i.get('model_code')} {i.get('where') or 'baris ' + str(i['row'])}")
    return out


async def variant_labels(db, parsed: dict) -> dict[str, str]:
    codes = {i.get("model_code") for i in parsed.get("bom_issues") or [] if i.get("model_code")}
    models = await db.rahaza_models.find({"code": {"$in": list(codes)}}, {"_id": 0, "id": 1, "code": 1}).to_list(5000)
    vmap = await load_model_variants(db, [m["id"] for m in models])
    return {m["code"]: _variants_label(vmap.get(m["id"]) or []) for m in models}


def petunjuk_sisa(parsed: dict, applied: dict | None) -> list[str]:
    t = parsed.get("totals") or {}
    c = parsed.get("bom_issue_counts") or {}
    lines = [
        "",
        "═══ DAMPAK BERKAS SEBELUMNYA (berkas yang Anda unggah) ═══",
        f"  BOM_AKSESORIS  DITERAPKAN : {t.get('bom_lines', 0)} baris bahan · {t.get('bom_groups', 0)} kelompok · {t.get('bom_models', 0)} model dari {t.get('bom_models_in_file', 0)} model di berkas.",
        "               Setiap (model, varian) di berkas: daftar aksesorisnya di BOM DIGANTI UTUH (potongan kain tetap); model di luar berkas tidak disentuh;",
        "               BOM varian yang belum ada dibuat (disalin dari varian saudara atau BOM dasar kosong); HPP dihitung ulang & diberi status tervalidasi/belum.",
    ]
    if applied:
        lines.append(f"               Hasil terapkan terakhir: {applied.get('bom_lines_appended', 0)} baris terpasang · {applied.get('boms_touched', 0)} BOM · "
                     f"HPP {applied.get('hpp_validated', 0)} tervalidasi / {applied.get('hpp_unvalidated', 0)} belum.")
    lines += [
        f"  BOM_AKSESORIS  DILEWATI   : {t.get('bom_skipped', 0)} catatan → " + " · ".join(f"{CATEGORIES.get(k, k)} {v}" for k, v in c.items()),
        f"  HARGA_JUAL_SKU            : {t.get('sku_prices', 0)} SKU harga_jual berbeda dari sistem → harga jual master diperbarui; "
        f"{t.get('sku_deactivate', 0)} SKU bertuliskan 'sudah tidak dijual' → varian + barang jadi + item katalognya DINONAKTIFKAN (tidak dihapus).",
        f"  TOKO                      : {t.get('stores', 0)} toko → rekening pencairan (kode akun kas/bank) diisi.",
        f"  REKENING · MATERIAL · GAJI: {t.get('accounts', 0)} rekening · {t.get('materials', 0)} harga material · {t.get('salaries', 0)} gaji berbeda dari sistem.",
        f"  STOK_AWAL                 : {t.get('stock_rows', 0)} baris (stok awal hanya dicatat sekali per item+lokasi).",
        "  Angka 0 = sudah sama dengan sistem (unggah ulang berkas yang sama tidak mengubah apa pun — importir idempoten).",
        "  Mode 'Terapkan hanya BOM' hanya menerapkan VARIAN_BARU + BOM_AKSESORIS; sheet lain butuh 'Terapkan semua sheet'.",
        "",
        "═══ BERKAS INI = SISANYA SAJA ═══",
        "  BOM_AKSESORIS  — hanya kelompok yang DILEWATI, ditulis ulang utuh. Kolom 'masalah' = alasan, 'tindakan' = yang harus Anda isi.",
        "                   Jangan hapus baris bahan yang sudah benar: saat unggah ulang, kelompok (model, varian) diganti utuh dengan isi berkas ini.",
        "  VARIAN_BARU    — model yang belum punya SKU: isi kode_warna (disarankan otomatis dari nama) & ukuran; SKU = MODEL-WARNA-UKURAN dibuat otomatis",
        "                   beserta barang jadinya; harga_jual opsional. Setelah itu kelompok BOM model tsb (sheet BOM_AKSESORIS) ikut diterapkan pada unggahan yang sama.",
        "  MATERIAL       — harga masih 0 (HPP belum tervalidasi) + material yang isi kemasannya (pcs per roll/pack) belum ada padahal dipakai BOM.",
        "                   Isi satuan_beli, isi_per_satuan_beli, harga_per_satuan_beli → HPP model pemakainya menjadi tervalidasi.",
        "  HARGA_JUAL_SKU — hanya SKU yang MASIH tanpa harga. Isi angka, atau tulis 'sudah tidak dijual' untuk menonaktifkan SKU.",
        "  Unggah balik berkas ini di layar yang sama → 'Terapkan (hanya BOM)' = VARIAN_BARU + BOM_AKSESORIS; 'Terapkan semua sheet' = termasuk MATERIAL/harga/gaji/dll.",
    ]
    return lines
