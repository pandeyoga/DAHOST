"""gap_workbook — SATU berkas Excel berisi semua data yang masih harus diisi manusia,
sudah berisi baris datanya (model/SKU/material/akun) supaya user tidak mengisi dari nol.

Sheet yang bisa DIUNGGAH BALIK ke layar "Impor Harga · Rekening · BOM":
  MATERIAL · BOM_AKSESORIS · MODEL · REKENING · TOKO (importir lama)
  HARGA_JUAL_SKU · STOK_AWAL_FG · STOK_AWAL_MATERIAL · GAJI (ditambahkan di sini)
Sheet SALDO_AWAL & PIUTANG/HUTANG diunggah ke layar "Saldo Awal" (berkas yang sama boleh dipakai).
"""
from __future__ import annotations

import io
import uuid
from datetime import datetime, timezone

import openpyxl
from openpyxl.styles import Font, PatternFill

from core.master_fill import MAT_COLS, REK_COLS, TOKO_COLS, BOM_COLS, MODEL_COLS, _head, _num
from core import opening_balance as _ob

SKU_PRICE_COLS = ["sku", "nama", "kode_model", "warna", "ukuran", "harga_saran_dari_saudara", "harga_jual", "keterangan"]
STOK_FG_COLS = ["sku", "nama", "kode_model", "warna", "ukuran", "lokasi_kode", "qty_awal", "hpp_satuan_sistem", "keterangan"]
STOK_MAT_COLS = ["kode", "nama", "tipe", "satuan_dasar", "lokasi_kode", "qty_awal", "harga_satuan_sistem", "keterangan"]
GAJI_COLS = ["kode_karyawan", "nama", "skema", "gaji_pokok_per_periode", "tarif_lembur_per_jam", "keterangan"]
DEFAULT_LOC_FG, DEFAULT_LOC_MAT = "GD-L1-RAK", "GD-L1"

PETUNJUK = [
    "DATA YANG MASIH HARUS DIISI — CV. Dewi Aditya (dibuat otomatis dari kondisi sistem saat ini)",
    "Baris yang tertulis di sini = data yang belum lengkap. Isi kolom KUNING saja; kolom lain hanya informasi.",
    "Baris yang kolom kuningnya dibiarkan kosong/0 TIDAK diubah — aman diunggah sebagian.",
    "",
    "UNGGAH ke Portal Keuangan → Master Akuntansi → Impor Harga · Rekening · BOM (berkas ini langsung):",
    "  MATERIAL          — material yang harganya masih 0: isi satuan_beli, isi_per_satuan_beli, harga_per_satuan_beli.",
    "  BOM_AKSESORIS     — aksesoris/bahan pendukung per model (benang, label, kancing…). Satu baris per bahan; kode_material lihat sheet REF_AKSESORIS.",
    "                      Baris ber-kode_model = awal kelompok; baris di bawahnya yang kode_model-nya KOSONG = bahan lain untuk kelompok yang sama.",
    "                      qty_per_pcs boleh ditulis dengan satuan (\"60 cm\", \"1 pcs\"). Kolom 'varian' (opsional) = warna/ukuran pemakai bahan ini,",
    "                      dipisah koma — lihat kolom 'varian_tersedia'. Kosong = semua varian model. Kode model sama dengan nama sama = model beda ukuran.",
    "  MODEL             — berat_gram per model (untuk ongkir).",
    "  HARGA_JUAL_SKU    — SKU barang jadi yang belum berharga; harga_saran = harga SKU saudara di model yang sama (salin ke harga_jual bila setuju).",
    "  STOK_AWAL_FG      — stok fisik barang jadi per SKU per tanggal go-live (qty). Nilai rupiahnya diambil dari SALDO_AWAL (akun Persediaan), bukan dari sini.",
    "  STOK_AWAL_MATERIAL— stok fisik kain & aksesoris per tanggal go-live (qty dalam satuan dasar).",
    "  REKENING / TOKO   — no. rekening & atas nama; rekening pencairan tiap toko (sekarang default 1-1201 Bank BCA).",
    "  GAJI              — gaji pokok per periode & tarif lembur per karyawan (skema monthly/daily/piece).",
    "",
    "UNGGAH ke Portal Keuangan → Master Akuntansi → Saldo Awal (berkas yang sama):",
    "  SALDO_AWAL        — neraca penutup pembukuan lama per tanggal go-live; PIUTANG_*/HUTANG_* rinciannya per pelanggan/vendor.",
    "",
    "Tidak ada sandi di berkas ini. Techpack/foto/SOP diunggah per model di Portal RnD.",
]
YELLOW = PatternFill("solid", fgColor="FFF2CC")


def _mark(ws, cols: list[str], fill_cols: list[str]) -> None:
    idx = [cols.index(c) + 1 for c in fill_cols if c in cols]
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for i in idx:
            row[i - 1].fill = YELLOW


def _widths(ws, widths):
    for i, w in enumerate(widths):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w


async def build_gap_workbook(db, parsed: dict | None = None, data: bytes | None = None, applied: dict | None = None) -> tuple[bytes, dict]:
    """parsed/data = berkas klien yang baru diterapkan → BOM_AKSESORIS hanya berisi kelompok yang dilewati (+VARIAN_BARU)."""
    from core import gap_sisa
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PETUNJUK"
    for r in PETUNJUK + (gap_sisa.petunjuk_sisa(parsed, applied) if parsed else []):
        ws.append([r])
    c = db
    n_fg = await c.rahaza_materials.count_documents({"type": "fg", "active": {"$ne": False}})
    n_fg_price = await c.rahaza_materials.count_documents({"type": "fg", "active": {"$ne": False}, "retail_price_master": {"$gt": 0}})
    n_var_off = await c.rahaza_model_variants.count_documents({"active": False})
    n_bom_acc = await c.rahaza_boms.count_documents({"active": {"$ne": False}, "materials.material_type": "accessory"})
    n_bom = await c.rahaza_boms.count_documents({"active": {"$ne": False}})
    n_store = await c.marketing_platform_accounts.count_documents({"coa_cash_code": {"$exists": True, "$nin": ["", None]}})
    n_store_all = await c.marketing_platform_accounts.count_documents({})
    n_hpp_ok = await c.rahaza_models.count_documents({"active": {"$ne": False}, "hpp_validation.status": "tervalidasi"})
    n_model = await c.rahaza_models.count_documents({"active": {"$ne": False}})
    for r in ["", "═══ KEADAAN SISTEM SAAT BERKAS INI DIBUAT ═══",
              f"  Model aktif {n_model} · HPP tervalidasi {n_hpp_ok} · BOM aktif {n_bom} (dengan aksesoris {n_bom_acc})",
              f"  SKU barang jadi aktif {n_fg} · sudah berharga jual {n_fg_price} · SKU dinonaktifkan ('sudah tidak dijual') {n_var_off}",
              f"  Toko dengan rekening pencairan {n_store}/{n_store_all}"]:
        ws.append([r])
    ws.column_dimensions["A"].width = 130
    stats: dict = {}
    need_pack = gap_sisa.materials_needing_pack(parsed) if parsed else {}

    # ── MATERIAL: hanya yang harganya 0 (tanpa panel potongan) + isi kemasan yang dibutuhkan BOM ──
    mats = await db.rahaza_materials.find({"active": {"$ne": False}, "type": {"$ne": "fg"}, "code": {"$not": {"$regex": "^CUT-"}}},
                                          {"_id": 0}).sort("code", 1).to_list(20000)
    ws = wb.create_sheet("MATERIAL")
    _head(ws, MAT_COLS)
    n = 0
    for m in mats:
        no_price = float(m.get("unit_cost") or 0) <= 0
        if not no_price and m.get("code") not in need_pack:
            continue
        ket = "; ".join(x for x in (("harga masih 0" if no_price else ""), need_pack.get(m.get("code"), "")) if x)
        ws.append([m.get("code"), m.get("name"), m.get("type"), m.get("category_name") or m.get("category"), m.get("unit"),
                   m.get("purchase_unit") or "", None, None, float(m.get("unit_cost") or 0), m.get("min_stock") or None, ket])
        n += 1
    _mark(ws, MAT_COLS, ["satuan_beli", "isi_per_satuan_beli", "harga_per_satuan_beli"])
    _widths(ws, (14, 40, 10, 16, 12, 12, 16, 20, 20, 10, 24))
    stats["material_tanpa_harga"] = n

    # ── MODEL & BOM_AKSESORIS ──
    models = await db.rahaza_models.find({"active": {"$ne": False}}, {"_id": 0, "id": 1, "code": 1, "name": 1, "category_name": 1,
                                                                       "weight_gram": 1, "hpp": 1}).sort("code", 1).to_list(5000)
    has_bom, has_acc = set(), set()
    async for b in db.rahaza_boms.find({"active": {"$ne": False}, "is_active": True}, {"_id": 0, "model_id": 1, "materials": 1}):
        has_bom.add(b["model_id"])
        if any((ln.get("material_type") or "").lower() not in ("fabric", "") and not ln.get("is_cut_panel") for ln in b.get("materials") or []):
            has_acc.add(b["model_id"])
    from core.bom_fill import load_model_variants, _variants_label, clean
    vmap = await load_model_variants(db, [m["id"] for m in models])
    stopped = {m["id"] for m in models if not vmap.get(m["id"]) and await db.rahaza_model_variants.count_documents({"model_id": m["id"]}) > 0}
    ws = wb.create_sheet("MODEL")
    _head(ws, MODEL_COLS + ["hpp_sistem", "varian_tersedia"])
    for m in models:
        ws.append([m["code"], m["name"], m.get("category_name"), float(m.get("weight_gram") or 0) or None,
                   "ya" if m["id"] in has_bom else "BELUM", "ya" if m["id"] in has_acc else "BELUM",
                   "DIHENTIKAN — semua SKU 'sudah tidak dijual' (HARGA_JUAL_SKU)" if m["id"] in stopped else ("" if m["id"] in has_bom else "belum punya BOM — isi BOM_AKSESORIS + kain di RnD"), float(m.get("hpp") or 0) or None,
                   "(semua SKU nonaktif)" if m["id"] in stopped else _variants_label(vmap.get(m["id"]) or [])])
    _mark(ws, MODEL_COLS, ["berat_gram"])
    _widths(ws, (14, 30, 16, 12, 12, 16, 44, 14, 60))
    stats["model_tanpa_berat"] = sum(1 for m in models if not m.get("weight_gram"))
    stats["model_tanpa_bom"] = sum(1 for m in models if m["id"] not in has_bom and m["id"] not in stopped)
    stats["model_tanpa_aksesoris"] = sum(1 for m in models if m["id"] not in has_acc and m["id"] not in stopped)
    stats["model_dihentikan"] = len(stopped)

    ws = wb.create_sheet("BOM_AKSESORIS")
    if parsed:
        _head(ws, BOM_COLS + gap_sisa.SISA_BOM_EXTRA)
        vlabel = await gap_sisa.variant_labels(db, parsed)
        sisa_rows = gap_sisa.sisa_bom_rows(parsed, gap_sisa.bom_rows_from_file(data or b""), vlabel)
        for r in sisa_rows:
            ws.append(r)
        in_file = {clean(r[0]).upper() for r in sisa_rows if r[0]} | {g["model_code"] for g in parsed.get("bom_groups") or []}
        stats["bom_sisa_baris"] = len(sisa_rows)
        stats["bom_sisa_kelompok"] = sum(1 for r in sisa_rows if r[0])
        for c in ws[1][len(BOM_COLS):]:
            c.fill = YELLOW
    else:
        _head(ws, BOM_COLS)
        in_file = set()
    for m in models:
        if m["id"] in has_acc or m["code"] in in_file or m["id"] in stopped:
            continue
        label = _variants_label(vmap.get(m["id"]) or [])
        for _ in range(3):  # 3 baris kosong per model — tambah baris sendiri bila perlu
            ws.append([m["code"], m["name"], "", "", None, "", "isi kode_material (lihat REF_AKSESORIS) & qty_per_pcs", "", label]
                      + (["belum ada aksesoris di BOM & tidak ada di berkas", "isi bahan aksesoris model ini"] if parsed else []))
    _mark(ws, BOM_COLS, ["kode_material", "qty_per_pcs", "varian"])
    _widths(ws, (14, 30, 16, 36, 12, 10, 48, 24, 60, 70, 70))
    ws.freeze_panes = "C2"
    if parsed:
        ws = wb.create_sheet("VARIAN_BARU")
        _head(ws, gap_sisa.VARIAN_BARU_COLS)
        vb = await gap_sisa.varian_baru_rows(db, parsed)
        for r in vb:
            ws.append(r)
        _mark(ws, gap_sisa.VARIAN_BARU_COLS, ["kode_warna", "ukuran", "harga_jual"])
        _widths(ws, (14, 24, 18, 12, 12, 14, 60, 90, 40))
        ws.freeze_panes = "C2"
        stats["varian_baru"] = len(vb)
    ws = wb.create_sheet("REF_AKSESORIS")
    _head(ws, ["kode_material", "nama", "tipe", "satuan_dasar", "harga_per_satuan_dasar"])
    for m in mats:
        if (m.get("type") or "").lower() != "fabric":
            ws.append([m.get("code"), m.get("name"), m.get("type"), m.get("unit"), float(m.get("unit_cost") or 0)])
    _widths(ws, (16, 44, 12, 12, 18))

    # ── HARGA_JUAL_SKU: FG tanpa harga + saran dari saudara ──
    fgs = await db.rahaza_materials.find({"type": "fg", "active": {"$ne": False}}, {"_id": 0}).sort("code", 1).to_list(20000)
    sib: dict = {}
    for f in fgs:
        p = float(f.get("retail_price_master") or 0)
        if p > 0:
            sib.setdefault(f.get("model_code"), set()).add(p)
    ws = wb.create_sheet("HARGA_JUAL_SKU")
    _head(ws, SKU_PRICE_COLS)
    n = 0
    for f in fgs:
        if float(f.get("retail_price_master") or 0) > 0:
            continue
        prices = sorted(sib.get(f.get("model_code")) or [])
        saran = prices[0] if len(prices) == 1 else None
        ket = "" if saran else ("model ini punya beberapa harga: " + " / ".join(f"{int(p):,}".replace(",", ".") for p in prices) if prices else "belum ada SKU berharga di model ini")
        ws.append([f.get("code"), f.get("name"), f.get("model_code"), f.get("color_name"), f.get("size_code"), saran, None,
                   ket or "salin harga_saran ke harga_jual bila setuju"])
        n += 1
    _mark(ws, SKU_PRICE_COLS, ["harga_jual"])
    _widths(ws, (26, 34, 12, 14, 10, 20, 14, 48))
    stats["sku_tanpa_harga"] = n

    # ── STOK_AWAL_FG & STOK_AWAL_MATERIAL ──
    ws = wb.create_sheet("STOK_AWAL_FG")
    _head(ws, STOK_FG_COLS)
    for f in fgs:
        ws.append([f.get("code"), f.get("name"), f.get("model_code"), f.get("color_name"), f.get("size_code"), DEFAULT_LOC_FG, None,
                   float(f.get("hpp") or 0) or None, ""])
    _mark(ws, STOK_FG_COLS, ["lokasi_kode", "qty_awal"])
    _widths(ws, (26, 34, 12, 14, 10, 12, 10, 16, 30))
    ws.freeze_panes = "B2"
    ws = wb.create_sheet("STOK_AWAL_MATERIAL")
    _head(ws, STOK_MAT_COLS)
    for m in mats:
        ws.append([m.get("code"), m.get("name"), m.get("type"), m.get("unit"), DEFAULT_LOC_MAT, None, float(m.get("unit_cost") or 0) or None, ""])
    _mark(ws, STOK_MAT_COLS, ["lokasi_kode", "qty_awal"])
    _widths(ws, (16, 44, 10, 12, 12, 10, 18, 30))
    ws.freeze_panes = "B2"
    ws = wb.create_sheet("REF_LOKASI")
    _head(ws, ["lokasi_kode", "nama", "tipe"])
    async for loc in db.rahaza_locations.find({"active": {"$ne": False}}, {"_id": 0, "code": 1, "name": 1, "type": 1}).sort("code", 1):
        ws.append([loc.get("code"), loc.get("name"), loc.get("type")])
    _widths(ws, (14, 30, 10))
    stats["sku_fg"], stats["material"] = len(fgs), len(mats)

    # ── REKENING & TOKO ──
    ws = wb.create_sheet("REKENING")
    _head(ws, REK_COLS)
    n = 0
    async for a in db.rahaza_cash_accounts.find({"active": {"$ne": False}}, {"_id": 0}).sort("coa_code", 1):
        ws.append([a.get("coa_code") or a.get("code"), a.get("name"), a.get("bank_name") or "", a.get("account_number") or "", a.get("account_holder") or ""])
        n += 1
    _mark(ws, REK_COLS, ["bank", "no_rekening", "atas_nama"])
    _widths(ws, (12, 34, 16, 22, 30))
    stats["rekening"] = n
    ws = wb.create_sheet("TOKO")
    _head(ws, TOKO_COLS + ["pic_email_sekarang"])
    async for t in db.marketing_platform_accounts.find({"status": {"$ne": "archived"}}, {"_id": 0}).sort("account_code", 1):
        ws.append([t.get("account_code"), t.get("account_name"), t.get("platform"), t.get("coa_cash_code") or "", t.get("pic_email") or t.get("pic_user_email") or ""])
    _mark(ws, TOKO_COLS, ["rekening_pencairan_kode_akun"])
    _widths(ws, (12, 30, 12, 28, 30))

    # ── GAJI ──
    ws = wb.create_sheet("GAJI")
    _head(ws, GAJI_COLS)
    emps = {e["id"]: e for e in await db.rahaza_employees.find({"active": {"$ne": False}}, {"_id": 0, "id": 1, "employee_code": 1, "name": 1}).to_list(5000)}
    n = 0
    async for p in db.rahaza_payroll_profiles.find({"active": {"$ne": False}}, {"_id": 0}):
        e = emps.get(p.get("employee_id")) or {}
        ws.append([e.get("employee_code"), e.get("name"), p.get("pay_scheme"), float(p.get("base_rate") or 0) or None,
                   float(p.get("overtime_rate") or 0) or None, "tarif lembur masih 0" if not p.get("overtime_rate") else ""])
        n += 1
    _mark(ws, GAJI_COLS, ["skema", "gaji_pokok_per_periode", "tarif_lembur_per_jam"])
    _widths(ws, (16, 30, 10, 22, 20, 30))
    stats["karyawan"] = n

    # ── SALDO_AWAL (+ rincian piutang/hutang) dari template resmi layar Saldo Awal ──
    ob_bytes, _ = await _ob.build_template(db)
    src = openpyxl.load_workbook(io.BytesIO(ob_bytes))
    for s in src.worksheets:
        if s.title == "PETUNJUK":
            continue
        ws = wb.create_sheet(s.title)
        for row in s.iter_rows(values_only=True):
            ws.append(list(row))
        for c in ws[1]:
            c.font = Font(bold=True)
        if s.title == "SALDO_AWAL":
            _mark(ws, ["kode_akun", "nama_akun", "tipe", "saldo_normal", "is_header", "debit", "kredit", "keterangan"], ["debit", "kredit"])
            _widths(ws, (12, 46, 12, 12, 10, 16, 16, 30))
        else:
            _widths(ws, (20, 30, 20, 16, 16, 30))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), stats


# ── PARSE & APPLY sheet tambahan (dipanggil dari master_fill) ────────────────────────
async def parse_extra_sheets(db, wb, errors: list) -> dict:
    out = {"sku_prices": [], "sku_deactivate": [], "stock_fg": [], "stock_mat": [], "salaries": [], "new_variants": [], "warnings": []}
    locs = {l["code"]: l for l in await db.rahaza_locations.find({}, {"_id": 0, "id": 1, "code": 1}).to_list(2000)}
    if "VARIAN_BARU" in wb.sheetnames:
        models = {m["code"]: m for m in await db.rahaza_models.find({"active": {"$ne": False}}, {"_id": 0, "id": 1, "code": 1, "name": 1}).to_list(5000)}
        colors = {c["code"].upper(): c for c in await db.rahaza_colors.find({}, {"_id": 0}).to_list(2000)}
        sizes = {s["code"].upper(): s for s in await db.rahaza_sizes.find({}, {"_id": 0}).to_list(200)}
        existing = {v["sku"] for v in await db.rahaza_model_variants.find({}, {"_id": 0, "sku": 1}).to_list(50000)}
        from utils.variant_ssot import build_variant_sku
        for i, r in enumerate(wb["VARIAN_BARU"].iter_rows(values_only=True, min_row=2), start=2):
            r = tuple(r) + (None,) * (7 - len(r))
            if not r or not r[0]:
                continue
            mcode, ccode, scode = str(r[0]).strip().upper(), str(r[3] or "").strip().upper(), str(r[4] or "").strip().upper()
            if not ccode or not scode:
                continue  # baris belum diisi → dilewati (aman unggah sebagian)
            if mcode not in models:
                errors.append(f"VARIAN_BARU baris {i}: model {mcode} tidak ada")
                continue
            if ccode not in colors:
                errors.append(f"VARIAN_BARU baris {i}: kode_warna {ccode} tidak ada di master warna")
                continue
            if scode not in sizes:
                errors.append(f"VARIAN_BARU baris {i}: ukuran {scode} tidak ada di master ukuran")
                continue
            try:
                harga = _num(r[5]) if r[5] not in (None, "") else 0
            except ValueError:
                errors.append(f"VARIAN_BARU baris {i}: harga_jual bukan angka")
                continue
            sku = build_variant_sku(mcode, ccode, scode)
            if sku in existing:
                continue  # sudah ada → idempoten
            existing.add(sku)
            out["new_variants"].append({"row": i, "sku": sku, "model": models[mcode], "color": colors[ccode], "size": sizes[scode], "harga": harga})
    if "HARGA_JUAL_SKU" in wb.sheetnames:
        fg = {f["code"]: f for f in await db.rahaza_materials.find({"type": "fg"}, {"_id": 0, "id": 1, "code": 1, "retail_price_master": 1, "active": 1}).to_list(20000)}
        variants = {}
        for v in await db.rahaza_model_variants.find({}, {"_id": 0, "id": 1, "sku": 1, "active": 1}).to_list(50000):
            if v["sku"] not in variants or v.get("active") is not False:
                variants[v["sku"]] = v
        for i, r in enumerate(wb["HARGA_JUAL_SKU"].iter_rows(values_only=True, min_row=2), start=2):
            if not r or not r[0]:
                continue
            sku = str(r[0]).strip()
            raw = r[6] if len(r) > 6 else None
            try:
                harga = _num(raw)
            except ValueError:
                txt = str(raw).strip().lower()
                if "tidak dijual" in txt or "tidak jual" in txt or txt in ("stop", "hapus", "nonaktif", "non aktif", "discontinue"):
                    v = variants.get(sku)
                    f = fg.get(sku)
                    if (v and v.get("active") is not False) or (f and f.get("active") is not False):
                        out["sku_deactivate"].append({"row": i, "sku": sku, "variant_id": (v or {}).get("id"), "fg_id": (f or {}).get("id"), "teks": str(raw).strip()})
                    continue  # sudah nonaktif → idempoten
                out["warnings"].append(f"HARGA_JUAL_SKU baris {i}: {sku} harga_jual berisi teks '{str(raw).strip()}' — tidak diubah "
                                       "(tulis 'sudah tidak dijual' agar SKU dinonaktifkan otomatis, atau angka harga)")
                continue
            if harga <= 0:
                continue
            if sku not in fg:
                errors.append(f"HARGA_JUAL_SKU baris {i}: SKU {sku} tidak ada")
                continue
            if abs(float(fg[sku].get("retail_price_master") or 0) - harga) > 1e-9:
                out["sku_prices"].append({"sku": sku, "id": fg[sku]["id"], "harga": harga})
    for sheet, key, coll_q, cols in (("STOK_AWAL_FG", "stock_fg", {"type": "fg"}, STOK_FG_COLS),
                                     ("STOK_AWAL_MATERIAL", "stock_mat", {"type": {"$ne": "fg"}}, STOK_MAT_COLS)):
        if sheet not in wb.sheetnames:
            continue
        items = {m["code"]: m for m in await db.rahaza_materials.find(coll_q, {"_id": 0, "id": 1, "code": 1, "unit": 1}).to_list(30000)}
        li, qi = cols.index("lokasi_kode"), cols.index("qty_awal")
        for i, r in enumerate(wb[sheet].iter_rows(values_only=True, min_row=2), start=2):
            if not r or not r[0]:
                continue
            code = str(r[0]).strip()
            try:
                qty = _num(r[qi]) if len(r) > qi else 0
            except ValueError:
                errors.append(f"{sheet} baris {i}: qty_awal bukan angka")
                continue
            if qty <= 0:
                continue
            loc = str(r[li]).strip() if len(r) > li and r[li] else ""
            if code not in items:
                errors.append(f"{sheet} baris {i}: kode {code} tidak ada di master")
                continue
            if loc not in locs:
                errors.append(f"{sheet} baris {i}: lokasi_kode '{loc}' tidak ada (lihat REF_LOKASI)")
                continue
            out[key].append({"code": code, "material_id": items[code]["id"], "location_id": locs[loc]["id"], "location_code": loc, "qty": qty, "unit": items[code].get("unit")})
    if "GAJI" in wb.sheetnames:
        emps = {e["employee_code"]: e for e in await db.rahaza_employees.find({}, {"_id": 0, "id": 1, "employee_code": 1}).to_list(5000)}
        profiles = {p["employee_id"]: p for p in await db.rahaza_payroll_profiles.find({}, {"_id": 0, "employee_id": 1, "base_rate": 1, "overtime_rate": 1, "pay_scheme": 1}).to_list(5000)}
        for i, r in enumerate(wb["GAJI"].iter_rows(values_only=True, min_row=2), start=2):
            if not r or not r[0]:
                continue
            code = str(r[0]).strip()
            if code not in emps:
                errors.append(f"GAJI baris {i}: kode_karyawan {code} tidak ada")
                continue
            try:
                base = _num(r[3]) if len(r) > 3 else 0
                ot = _num(r[4]) if len(r) > 4 else 0
            except ValueError:
                errors.append(f"GAJI baris {i}: gaji/tarif bukan angka")
                continue
            scheme = (str(r[2]).strip().lower() if len(r) > 2 and r[2] else "") or None
            if scheme and scheme not in ("monthly", "daily", "piece", "weekly"):
                errors.append(f"GAJI baris {i}: skema '{scheme}' tidak dikenal (monthly/daily/weekly/piece)")
                continue
            cur = profiles.get(emps[code]["id"]) or {}
            changed = (base > 0 and abs(base - float(cur.get("base_rate") or 0)) > 1e-9) or \
                      (ot > 0 and abs(ot - float(cur.get("overtime_rate") or 0)) > 1e-9) or \
                      (scheme and scheme != cur.get("pay_scheme"))
            if changed:
                out["salaries"].append({"employee_code": code, "employee_id": emps[code]["id"], "base": base, "ot": ot, "scheme": scheme})
    return out


async def apply_new_variants(db, new_variants: list[dict], user: dict | None) -> dict:
    """Buat varian + barang jadi dari sheet VARIAN_BARU (idempoten: SKU yang sudah ada dilewati saat parse)."""
    from utils.variant_ssot import ensure_fg_material
    now = datetime.now(timezone.utc)
    created = []
    for v in new_variants:
        m, c, s = v["model"], v["color"], v["size"]
        doc = {"id": str(uuid.uuid4()), "model_id": m["id"], "model_code": m["code"], "model_name": m.get("name"),
               "size_id": s["id"], "size_code": s["code"], "color_id": c.get("id"), "color_code": c["code"], "color_name": c.get("name") or c["code"],
               "color_hex": c.get("hex"), "sku": v["sku"], "barcode": "", "notes": "", "active": True, "created_at": now, "updated_at": now,
               "created_from": "varian_baru_excel", "created_by": str((user or {}).get("id") or "system")}
        await db.rahaza_model_variants.insert_one(doc)
        fg = await ensure_fg_material(db, doc, user=user)
        if v["harga"] > 0 and fg:
            await db.rahaza_materials.update_one({"id": fg["id"]}, {"$set": {"retail_price_master": v["harga"], "updated_at": now}})
        created.append(v["sku"])
    return {"variants_created": len(created), "variants_created_skus": created[:50]}


async def apply_extra(db, parsed: dict, user: dict | None) -> dict:
    from core import stock_service
    now = datetime.now(timezone.utc)
    actor = {"id": str((user or {}).get("id") or "system"), "email": (user or {}).get("email", "")}
    for p in parsed.get("sku_prices") or []:
        await db.rahaza_materials.update_one({"id": p["id"]}, {"$set": {"retail_price_master": p["harga"], "updated_at": now}})
    n_deact = 0
    for d in parsed.get("sku_deactivate") or []:
        if d.get("variant_id"):
            from core import product_master as pm
            await pm.deactivate_catalog_items_for_variant(db, d["variant_id"])
            note = f"Dinonaktifkan dari Excel HARGA_JUAL_SKU: {d.get('teks', '')}"
            await db.rahaza_model_variants.update_many({"sku": d["sku"]}, {"$set": {"active": False, "updated_at": now, "notes": note}})
            v = await db.rahaza_model_variants.find_one({"id": d["variant_id"]}, {"_id": 0, "model_id": 1, "size_id": 1, "color_code": 1})
            if v:  # BOM ikut nonaktif — kalau tidak, sinkron master saat start akan membuat ulang variannya
                await db.rahaza_boms.update_many({"model_id": v["model_id"], "size_id": v.get("size_id"), "color_code": v.get("color_code")},
                                                 {"$set": {"active": False, "updated_at": now, "deactivated_reason": note}})
        if d.get("fg_id"):
            await db.rahaza_materials.update_one({"id": d["fg_id"]}, {"$set": {"active": False, "updated_at": now}})
        n_deact += 1
    n_stock = 0
    for key in ("stock_fg", "stock_mat"):
        for s in parsed.get(key) or []:
            already = await db.rahaza_material_movements.find_one({"type": "opening", "material_id": s["material_id"], "to_location_id": s["location_id"]}, {"_id": 0, "id": 1})
            if already:
                continue  # stok awal per item+lokasi hanya sekali — koreksi lewat Penyesuaian Stok
            await stock_service.add(s["material_id"], s["location_id"], s["qty"], ref={"source": "opening_stock", "reason": "Stok awal go-live (Excel)"}, actor=actor, db=db)
            await db.rahaza_material_movements.insert_one({"id": str(uuid.uuid4()), "created_at": now, "timestamp": now, "created_by": actor["id"],
                                                           "type": "opening", "material_id": s["material_id"], "qty": s["qty"], "from_location_id": None,
                                                           "to_location_id": s["location_id"], "ref_type": "opening_stock", "ref_id": None,
                                                           "notes": "Stok awal go-live dari Excel (nilai GL dari SALDO_AWAL)"})
            n_stock += 1
    for g in parsed.get("salaries") or []:
        upd = {"updated_at": now}
        if g["base"] > 0:
            upd["base_rate"] = g["base"]
        if g["ot"] > 0:
            upd["overtime_rate"] = g["ot"]
        if g["scheme"]:
            upd["pay_scheme"] = g["scheme"]
        await db.rahaza_payroll_profiles.update_one({"employee_id": g["employee_id"]}, {"$set": upd})
    return {"sku_prices_updated": len(parsed.get("sku_prices") or []), "sku_deactivated": n_deact, "opening_stock_rows": n_stock,
            "salaries_updated": len(parsed.get("salaries") or [])}
