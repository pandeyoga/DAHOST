"""core.fin_statements — SSOT perhitungan laporan keuangan turunan buku besar.

Dipakai `routes/rahaza_fin_reports.py` (neraca saldo, neraca lajur, laba rugi 12 bulan,
neraca 12 bulan, arus kas per akun) DAN ekspor Excel supaya layar & berkas tidak pernah beda angka.
Semua angka bersumber dari `rahaza_journal_lines` (hanya jurnal POSTED yang dimirror ke sana).
"""
from __future__ import annotations

import calendar
import io
from datetime import date

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PL_TYPES = ("REVENUE", "OTHER_INCOME", "COGS", "EXPENSE", "OTHER_EXPENSE")
BS_TYPES = ("ASSET", "LIABILITY", "EQUITY")
CREDIT_TYPES = ("LIABILITY", "EQUITY", "REVENUE", "OTHER_INCOME")
CF_GROUPS = ("OPR", "INV", "PND")
CF_LABELS = {"OPR": "Aktivitas Operasi", "INV": "Aktivitas Investasi", "PND": "Aktivitas Pendanaan"}
MONTH_ID = ["Jan", "Feb", "Mar", "Apr", "Mei", "Jun", "Jul", "Agu", "Sep", "Okt", "Nov", "Des"]
_NOT_PL = {"source_module": {"$ne": "year_end_close"}}  # M-09: jurnal penutup bukan L/R


def default_cash_flow_group(acc: dict) -> str:
    """Kelompok arus kas bawaan dari kode/tipe akun (dipakai bila akun belum diberi tag).
    1-2xxx aset tetap → INV · 2-2xxx liabilitas jangka panjang & ekuitas → PND · lainnya → OPR."""
    code = acc.get("code") or ""
    t = (acc.get("type") or "").upper()
    if code.startswith("1-2") or (acc.get("flags") or {}).get("is_fixed_asset") or (acc.get("flags") or {}).get("is_accum_dep"):
        return "INV"
    if t == "EQUITY" or code.startswith("2-2") or (acc.get("flags") or {}).get("is_long_term"):
        return "PND"
    return "OPR"


async def ensure_cash_flow_groups(db) -> int:
    """Backfill `cash_flow_group` untuk akun yang belum punya tag (idempoten). Return jumlah diisi."""
    n = 0
    async for acc in db.rahaza_coa_accounts.find({"cash_flow_group": {"$nin": list(CF_GROUPS)}}, {"_id": 0}):
        await db.rahaza_coa_accounts.update_one({"id": acc["id"]}, {"$set": {"cash_flow_group": default_cash_flow_group(acc)}})
        n += 1
    return n


async def cash_account_codes(db) -> set:
    """Akun kas/bank/dompet = leaf di bawah 1-1100 (Kas) / 1-1200 (Bank), flag is_cash/is_bank,
    atau gl_account_code rekening di Kas & Bank."""
    accs = {a["code"]: a async for a in db.rahaza_coa_accounts.find({}, {"_id": 0, "code": 1, "parent_code": 1, "flags": 1, "is_group": 1})}

    def under_cash(code: str) -> bool:
        seen = 0
        while code and seen < 8:
            if code in ("1-1100", "1-1200"):
                return True
            code = (accs.get(code) or {}).get("parent_code")
            seen += 1
        return False

    out = set()
    for code, a in accs.items():
        if a.get("is_group"):
            continue
        fl = a.get("flags") or {}
        if fl.get("is_cash") or fl.get("is_bank") or under_cash(code):
            out.add(code)
    async for r in db.rahaza_cash_accounts.find({"gl_account_code": {"$ne": None}}, {"_id": 0, "gl_account_code": 1}):
        if r.get("gl_account_code"):
            out.add(r["gl_account_code"])
    return out


async def _sum_by_account(db, match: dict) -> dict:
    pipe = [{"$match": match}, {"$group": {"_id": "$account_code", "debit": {"$sum": "$debit"}, "credit": {"$sum": "$credit"}}}]
    return {r["_id"]: r async for r in db.rahaza_journal_lines.aggregate(pipe)}


# ═══════════════════════════════════════════════════════════════════════════
# NERACA SALDO (inti) + NERACA LAJUR
# ═══════════════════════════════════════════════════════════════════════════
async def compute_trial_balance(db, from_date: str | None, to_date: str, show_zero: bool = False) -> dict:
    match_period = {"date": {"$lte": to_date}}
    if from_date:
        match_period["date"]["$gte"] = from_date
    accounts = await db.rahaza_coa_accounts.find({}, {"_id": 0}).sort("code", 1).to_list(5000)
    period_rows = await _sum_by_account(db, match_period)
    opening_rows = await _sum_by_account(db, {"date": {"$lt": from_date}}) if from_date else {}

    rows, totals = [], {k: 0.0 for k in ("opening_debit", "opening_credit", "period_debit", "period_credit", "end_debit", "end_credit")}
    for acc in accounts:
        if acc.get("is_group"):
            continue
        code = acc["code"]
        op = opening_rows.get(code, {})
        pr = period_rows.get(code, {})
        opening_net = float(op.get("debit") or 0) - float(op.get("credit") or 0)
        opening_debit, opening_credit = (opening_net, 0.0) if opening_net >= 0 else (0.0, -opening_net)
        period_debit, period_credit = round(float(pr.get("debit") or 0), 2), round(float(pr.get("credit") or 0), 2)
        end_net = opening_net + period_debit - period_credit
        end_debit, end_credit = (round(end_net, 2), 0.0) if end_net >= 0 else (0.0, round(-end_net, 2))
        has_amount = (opening_debit + opening_credit + period_debit + period_credit + end_debit + end_credit) != 0
        if not has_amount and (not show_zero or not acc.get("active", True)):
            continue
        row = {"code": code, "name": acc["name"], "type": acc["type"], "normal_balance": acc.get("normal_balance"),
               "active": acc.get("active", True),
               "opening_debit": round(opening_debit, 2), "opening_credit": round(opening_credit, 2),
               "period_debit": period_debit, "period_credit": period_credit,
               "end_debit": end_debit, "end_credit": end_credit}
        rows.append(row)
        for k in totals:
            totals[k] += row[k]
    totals = {k: round(v, 2) for k, v in totals.items()}
    return {"meta": {"from": from_date, "to": to_date}, "rows": rows, "totals": totals,
            "balanced": totals["end_debit"] == totals["end_credit"]}


async def compute_worksheet(db, from_date: str | None, to_date: str, show_zero: bool = False) -> dict:
    """Neraca lajur ala sheet N-Lajur: saldo awal · mutasi · saldo akhir · kolom L/R · kolom Neraca.
    Laba bersih = Σ kredit L/R − Σ debit L/R, ditambahkan sebagai baris penyeimbang (debit L/R + kredit Neraca)."""
    tb = await compute_trial_balance(db, from_date, to_date, show_zero)
    rows = []
    t = {k: 0.0 for k in ("pl_debit", "pl_credit", "bs_debit", "bs_credit")}
    for r in tb["rows"]:
        is_pl = r["type"] in PL_TYPES
        row = dict(r, statement="pl" if is_pl else "bs",
                   pl_debit=r["end_debit"] if is_pl else 0.0, pl_credit=r["end_credit"] if is_pl else 0.0,
                   bs_debit=0.0 if is_pl else r["end_debit"], bs_credit=0.0 if is_pl else r["end_credit"])
        rows.append(row)
        for k in t:
            t[k] += row[k]
    net_income = round(t["pl_credit"] - t["pl_debit"], 2)
    balancing = {"label": "Laba Bersih" if net_income >= 0 else "Rugi Bersih", "net_income": net_income,
                 "pl_debit": net_income if net_income > 0 else 0.0, "pl_credit": -net_income if net_income < 0 else 0.0,
                 "bs_debit": -net_income if net_income < 0 else 0.0, "bs_credit": net_income if net_income > 0 else 0.0}
    grand = {"pl_debit": round(t["pl_debit"] + balancing["pl_debit"], 2), "pl_credit": round(t["pl_credit"] + balancing["pl_credit"], 2),
             "bs_debit": round(t["bs_debit"] + balancing["bs_debit"], 2), "bs_credit": round(t["bs_credit"] + balancing["bs_credit"], 2)}
    return {"meta": tb["meta"], "rows": rows,
            "totals": {**tb["totals"], **{k: round(v, 2) for k, v in t.items()}},
            "balancing": balancing, "grand_totals": grand,
            "balanced": tb["balanced"] and grand["pl_debit"] == grand["pl_credit"] and grand["bs_debit"] == grand["bs_credit"]}


# ═══════════════════════════════════════════════════════════════════════════
# LABA RUGI 12 BULAN · NERACA 12 BULAN
# ═══════════════════════════════════════════════════════════════════════════
def _months(year: int) -> list[str]:
    return [f"{year}-{m:02d}" for m in range(1, 13)]


def _month_end(year: int, m: int) -> str:
    return date(year, m, calendar.monthrange(year, m)[1]).isoformat()


async def compute_profit_loss_monthly(db, year: int) -> dict:
    months = _months(year)
    pipe = [{"$match": {"date": {"$gte": f"{year}-01-01", "$lte": f"{year}-12-31"}, **_NOT_PL}},
            {"$group": {"_id": {"a": "$account_code", "m": {"$substr": ["$date", 0, 7]}},
                        "debit": {"$sum": "$debit"}, "credit": {"$sum": "$credit"}}}]
    cell = {}
    async for r in db.rahaza_journal_lines.aggregate(pipe):
        cell[(r["_id"]["a"], r["_id"]["m"])] = (float(r["debit"] or 0), float(r["credit"] or 0))
    accounts = await db.rahaza_coa_accounts.find({"type": {"$in": list(PL_TYPES)}, "is_group": False}, {"_id": 0}).sort("code", 1).to_list(5000)
    groups = {k: {"label": lbl, "accounts": [], "months": [0.0] * 12, "total": 0.0} for k, lbl in (
        ("revenue", "Pendapatan"), ("cogs", "Harga Pokok Penjualan (HPP)"), ("expense", "Beban Operasional"),
        ("other_income", "Pendapatan Lain-lain"), ("other_expense", "Beban Lain-lain"))}
    key_of = {"REVENUE": "revenue", "COGS": "cogs", "EXPENSE": "expense", "OTHER_INCOME": "other_income", "OTHER_EXPENSE": "other_expense"}
    for acc in accounts:
        vals = []
        for m in months:
            d, c = cell.get((acc["code"], m), (0.0, 0.0))
            vals.append(round(c - d, 2) if acc["type"] in CREDIT_TYPES else round(d - c, 2))
        if not any(vals) and not acc.get("active", True):
            continue
        g = groups[key_of[acc["type"]]]
        g["accounts"].append({"code": acc["code"], "name": acc["name"], "type": acc["type"], "months": vals, "total": round(sum(vals), 2)})
        g["months"] = [round(a + b, 2) for a, b in zip(g["months"], vals)]
    for g in groups.values():
        g["total"] = round(sum(g["months"]), 2)
    gm, cm, em, oim, oem = (groups[k]["months"] for k in ("revenue", "cogs", "expense", "other_income", "other_expense"))
    gross = [round(a - b, 2) for a, b in zip(gm, cm)]
    operating = [round(a - b, 2) for a, b in zip(gross, em)]
    net = [round(a + b - c, 2) for a, b, c in zip(operating, oim, oem)]
    cum = []
    run = 0.0
    for v in net:
        run = round(run + v, 2)
        cum.append(run)
    lines = {"gross_profit": gross, "operating_income": operating, "net_income": net, "net_income_cumulative": cum}
    return {"meta": {"year": year, "months": months, "month_labels": MONTH_ID}, "groups": groups,
            "lines": {k: {"months": v, "total": round(sum(v), 2) if k != "net_income_cumulative" else cum[-1]} for k, v in lines.items()}}


async def compute_balance_sheet_monthly(db, year: int) -> dict:
    """Saldo akun neraca per AKHIR tiap bulan (kumulatif seluruh histori ≤ akhir bulan), termasuk
    Laba/Rugi Tahun Berjalan komputasi supaya Aset = Liabilitas + Ekuitas tiap kolom (sama dgn /balance-sheet)."""
    months = _months(year)
    pipe = [{"$match": {"date": {"$lte": f"{year}-12-31"}}},
            {"$group": {"_id": {"a": "$account_code", "m": {"$cond": [{"$lt": ["$date", f"{year}-01-01"]}, "open", {"$substr": ["$date", 0, 7]}]}},
                        "debit": {"$sum": "$debit"}, "credit": {"$sum": "$credit"}}}]
    per_acc: dict = {}
    async for r in db.rahaza_journal_lines.aggregate(pipe):
        per_acc.setdefault(r["_id"]["a"], {})[r["_id"]["m"]] = float(r["debit"] or 0) - float(r["credit"] or 0)
    accounts = await db.rahaza_coa_accounts.find({"is_group": False}, {"_id": 0}).sort("code", 1).to_list(5000)
    from routes.rahaza_fin_reports import (
        _bs_type,
    )
    sections = {"assets": {"label": "Aset", "accounts": [], "months": [0.0] * 12},
                "liabilities": {"label": "Liabilitas", "accounts": [], "months": [0.0] * 12},
                "equity": {"label": "Ekuitas", "accounts": [], "months": [0.0] * 12}}
    earnings = [0.0] * 12
    for acc in accounts:
        buckets = per_acc.get(acc["code"])
        if not buckets:
            continue
        run = buckets.get("open", 0.0)
        cum = []
        for m in months:
            run += buckets.get(m, 0.0)
            cum.append(run)
        t = _bs_type(acc)
        if t in PL_TYPES:
            earnings = [round(e - v, 2) for e, v in zip(earnings, cum)]  # laba = Σ(K − D) semua akun L/R
            continue
        vals = [round(v if t == "ASSET" else -v, 2) for v in cum]
        if not any(vals):
            continue
        key = {"ASSET": "assets", "LIABILITY": "liabilities", "EQUITY": "equity"}[t]
        sections[key]["accounts"].append({"code": acc["code"], "name": acc["name"], "months": vals})
        sections[key]["months"] = [round(a + b, 2) for a, b in zip(sections[key]["months"], vals)]
    if any(earnings):
        sections["equity"]["accounts"].append({"code": "LABA_BERJALAN", "name": "Laba/Rugi Tahun Berjalan (Komputasi)", "months": earnings, "computed": True})
        sections["equity"]["months"] = [round(a + b, 2) for a, b in zip(sections["equity"]["months"], earnings)]
    liab_eq = [round(a + b, 2) for a, b in zip(sections["liabilities"]["months"], sections["equity"]["months"])]
    diff = [round(a - b, 2) for a, b in zip(sections["assets"]["months"], liab_eq)]
    return {"meta": {"year": year, "months": months, "month_labels": MONTH_ID, "as_of": [_month_end(year, m) for m in range(1, 13)]},
            "sections": sections, "totals": {"liab_plus_equity": liab_eq, "current_earnings": earnings, "diff": diff},
            "balanced": not any(diff)}


# ═══════════════════════════════════════════════════════════════════════════
# ARUS KAS PER AKUN (tag OPR/INV/PND akun lawan, dari buku besar)
# ═══════════════════════════════════════════════════════════════════════════
async def compute_cash_flow_by_account(db, from_date: str, to_date: str) -> dict:
    cash_codes = await cash_account_codes(db)
    if not cash_codes:
        return {"cash_account_codes": [], "groups": {}, "totals": {}, "note": "Belum ada akun kas/bank di bagan akun."}
    codes = sorted(cash_codes)
    # Jurnal saldo awal (opening_balance) = KAS AWAL, bukan arus kas periode (ala sheet Arus-K)
    in_period = {"date": {"$gte": from_date, "$lte": to_date}, "source_module": {"$ne": "opening_balance"}}
    je_ids = await db.rahaza_journal_lines.distinct("je_id", {"account_code": {"$in": codes}, **in_period})
    opening_cash = 0.0
    async for r in db.rahaza_journal_lines.aggregate([
        {"$match": {"account_code": {"$in": codes}, "$or": [{"date": {"$lt": from_date}}, {"source_module": "opening_balance", "date": {"$lte": to_date}}]}},
        {"$group": {"_id": None, "d": {"$sum": "$debit"}, "c": {"$sum": "$credit"}}}]):
        opening_cash = round(float(r["d"] or 0) - float(r["c"] or 0), 2)
    accounts = {a["code"]: a async for a in db.rahaza_coa_accounts.find({}, {"_id": 0, "code": 1, "name": 1, "type": 1, "cash_flow_group": 1, "flags": 1})}
    groups = {g: {"label": CF_LABELS[g], "items": [], "inflow": 0.0, "outflow": 0.0, "total": 0.0} for g in CF_GROUPS}
    cash_in = cash_out = 0.0
    if je_ids:
        pipe = [{"$match": {"je_id": {"$in": je_ids}}},
                {"$group": {"_id": "$account_code", "debit": {"$sum": "$debit"}, "credit": {"$sum": "$credit"}, "count": {"$sum": 1}}}]
        async for r in db.rahaza_journal_lines.aggregate(pipe):
            code = r["_id"]
            d, c = float(r["debit"] or 0), float(r["credit"] or 0)
            if code in cash_codes:
                cash_in += d
                cash_out += c
                continue
            acc = accounts.get(code) or {"code": code, "name": "(akun tidak ada di bagan)", "type": ""}
            grp = acc.get("cash_flow_group") if acc.get("cash_flow_group") in CF_GROUPS else default_cash_flow_group(acc)
            g = groups[grp]
            # akun lawan dikredit ⇒ kas bertambah (masuk); didebit ⇒ kas berkurang (keluar)
            g["items"].append({"code": code, "name": acc.get("name"), "type": acc.get("type"), "inflow": round(c, 2), "outflow": round(d, 2),
                               "net": round(c - d, 2), "count": r["count"], "tag_source": "account" if acc.get("cash_flow_group") in CF_GROUPS else "default"})
            g["inflow"] += c
            g["outflow"] += d
    for g in groups.values():
        g["items"].sort(key=lambda x: x["code"])
        g["inflow"], g["outflow"] = round(g["inflow"], 2), round(g["outflow"], 2)
        g["total"] = round(g["inflow"] - g["outflow"], 2)
    net = round(sum(g["total"] for g in groups.values()), 2)
    cash_net = round(cash_in - cash_out, 2)
    return {"cash_account_codes": codes, "groups": groups,
            "totals": {"OPR": groups["OPR"]["total"], "INV": groups["INV"]["total"], "PND": groups["PND"]["total"],
                       "net_change_in_cash": net, "cash_gl_net_change": cash_net, "journals": len(je_ids),
                       "opening_cash_gl": opening_cash, "closing_cash_gl": round(opening_cash + cash_net, 2),
                       "unexplained": round(cash_net - net, 2)}}


# ═══════════════════════════════════════════════════════════════════════════
# EKSPOR EXCEL
# ═══════════════════════════════════════════════════════════════════════════
_HDR = Font(bold=True, color="FFFFFF")
_FILL = PatternFill("solid", fgColor="1F4E78")
_BOLD = Font(bold=True)
_NUM = '#,##0;[Red]-#,##0;"-"'


def _sheet(wb, title, header, widths=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill, c.alignment = _HDR, _FILL, Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"
    return ws


def _fmt_numbers(ws, first_col: int):
    for row in ws.iter_rows(min_row=2):
        for c in row[first_col - 1:]:
            if isinstance(c.value, (int, float)):
                c.number_format = _NUM


def _bold_row(ws):
    for c in ws[ws.max_row]:
        c.font = _BOLD


def _finish(wb) -> bytes:
    if "Sheet" in wb.sheetnames and len(wb.sheetnames) > 1:
        del wb["Sheet"]
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xlsx_profit_loss_monthly(data: dict) -> bytes:
    wb = openpyxl.Workbook()
    hdr = ["Kode", "Akun"] + [f"{lbl} {data['meta']['year']}" for lbl in data["meta"]["month_labels"]] + ["Total"]
    ws = _sheet(wb, f"Laba-12 {data['meta']['year']}", hdr, [12, 42] + [14] * 13)
    for key in ("revenue", "cogs", "expense", "other_income", "other_expense"):
        g = data["groups"][key]
        ws.append(["", g["label"].upper()]); _bold_row(ws)
        for a in g["accounts"]:
            ws.append([a["code"], a["name"], *a["months"], a["total"]])
        ws.append(["", f"Total {g['label']}", *g["months"], g["total"]]); _bold_row(ws)
        if key == "cogs":
            ln = data["lines"]["gross_profit"]; ws.append(["", "LABA KOTOR", *ln["months"], ln["total"]]); _bold_row(ws)
        if key == "expense":
            ln = data["lines"]["operating_income"]; ws.append(["", "LABA OPERASI", *ln["months"], ln["total"]]); _bold_row(ws)
    ln = data["lines"]["net_income"]; ws.append(["", "LABA BERSIH", *ln["months"], ln["total"]]); _bold_row(ws)
    ln = data["lines"]["net_income_cumulative"]; ws.append(["", "Laba bersih akumulasi s/d bulan", *ln["months"], ln["total"]])
    _fmt_numbers(ws, 3)
    return _finish(wb)


def xlsx_balance_sheet_monthly(data: dict) -> bytes:
    wb = openpyxl.Workbook()
    hdr = ["Kode", "Akun"] + [f"{lbl} {data['meta']['year']}" for lbl in data["meta"]["month_labels"]]
    ws = _sheet(wb, f"Neraca-12 {data['meta']['year']}", hdr, [12, 42] + [14] * 12)
    for key in ("assets", "liabilities", "equity"):
        s = data["sections"][key]
        ws.append(["", s["label"].upper()]); _bold_row(ws)
        for a in s["accounts"]:
            ws.append([a["code"], a["name"], *a["months"]])
        ws.append(["", f"Total {s['label']}", *s["months"]]); _bold_row(ws)
    ws.append(["", "TOTAL LIABILITAS + EKUITAS", *data["totals"]["liab_plus_equity"]]); _bold_row(ws)
    ws.append(["", "Selisih (Aset − L&E)", *data["totals"]["diff"]])
    _fmt_numbers(ws, 3)
    return _finish(wb)


def xlsx_worksheet(data: dict) -> bytes:
    wb = openpyxl.Workbook()
    hdr = ["Kode", "Akun", "Saldo Awal D", "Saldo Awal K", "Mutasi D", "Mutasi K", "Saldo Akhir D", "Saldo Akhir K",
           "Laba Rugi D", "Laba Rugi K", "Neraca D", "Neraca K"]
    ws = _sheet(wb, "N-Lajur", hdr, [12, 42] + [15] * 10)
    for r in data["rows"]:
        ws.append([r["code"], r["name"], r["opening_debit"], r["opening_credit"], r["period_debit"], r["period_credit"],
                   r["end_debit"], r["end_credit"], r["pl_debit"], r["pl_credit"], r["bs_debit"], r["bs_credit"]])
    t = data["totals"]
    ws.append(["", "JUMLAH", t["opening_debit"], t["opening_credit"], t["period_debit"], t["period_credit"], t["end_debit"], t["end_credit"],
               t["pl_debit"], t["pl_credit"], t["bs_debit"], t["bs_credit"]]); _bold_row(ws)
    b = data["balancing"]
    ws.append(["", b["label"], None, None, None, None, None, None, b["pl_debit"], b["pl_credit"], b["bs_debit"], b["bs_credit"]])
    g = data["grand_totals"]
    ws.append(["", "TOTAL", None, None, None, None, None, None, g["pl_debit"], g["pl_credit"], g["bs_debit"], g["bs_credit"]]); _bold_row(ws)
    _fmt_numbers(ws, 3)
    return _finish(wb)


def xlsx_cash_flow(data: dict) -> bytes:
    wb = openpyxl.Workbook()
    ws = _sheet(wb, "Arus-K (kategori)", ["Aktivitas", "Kategori", "Masuk", "Keluar", "Net", "Jumlah"], [24, 30, 16, 16, 16, 10])
    for act in data["activities"].values():
        for r in act["items"]:
            ws.append([act["label"], r["label"], r["inflow"], r["outflow"], r["net"], r["count"]])
        ws.append([f"Total {act['label']}", None, None, None, act["total"]]); _bold_row(ws)
    t = data["totals"]
    for lbl, key in (("Perubahan bersih kas", "net_change_in_cash"), ("Kas awal", "opening_cash"), ("Kas akhir", "closing_cash")):
        ws.append([lbl, None, None, None, t[key]]); _bold_row(ws)
    _fmt_numbers(ws, 3)
    ba = data.get("by_account") or {}
    ws2 = _sheet(wb, "Arus-K (akun)", ["Kelompok", "Kode", "Akun", "Masuk", "Keluar", "Net", "Sumber tag"], [22, 12, 40, 16, 16, 16, 12])
    for g in (ba.get("groups") or {}).values():
        for r in g["items"]:
            ws2.append([g["label"], r["code"], r["name"], r["inflow"], r["outflow"], r["net"], r["tag_source"]])
        ws2.append([f"Total {g['label']}", None, None, g["inflow"], g["outflow"], g["total"]]); _bold_row(ws2)
    bt = ba.get("totals") or {}
    ws2.append(["Perubahan bersih kas (akun lawan)", None, None, None, None, bt.get("net_change_in_cash", 0)]); _bold_row(ws2)
    ws2.append(["Perubahan kas menurut GL akun kas", None, None, None, None, bt.get("cash_gl_net_change", 0)])
    _fmt_numbers(ws2, 4)
    return _finish(wb)
