"""core.journal_import — format Excel pencatatan/penjurnalan yang bisa DIIMPOR ke sistem.

Permintaan owner 2026-09-24: "data finance tidak bisa sepenuhnya detail dimasukkan ke sistem,
saya butuh format excel pencatatan yang nanti bisa di-import".

Dua sheet yang dibaca (boleh salah satu saja):
  * MUTASI_KAS_BANK — satu baris = satu mutasi rekening (gaya "Lap. Mutasi Rekening" owner):
      tanggal | kode_kas_bank | keterangan | pemasukan | pengeluaran | kode_akun_lawan | referensi
    → jurnal 2 baris: pemasukan = D kas/bank, K akun lawan; pengeluaran = D akun lawan, K kas/bank.
  * JURNAL_UMUM — jurnal bebas banyak baris, dikelompokkan per `no_jurnal`:
      no_jurnal | tanggal | kode_akun | debit | kredit | keterangan | referensi
Setiap baris punya `referensi` opsional; kunci anti-dobel = sha1(isi berkas) + nomor baris,
disimpan di `source_ref` jurnal (`import:<sha8>:<sheet>:<baris>`) → impor ulang berkas yang sama
DILEWATI, bukan digandakan.
"""
from __future__ import annotations

import hashlib
import io
from datetime import date, datetime
from uuid import uuid4

import openpyxl
from openpyxl.styles import Font, PatternFill

SOURCE_MODULE = "journal_import"
SHEET_MUTASI = "MUTASI_KAS_BANK"
SHEET_UMUM = "JURNAL_UMUM"
COLS_MUTASI = ["tanggal", "kode_kas_bank", "keterangan", "pemasukan", "pengeluaran", "kode_akun_lawan", "referensi"]
COLS_UMUM = ["no_jurnal", "tanggal", "kode_akun", "debit", "kredit", "keterangan", "referensi"]
_HDR = Font(bold=True, color="FFFFFF")
_FILL = PatternFill("solid", fgColor="1F4E78")

PETUNJUK = [
    "FORMAT IMPOR PENCATATAN / PENJURNALAN — CV. Dewi Aditya",
    "1. Sheet MUTASI_KAS_BANK: satu baris = satu mutasi rekening. Isi kode_kas_bank (lihat sheet DAFTAR_AKUN, kelompok kas/bank),",
    "   keterangan, lalu SALAH SATU: pemasukan atau pengeluaran (Rp, tanpa titik), dan kode_akun_lawan (akun biaya/pendapatan/",
    "   hutang/piutang/modal). Contoh: token listrik 1.003.000 dari BCA Dekka → kode_kas_bank 1-1211, pengeluaran 1003000, akun lawan 6-xxxx.",
    "   Transfer antar rekening sendiri: tulis SATU baris pengeluaran dari rekening asal dengan kode_akun_lawan = kode rekening tujuan.",
    "2. Sheet JURNAL_UMUM: untuk jurnal yang lebih dari 2 baris. Baris dengan no_jurnal yang sama menjadi satu jurnal; total debit harus = kredit.",
    "3. tanggal: YYYY-MM-DD atau sel bertipe tanggal Excel. referensi: nomor bukti/nota (opsional).",
    "4. Unggah di Portal Keuangan → Akuntansi → Impor Jurnal (Excel): PRATINJAU dulu (kesalahan ditampilkan per baris), baru Posting.",
    "5. Mengunggah berkas yang sama dua kali AMAN — baris yang sudah masuk dilewati (kunci: isi berkas + nomor baris).",
    "6. Kode akun hanya boleh akun DETAIL (bukan header) yang aktif — lihat sheet DAFTAR_AKUN.",
]


def _head(ws, cols):
    ws.append(cols)
    for c in ws[1]:
        c.font, c.fill = _HDR, _FILL
    for i in range(len(cols)):
        ws.column_dimensions[chr(65 + i)].width = 22


async def build_template(db) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PETUNJUK"
    for r in PETUNJUK:
        ws.append([r])
    ws.column_dimensions["A"].width = 130
    ws = wb.create_sheet(SHEET_MUTASI)
    _head(ws, COLS_MUTASI)
    ws.append(["2026-09-24", "1-1219", "KR Otomatis DA Store Karanganyar", 191152, None, "4-1100", ""])
    ws.append(["2026-09-24", "1-1211", "Token listrik IDPEL 32177959940", None, 1003000, "6-2200", ""])
    ws.append(["2026-09-24", "1-1211", "Pinjaman modal tanpa bunga → CV DA Official (transfer antar rekening)", None, 9400000, "1-1219", ""])
    ws = wb.create_sheet(SHEET_UMUM)
    _head(ws, COLS_UMUM)
    ws.append(["JU-001", "2026-09-24", "5-1100", 42629400, None, "Pembelian kain rayon RAYONINDO", "INV-123"])
    ws.append(["JU-001", "2026-09-24", "1-1219", None, 42629400, "Bayar via BCA DA Official", "INV-123"])
    ws = wb.create_sheet("DAFTAR_AKUN")
    _head(ws, ["kode_akun", "nama_akun", "tipe", "kelompok"])
    async for a in db.rahaza_coa_accounts.find({"active": True, "is_group": {"$ne": True}}, {"_id": 0}).sort("code", 1):
        kb = "KAS/BANK" if str(a.get("code", "")).startswith(("1-11", "1-12")) else ""
        ws.append([a.get("code"), a.get("name"), a.get("type"), kb])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _num(v) -> float:
    if v in (None, "", "-"):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace("Rp", "").replace(".", "").replace(",", ".").replace(" ", "")
    return float(s) if s else 0.0


def _date(v) -> str:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v or "").strip()
    if not s:
        raise ValueError("tanggal kosong")
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError(f"tanggal '{s}' tidak dikenal (pakai YYYY-MM-DD)")


async def parse_workbook(db, data: bytes) -> dict:
    sha = hashlib.sha1(data).hexdigest()[:8]
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    acc = {a["code"]: a async for a in db.rahaza_coa_accounts.find({"active": True}, {"_id": 0, "code": 1, "name": 1, "type": 1, "is_group": 1})}
    done = {j["source_ref"] async for j in db.rahaza_journal_entries.find(
        {"source_module": SOURCE_MODULE, "status": {"$ne": "voided"}, "source_ref": {"$regex": f"^import:{sha}:"}}, {"_id": 0, "source_ref": 1})}
    journals, errors, skipped = [], [], 0

    def _acc(code, where):
        code = str(code or "").strip()
        a = acc.get(code)
        if not a:
            errors.append(f"{where}: kode akun '{code}' tidak ada / nonaktif")
            return None
        if a.get("is_group"):
            errors.append(f"{where}: akun '{code}' adalah header — pakai akun anaknya")
            return None
        return a

    def _line(a, d, c, desc):
        return {"account_code": a["code"], "account_name": a.get("name", ""), "account_type": a.get("type", ""),
                "debit": round(float(d), 2), "credit": round(float(c), 2), "description": desc}

    if SHEET_MUTASI in wb.sheetnames:
        ws = wb[SHEET_MUTASI]
        for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not r or all(v in (None, "") for v in r[:6]):
                continue
            where = f"{SHEET_MUTASI} baris {i}"
            ref = f"import:{sha}:mutasi:{i}"
            if ref in done:
                skipped += 1
                continue
            r = list(r) + [None] * (7 - len(r))
            try:
                tgl = _date(r[0])
            except ValueError as e:
                errors.append(f"{where}: {e}")
                continue
            bank, lawan = _acc(r[1], where), _acc(r[5], where)
            masuk, keluar = _num(r[3]), _num(r[4])
            if not bank or not lawan:
                continue
            if (masuk > 0) == (keluar > 0):
                errors.append(f"{where}: isi SALAH SATU pemasukan atau pengeluaran (> 0)")
                continue
            desc = str(r[2] or "").strip() or "Mutasi kas/bank"
            amt = masuk or keluar
            lines = ([_line(bank, amt, 0, desc), _line(lawan, 0, amt, desc)] if masuk
                     else [_line(lawan, amt, 0, desc), _line(bank, 0, amt, desc)])
            journals.append({"key": ref, "sheet": SHEET_MUTASI, "row": i, "date": tgl, "memo": desc,
                             "reference": str(r[6] or "").strip(), "lines": lines, "total": amt})

    if SHEET_UMUM in wb.sheetnames:
        ws = wb[SHEET_UMUM]
        groups: dict = {}
        for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not r or all(v in (None, "") for v in r[:5]):
                continue
            r = list(r) + [None] * (7 - len(r))
            no = str(r[0] or "").strip()
            if not no:
                errors.append(f"{SHEET_UMUM} baris {i}: no_jurnal kosong")
                continue
            groups.setdefault(no, []).append((i, r))
        for no, rows in groups.items():
            first = rows[0][0]
            ref = f"import:{sha}:umum:{first}:{no}"
            if ref in done:
                skipped += 1
                continue
            lines, memo, refno, tgl = [], "", "", None
            ok = True
            for i, r in rows:
                where = f"{SHEET_UMUM} baris {i} ({no})"
                try:
                    tgl = tgl or _date(r[1])
                except ValueError as e:
                    errors.append(f"{where}: {e}")
                    ok = False
                    continue
                a = _acc(r[2], where)
                if not a:
                    ok = False
                    continue
                d, c = _num(r[3]), _num(r[4])
                if (d > 0) == (c > 0):
                    errors.append(f"{where}: isi SALAH SATU debit atau kredit (> 0)")
                    ok = False
                    continue
                desc = str(r[5] or "").strip()
                memo = memo or desc
                refno = refno or str(r[6] or "").strip()
                lines.append(_line(a, d, c, desc))
            if not ok:
                continue
            td, tc = sum(l["debit"] for l in lines), sum(l["credit"] for l in lines)
            if abs(td - tc) > 0.5:
                errors.append(f"{SHEET_UMUM} {no}: Debit Rp {td:,.0f} ≠ Kredit Rp {tc:,.0f}")
                continue
            if len(lines) < 2:
                errors.append(f"{SHEET_UMUM} {no}: minimal 2 baris")
                continue
            journals.append({"key": ref, "sheet": SHEET_UMUM, "row": first, "date": tgl, "memo": memo or no,
                             "reference": refno, "lines": lines, "total": td, "no_jurnal": no})

    if SHEET_MUTASI not in wb.sheetnames and SHEET_UMUM not in wb.sheetnames:
        errors.append(f"Berkas tidak punya sheet {SHEET_MUTASI} maupun {SHEET_UMUM} — unduh template dulu.")
    return {"ok": not errors, "file_sha": sha, "errors": errors, "journals": journals, "skipped_existing": skipped,
            "totals": {"journals": len(journals), "amount": round(sum(j["total"] for j in journals), 2)}}


async def post_journals(db, user: dict, journals: list, filename: str) -> dict:
    from routes.rahaza_journals import _check_period_open, _gen_je_number, _mirror_lines, _validate_lines
    posted, failed = [], []
    for j in journals:
        try:
            d = date.fromisoformat(j["date"])
            total_d, total_c = await _validate_lines(db, j["lines"])
            await _check_period_open(db, d)
            now = datetime.utcnow()
            doc = {"id": str(uuid4()), "je_number": await _gen_je_number(db, d), "date": d.isoformat(),
                   "memo": j["memo"] + (f" [{j['reference']}]" if j.get("reference") else ""),
                   "source_module": SOURCE_MODULE, "source_ref": j["key"], "import_file": filename,
                   "status": "posted", "total_debit": total_d, "total_credit": total_c,
                   "created_at": now, "updated_at": now, "created_by": user["id"], "created_by_name": user.get("name", ""),
                   "posted_at": now, "posted_by": user["id"], "voided_at": None, "voided_by": None,
                   "flags": {"imported": True},
                   "lines": [{"line_id": str(uuid4()), **ln, "cost_center_id": None} for ln in j["lines"]]}
            await db.rahaza_journal_entries.insert_one(doc)
            await _mirror_lines(db, doc)
            posted.append({"je_number": doc["je_number"], "date": doc["date"], "memo": doc["memo"], "total": total_d})
        except Exception as e:  # noqa: BLE001 — satu baris gagal tidak menghentikan sisanya
            detail = getattr(e, "detail", None) or str(e)
            failed.append({"sheet": j["sheet"], "row": j["row"], "memo": j["memo"], "error": str(detail)})
    return {"posted": posted, "failed": failed, "posted_count": len(posted), "failed_count": len(failed)}
