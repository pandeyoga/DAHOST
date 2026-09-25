"""core.laporan_sisa — LAPORAN_SISA_DA.xlsx: semua kekurangan yang TIDAK diimpor, satu sheet per kategori.

Sumber: hasil parse berkas unggahan (kelompok/baris yang dilewati, HARGA_JUAL_SKU, GAJI) + kondisi DB saat ini
(material harga 0, HPP belum tervalidasi, berat model, rekening, toko). Laporan saja — tidak mengubah data.
"""
from __future__ import annotations

import io
import re
from datetime import datetime, timezone

import openpyxl
from openpyxl.styles import Font, PatternFill

from core.bom_fill import CATEGORIES, _SPLIT_RE, clean
from core.master_fill import _head

YELLOW = PatternFill("solid", fgColor="FFF2CC")
DEFAULT_CASH = "1-1201"


def _widths(ws, widths):
    for i, w in enumerate(widths):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w


def _sheet(wb, title, cols, rows, widths=None, fill_cols=()):
    ws = wb.create_sheet(title)
    _head(ws, cols)
    for r in rows:
        ws.append(list(r))
    if widths:
        _widths(ws, widths)
    for c in fill_cols:
        idx = cols.index(c) + 1
        for row in ws.iter_rows(min_row=2, max_row=max(ws.max_row, 2)):
            row[idx - 1].fill = YELLOW
    ws.freeze_panes = "A2"
    return ws


def _rows(wb, name):
    if not wb or name not in wb.sheetnames:
        return []
    return [tuple(r) + (None,) * (10 - len(r)) for r in wb[name].iter_rows(values_only=True, min_row=2) if r and r[0] not in (None, "")]


def sku_price_report(wb) -> list[dict]:
    out = []
    for i, r in enumerate(_rows(wb, "HARGA_JUAL_SKU"), start=2):
        sku, nama, model, warna, ukuran, saran, harga, ket = r[:8]
        if isinstance(harga, str) and clean(harga) and not re.match(r"^[\d.,\s]+$", clean(harga)):
            status = "sudah tidak dijual — perlu verifikasi (laporan saja, SKU tidak dinonaktifkan)"
        elif harga in (None, "", 0) and saran not in (None, "", 0):
            status = "saran belum dikonfirmasi"
        elif harga in (None, "", 0) and "beberapa harga" in clean(ket):
            status = "harga ambigu (model punya beberapa harga)"
        elif harga in (None, "", 0):
            status = "kosong"
        else:
            continue
        out.append({"row": i, "sku": sku, "nama": nama, "model": model, "warna": warna, "ukuran": ukuran, "saran": saran, "isian": clean(harga), "status": status, "ket": clean(ket)})
    return out


def gaji_report(wb) -> list[dict]:
    out = []
    for i, r in enumerate(_rows(wb, "GAJI"), start=2):
        kode, nama, skema, gaji, lembur = r[:5]
        masalah = []
        if skema and clean(skema).lower() not in ("monthly", "daily", "weekly", "piece"):
            masalah.append(f"skema '{skema}' tidak dikenal (monthly/daily/weekly/piece)")
        for label, v in (("gaji_pokok_per_periode", gaji), ("tarif_lembur_per_jam", lembur)):
            if isinstance(v, str) and clean(v) and not re.match(r"^[\d.,\s]+$", clean(v).replace("Rp", "")):
                masalah.append(f"{label} berisi teks '{clean(v)}'")
        if masalah:
            out.append({"row": i, "kode": kode, "nama": nama, "masalah": "; ".join(masalah)})
    return out


async def build_laporan_sisa(db, parsed: dict | None, data: bytes | None) -> tuple[bytes, dict]:
    parsed = parsed or {}
    src = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True) if data else None
    issues = parsed.get("bom_issues") or []
    by_cat: dict[str, list] = {}
    for i in issues:
        by_cat.setdefault(i["kategori"], []).append(i)
    wb = openpyxl.Workbook()
    ring = wb.active
    ring.title = "RINGKASAN"
    ring.append([f"LAPORAN SISA KEKURANGAN DATA — CV. Dewi Aditya · dibuat {datetime.now(timezone.utc):%Y-%m-%d %H:%M} UTC"])
    ring.append(["Hanya BOM_AKSESORIS yang diterapkan. Semua di bawah ini adalah CATATAN untuk diisi/diputuskan, bukan perubahan otomatis."])
    ring.append([])
    ring.append(["#", "Kategori", "Jumlah", "Sheet"])
    for c in ring[4]:
        c.font = Font(bold=True)
    ring.column_dimensions["B"].width = 70
    ring.column_dimensions["D"].width = 28
    summary: dict = {}

    def add_summary(n, label, count, sheet):
        ring.append([n, label, count, sheet])
        summary[sheet] = count

    # 1. model tanpa SKU + VARIAN_BARU
    tanpa_sku: dict[str, dict] = {}
    for i in by_cat.get("model_tanpa_sku", []):
        d = tanpa_sku.setdefault(i["model_code"], {"name": i.get("model_name"), "rows": [], "colors": []})
        d["rows"].append(i["row"])
        for tok in _SPLIT_RE.split(i.get("varian_raw") or ""):
            if tok.strip() and tok.strip().title() not in d["colors"]:
                d["colors"].append(tok.strip().title())
    _sheet(wb, "1_MODEL_TANPA_SKU", ["kode_model", "nama_model", "warna_terdeteksi_dari_kolom_varian", "baris_di_berkas", "tindakan"],
           [(k, v["name"], ", ".join(v["colors"]), ", ".join(map(str, v["rows"])), "buat varian di RnD → Master Produk → Varian (atau isi sheet VARIAN_BARU) lalu unggah ulang BOM")
            for k, v in sorted(tanpa_sku.items())], (14, 24, 44, 24, 70))
    varian_baru = [(k, v["name"], c, "", "", "", "") for k, v in sorted(tanpa_sku.items()) for c in (v["colors"] or [""])]
    _sheet(wb, "VARIAN_BARU", ["kode_model", "nama_model", "warna", "kode_warna", "ukuran", "harga_jual", "keterangan"], varian_baru,
           (14, 24, 20, 12, 14, 14, 30), fill_cols=("kode_warna", "ukuran", "harga_jual"))
    add_summary(1, "Model belum punya varian/SKU (BOM-nya belum bisa masuk)", len(tanpa_sku), "1_MODEL_TANPA_SKU · VARIAN_BARU")

    # 2. kelompok tanpa varian
    ktv = by_cat.get("kelompok_tanpa_varian", [])
    _sheet(wb, "2_KELOMPOK_TANPA_VARIAN", ["baris", "kode_model", "nama_model", "material_kelompok", "varian_tersedia", "isi_kolom_varian"],
           [(i["row"], i["model_code"], i.get("model_name"), i.get("materials"), i.get("varian_tersedia"), "") for i in ktv],
           (8, 14, 24, 70, 60, 30), fill_cols=("isi_kolom_varian",))
    add_summary(2, "Kelompok tanpa kolom varian pada model banyak-kelompok (dilewati)", len({i["model_code"] for i in ktv}), "2_KELOMPOK_TANPA_VARIAN (model)")

    # 3. satuan / kode / qty
    inval = by_cat.get("satuan_tak_valid", []) + by_cat.get("kode_tak_dikenal", []) + by_cat.get("qty_kosong", []) + \
        by_cat.get("model_tak_dikenal", []) + by_cat.get("baris_tanpa_model", [])
    inval.sort(key=lambda i: i["row"])
    _sheet(wb, "3_SATUAN_KODE_TAK_VALID", ["baris", "kategori", "kode_model", "kode_material", "nama_material", "qty_per_pcs", "satuan", "masalah"],
           [(i["row"], CATEGORIES.get(i["kategori"], i["kategori"]), i.get("model_code"), i.get("code"), i.get("name"), i.get("qty_raw"), i.get("unit_raw"), i["detail"]) for i in inval],
           (8, 36, 12, 16, 36, 12, 10, 90))
    add_summary(3, "Baris dengan satuan/kode/qty tidak valid (dilewati)", len(inval), "3_SATUAN_KODE_TAK_VALID")

    # 4. warna tak dikenal
    wtd = by_cat.get("warna_tak_dikenal", [])
    _sheet(wb, "4_WARNA_TAK_DIKENAL", ["baris", "kode_model", "nama_model", "token_varian", "kolom_varian_lengkap", "keterangan"],
           [(i["row"], i["model_code"], i.get("model_name"), i.get("token"), i.get("varian_raw"), i["detail"]) for i in wtd], (8, 12, 24, 20, 40, 90))
    add_summary(4, "Warna/ukuran di kolom varian tidak dikenal", len(wtd), "4_WARNA_TAK_DIKENAL")

    # 5. MATERIAL harga 0 — blocker HPP
    used: dict[str, set] = {}
    async for b in db.rahaza_boms.find({"active": {"$ne": False}, "is_active": True}, {"_id": 0, "model_id": 1, "materials": 1}):
        for ln in b.get("materials") or []:
            if ln.get("material_id"):
                used.setdefault(ln["material_id"], set()).add(b["model_id"])
    mats0 = [m async for m in db.rahaza_materials.find({"active": {"$ne": False}, "type": {"$ne": "fg"}, "code": {"$not": {"$regex": "^CUT-"}},
                                                        "$or": [{"unit_cost": {"$in": [None, 0]}}, {"unit_cost": {"$exists": False}}]}, {"_id": 0}).sort("code", 1)]
    _sheet(wb, "5_MATERIAL_HARGA_0", ["kode", "nama", "tipe", "satuan_dasar", "isi_per_satuan_beli_sekarang", "dipakai_di_n_model_bom", "satuan_beli", "isi_per_satuan_beli", "harga_per_satuan_beli"],
           [(m["code"], m.get("name"), m.get("type"), m.get("unit"), m.get("pack_size") if (m.get("pack_size") or 0) > 1 else "", len(used.get(m["id"], ())), "", "", "") for m in mats0],
           (16, 44, 10, 12, 20, 18, 12, 18, 20), fill_cols=("satuan_beli", "isi_per_satuan_beli", "harga_per_satuan_beli"))
    add_summary(5, "MATERIAL harga 0 — BLOCKER HPP (isi lewat sheet MATERIAL)", len(mats0), "5_MATERIAL_HARGA_0")

    # 6. HARGA_JUAL_SKU (laporan saja)
    hj = sku_price_report(src)
    _sheet(wb, "6_HARGA_JUAL_SKU", ["baris", "sku", "nama", "kode_model", "warna", "ukuran", "status", "harga_saran", "isian_di_berkas", "keterangan"],
           [(h["row"], h["sku"], h["nama"], h["model"], h["warna"], h["ukuran"], h["status"], h["saran"], h["isian"], h["ket"]) for h in hj],
           (8, 26, 32, 12, 14, 10, 52, 14, 22, 40))
    c_saran = sum(1 for h in hj if h["status"].startswith("saran"))
    c_amb = sum(1 for h in hj if h["status"].startswith("harga ambigu"))
    c_stop = sum(1 for h in hj if h["status"].startswith("sudah"))
    add_summary(6, f"HARGA_JUAL_SKU: {c_saran} saran belum dikonfirmasi · {c_amb} ambigu · {c_stop} 'sudah tidak dijual' (laporan saja, SKU TIDAK dinonaktifkan)",
                len(hj), "6_HARGA_JUAL_SKU")

    # 7. lainnya: berat, toko, rekening, gaji
    lain = []
    async for m in db.rahaza_models.find({"active": {"$ne": False}, "$or": [{"weight_gram": {"$in": [None, 0]}}, {"weight_gram": {"$exists": False}}]}, {"_id": 0, "code": 1, "name": 1}).sort("code", 1):
        lain.append(("MODEL", m["code"], m.get("name"), "berat_gram kosong", "sheet MODEL kolom berat_gram"))
    n_berat = len(lain)
    toko_all = await db.marketing_platform_accounts.find({"status": {"$ne": "archived"}}, {"_id": 0, "account_code": 1, "account_name": 1, "coa_cash_code": 1}).sort("account_code", 1).to_list(100)
    toko_def = [t for t in toko_all if (t.get("coa_cash_code") or DEFAULT_CASH) == DEFAULT_CASH]
    for t in toko_def:
        lain.append(("TOKO", t["account_code"], t.get("account_name"), f"rekening pencairan masih default {DEFAULT_CASH}", "sheet TOKO kolom rekening_pencairan_kode_akun"))
    rek = [a async for a in db.rahaza_cash_accounts.find({"active": {"$ne": False}, "$or": [{"account_number": {"$in": [None, ""]}}, {"account_number": {"$exists": False}}]},
                                                          {"_id": 0, "coa_code": 1, "gl_account_code": 1, "code": 1, "name": 1})]
    for a in rek:
        lain.append(("REKENING", a.get("gl_account_code") or a.get("coa_code") or a.get("code"), a.get("name"), "no_rekening / atas_nama kosong", "sheet REKENING"))
    gj = gaji_report(src)
    for g in gj:
        lain.append(("GAJI", g["kode"], g["nama"], g["masalah"], f"sheet GAJI baris {g['row']}"))
    _sheet(wb, "7_LAINNYA", ["sheet", "kode", "nama", "kekurangan", "tempat_pengisian"], lain, (12, 16, 30, 60, 44))
    add_summary(7, f"Lainnya: {n_berat} MODEL berat kosong · TOKO {len(toko_def)} dari {len(toko_all)} rekening default · {len(rek)} REKENING kosong · {len(gj)} GAJI format salah",
                len(lain), "7_LAINNYA")

    # 8. HPP belum tervalidasi
    hv = [m async for m in db.rahaza_models.find({"active": {"$ne": False}, "hpp_validation.status": "belum_tervalidasi"}, {"_id": 0, "code": 1, "name": 1, "hpp": 1, "hpp_validation": 1}).sort("code", 1)]
    _sheet(wb, "8_HPP_BELUM_TERVALIDASI", ["kode_model", "nama_model", "hpp_sistem", "alasan"],
           [(m["code"], m.get("name"), float(m.get("hpp") or 0) or None, "; ".join(m["hpp_validation"].get("reasons") or [])) for m in hv], (14, 24, 14, 100))
    add_summary(8, "HPP belum tervalidasi (BOM memakai material harga 0 / kemasan tanpa isi / BOM belum ada)", len(hv), "8_HPP_BELUM_TERVALIDASI")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), summary
